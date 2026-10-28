"""
ỨNG DỤNG XÂY DỰNG MA TRẬN, BẢN ĐẶC TẢ & NGÂN HÀNG CÂU HỎI MÔN SINH HỌC GDPT 2018 - PHIÊN BẢN v6
Hỗ trợ:
 - Tự động cài đặt thư viện và Tự động mở trình duyệt web khi chạy trực tiếp bằng Python.
 - Tải ảnh minh họa cho CÂU HỎI và TỪNG ĐÁP ÁN (A, B, C, D).
 - Đáp án mặc định cho câu hỏi trắc nghiệm luôn chuẩn hóa là "A".
 - TỰ ĐỘNG TẠO ĐỀ THI NGẪU NHIÊN từ Ngân hàng câu hỏi theo Ma trận được chọn.
 - Xuất Đề thi, Đáp án và Bản đặc tả ra file Word (.docx) và Excel (.xlsx).
"""

import os
import sys
import subprocess

def ensure_dependencies():
    """Tự động kiểm tra và cài đặt các thư viện cần thiết nếu máy chưa có."""
    required = {
        "streamlit": "streamlit",
        "pandas": "pandas",
        "openpyxl": "openpyxl",
        "docx": "python-docx"
    }
    missing = []
    for module_name, pip_name in required.items():
        try:
            __import__(module_name)
        except ImportError:
            missing.append(pip_name)
    
    if missing:
        print("=" * 65)
        print("[*] ĐANG TỰ ĐỘNG CÀI ĐẶT THƯ VIỆN CẦN THIẾT:")
        print(f"    Cần cài: {', '.join(missing)}")
        print("    Xin vui lòng đợi trong giây lát...")
        print("=" * 65)
        try:
            subprocess.check_call([sys.executable, "-m", "pip", "install", "--upgrade"] + missing)
            print("[✓] Cai dat thu vien thanh cong!")
        except Exception as e:
            print(f"[!] Lưu ý khi cài đặt: {e}")

ensure_dependencies()

def configure_streamlit():
    """Tự động cấu hình để Streamlit không bao giờ hiện prompt hỏi email."""
    try:
        s_dir = os.path.expanduser("~/.streamlit")
        os.makedirs(s_dir, exist_ok=True)
        cred = os.path.join(s_dir, "credentials.toml")
        if not os.path.exists(cred):
            with open(cred, "w", encoding="utf-8") as cf:
                cf.write("[general]\nemail = \"\"\n")
        cfg = os.path.join(s_dir, "config.toml")
        if not os.path.exists(cfg):
            with open(cfg, "w", encoding="utf-8") as cfgf:
                cfgf.write("[browser]\ngatherUsageStats = false\n")
    except Exception:
        pass

configure_streamlit()

def auto_launch_streamlit():
    """Tự động chuyển từ chạy Python thông thường sang Streamlit và mở trình duyệt."""
    try:
        from streamlit.runtime import exists
        if exists():
            return # Đang chạy chuẩn trong Streamlit engine
    except Exception:
        pass
        
    if os.environ.get("BIO_MATRIX_APP_RUNNING") == "1":
        return

    print("=" * 65)
    print("[*] ĐANG KHỞI ĐỘNG ỨNG DỤNG XÂY DỰNG MA TRẬN & BẢN ĐẶC TẢ SINH HỌC...")
    print("    Trình duyệt web (Chrome / Edge) sẽ tự động mở lên trong giây lát.")
    print("=" * 65)
    
    script_path = os.path.abspath(__file__) if "__file__" in globals() else os.path.abspath(sys.argv[0])
    os.environ["BIO_MATRIX_APP_RUNNING"] = "1"
    
    try:
        cmd = [sys.executable, "-m", "streamlit", "run", script_path]
        subprocess.run(cmd)
        sys.exit(0)
    except Exception as e:
        print(f"[!] Không thể tự động chạy Streamlit: {e}")

if __name__ == "__main__":
    auto_launch_streamlit()

import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import io
import json
import base64
import random
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# --- CẤU HÌNH TRANG STREAMLIT ---
st.set_page_config(
    page_title="SinhHoc Matrix, Spec & Exam Builder 2018",
    page_icon="🧬",
    layout="wide",
    initial_sidebar_state="expanded"
)

# --- CƠ SỞ DỮ LIỆU CHỈ BÁO NĂNG LỰC SINH HỌC ---
COMPETENCY_INDICATORS = {
    "NT1": {"name": "Nhận biết, kể tên, phát biểu đối tượng/khái niệm/quy luật", "level": "Nhận biết", "group": "Nhận thức Sinh học"},
    "NT2": {"name": "Trình bày đặc điểm, vai trò, quá trình sống (lời, sơ đồ, biểu đồ)", "level": "Nhận biết", "group": "Nhận thức Sinh học"},
    "NT3": {"name": "Phân loại đối tượng, hiện tượng sống theo tiêu chí", "level": "Thông hiểu", "group": "Nhận thức Sinh học"},
    "NT4": {"name": "Phân tích đặc điểm đối tượng/quá trình theo logic", "level": "Thông hiểu", "group": "Nhận thức Sinh học"},
    "NT5": {"name": "So sánh, lựa chọn đối tượng, cơ chế, quá trình sống", "level": "Thông hiểu", "group": "Nhận thức Sinh học"},
    "NT6": {"name": "Giải thích mối quan hệ (nguyên nhân-kết quả, cấu tạo-chức năng)", "level": "Thông hiểu", "group": "Nhận thức Sinh học"},
    "NT7": {"name": "Nhận ra và chỉnh sửa điểm sai, phản biện thảo luận", "level": "Vận dụng", "group": "Nhận thức Sinh học"},
    "NT8": {"name": "Tìm từ khóa, dùng thuật ngữ khoa học, kết nối logic thông tin", "level": "Vận dụng", "group": "Nhận thức Sinh học"},
    "TH1": {"name": "Đề xuất vấn đề liên quan đến thế giới sống, đặt câu hỏi", "level": "Thông hiểu", "group": "Tìm hiểu thế giới sống"},
    "TH2": {"name": "Đưa ra phán đoán và xây dựng giả thuyết nghiên cứu", "level": "Thông hiểu", "group": "Tìm hiểu thế giới sống"},
    "TH3": {"name": "Lập kế hoạch thực hiện, lựa chọn phương pháp nghiên cứu", "level": "Thông hiểu", "group": "Tìm hiểu thế giới sống"},
    "TH4": {"name": "Thực hiện kế hoạch, thu thập, xử lý và phân tích dữ liệu thực nghiệm", "level": "Vận dụng", "group": "Tìm hiểu thế giới sống"},
    "TH5": {"name": "Viết, trình bày báo cáo nghiên cứu và thảo luận kết quả", "level": "Vận dụng", "group": "Tìm hiểu thế giới sống"},
    "VD1": {"name": "Giải thích, đánh giá hiện tượng thực tiễn, mô hình công nghệ", "level": "Vận dụng", "group": "Vận dụng kiến thức, kĩ năng"},
    "VD2": {"name": "Đề xuất, thực hiện giải pháp bảo vệ sức khỏe, môi trường, phát triển bền vững", "level": "Vận dụng", "group": "Vận dụng kiến thức, kĩ năng"}
}


DEFAULT_CURRICULUM = {
    "Sinh học 10": {
        "Phần Mở đầu": {
            "Bài 1: Giới thiệu khái quát môn Sinh học": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được đối tượng và các lĩnh vực nghiên cứu của sinh học."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được mục tiêu môn Sinh học và định nghĩa về phát triển bền vững."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được vai trò của sinh học với cuộc sống hằng ngày, với sự phát triển kinh tế – xã hội và những vấn đề toàn cầu; mối quan hệ giữa sinh học với các vấn đề xã hội (đạo đức sinh học, kinh tế, công nghệ)."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Nêu được triển vọng phát triển sinh học trong tương lai, kể tên các ngành nghề liên quan đến sinh học và thành tựu công nghệ của một số ngành nghề chủ chốt."
                }
            ],
            "Bài 2: Phương pháp nghiên cứu và học tập môn Sinh học": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Kể tên và mô tả được các phương pháp nghiên cứu và học tập môn Sinh học (quan sát, làm việc trong phòng thí nghiệm, thực nghiệm khoa học)."
                },
                {
                    "code": "TH1",
                    "level": "Thông hiểu",
                    "text": "Mô tả được các bước trong tiến trình nghiên cứu khoa học: quan sát và đặt câu hỏi, hình thành giả thuyết, thiết kế và tiến hành thí nghiệm, phân tích kết quả và thảo luận, kết luận."
                },
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Lập được kế hoạch nghiên cứu một vấn đề sinh học đơn giản trong học tập hoặc thực tiễn."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Biết cách thu thập, ghi chép và xử lí dữ liệu nghiên cứu thực nghiệm sinh học."
                }
            ],
            "Bài 3: Các cấp độ tổ chức của thế giới sống": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Phát biểu được khái niệm cấp độ tổ chức sống và kể tên các cấp độ tổ chức sống cơ bản (tế bào, cơ thể, quần thể, quần xã, hệ sinh thái)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được các đặc điểm chung của các cấp độ tổ chức sống: tổ chức theo nguyên tắc thứ bậc, hệ thống mở và tự điều chỉnh, liên tục tiến hóa."
                },
                {
                    "code": "NT5",
                    "level": "Thông hiểu",
                    "text": "Dựa vào sơ đồ, phân biệt được cấp độ tổ chức sống cơ bản và cấp độ không cơ bản."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được mối quan hệ thứ bậc giữa các cấp độ tổ chức sống và giải thích được đặc tính nổi trội của hệ thống sống."
                }
            ]
        },
        "Chương 1: Thành phần hóa học của tế bào": {
            "Bài 4: Các nguyên tố hóa học và nước": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Kể tên được các nguyên tố hóa học chính cấu tạo nên tế bào; phân biệt được nguyên tố đa lượng và vi lượng."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Nêu được cấu tạo hóa học và tính chất vật lý, hóa học đặc biệt của phân tử nước (tính phân cực, liên kết hydrogengen)."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được vai trò sinh học của nước đối với tế bào và sinh vật (dung môi hoà tan, môi trường phản ứng, điều hòa nhiệt độ)."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Vận dụng kiến thức về vai trò của nước và các nguyên tố khoáng để giải thích các hiện tượng thực tế và xây dựng chế độ uống nước, bổ sung khoáng chất hợp lí."
                }
            ],
            "Bài 5: Các phân tử sinh học": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm phân tử sinh học; kể tên 4 nhóm đại phân tử chính trong tế bào."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được thành phần cấu tạo, cấu trúc và vai trò của carbohydrate, lipid, protein và nucleic acid."
                },
                {
                    "code": "NT3",
                    "level": "Thông hiểu",
                    "text": "Phân loại được các dạng carbohydrate (đường đơn, đường đôi, đường đa) và các dạng lipid (mỡ, phospholipid, steroid, carotenoid)."
                },
                {
                    "code": "NT5",
                    "level": "Thông hiểu",
                    "text": "So sánh được cấu trúc và chức năng giữa DNA và RNA; phân tích các bậc cấu trúc không gian của protein."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được mối quan hệ giữa cấu tạo và chức năng của các phân tử sinh học; nêu được nguồn thực phẩm cung cấp các phân tử sinh học cho cơ thể."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Vận dụng được kiến thức về các phân tử sinh học để giải thích các hiện tượng và ứng dụng trong thực tiễn (vai trò của DNA trong xác định huyết thống, truy tìm tội phạm; chế độ dinh dưỡng hợp lí)."
                }
            ],
            "Bài 6: Thực hành: Nhận biết một số phân tử sinh học": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Mô tả được nguyên lí và các bước tiến hành thí nghiệm nhận biết đường khử, tinh bột, protein và lipid bằng thuốc thử đặc trưng."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thực hiện được các thí nghiệm nhận biết đường khử (thuốc thử Benedict/Fehling), tinh bột (Lugol/iodine), protein (Biuret) và lipid (nhũ tương)."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Quan sát, ghi chép hiện tượng, giải thích kết quả và viết được báo cáo thực hành nhận biết các phân tử sinh học."
                }
            ]
        },
        "Chương 2: Cấu trúc tế bào": {
            "Bài 7: Tế bào nhân sơ": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Mô tả được các thành phần cấu tạo chính của tế bào nhân sơ: màng sinh chất, thành tế bào, tế bào chất, vùng nhân, vỏ nhầy, roi, lông."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Nêu được kích thước nhỏ đem lại ưu thế gì cho tế bào vi khuẩn (tỉ lệ S/V lớn, trao đổi chất nhanh, sinh sản nhanh)."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được cơ chế tác động của một số loại kháng sinh lên thành peptidoglycan của vi khuẩn để ứng dụng trong y học."
                }
            ],
            "Bài 8: Tế bào nhân thực": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Kể tên và xác định được vị trí của các bào quan trong tế bào nhân thực trên sơ đồ hoặc hình ảnh."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được cấu trúc và chức năng của nhân tế bào, ribosome, lưới nội chất, bộ máy Golgi, ti thể, lục lạp, lysosome, peroxisome, không bào, khung xương tế bào và màng sinh chất."
                },
                {
                    "code": "NT5",
                    "level": "Thông hiểu",
                    "text": "So sánh được cấu trúc giữa tế bào nhân sơ và tế bào nhân thực; giữa tế bào thực vật và tế bào động vật."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được sự phối hợp hoạt động của các bào quan trong quá trình tổng hợp, biến đổi và xuất khẩu protein."
                }
            ],
            "Bài 9: Thực hành: Quan sát tế bào": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Mô tả được quy trình làm tiêu bản tạm thời để quan sát tế bào thực vật (tế bào vảy hành) và tế bào động vật/người (tế bào niêm mạc miệng)."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thực hiện thành thạo kĩ năng sử dụng kính hiển vi quang học; làm được tiêu bản và quan sát rõ hình dạng tế bào."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Vẽ được hình dạng tế bào quan sát được dưới kính hiển vi và viết báo cáo thu hoạch."
                }
            ]
        },
        "Chương 3: Trao đổi chất qua màng và truyền tin tế bào": {
            "Bài 10: Trao đổi chất qua màng tế bào": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm trao đổi chất qua màng tế bào."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được hiện tượng nhập bào và xuất bào thông qua biến dạng của màng tế bào; lấy được ví dụ minh hoạ."
                },
                {
                    "code": "NT3",
                    "level": "Thông hiểu",
                    "text": "Phân biệt được các hình thức vận chuyển các chất qua màng: khuếch tán đơn giản, khuếch tán tăng cường, vận chuyển chủ động. Nêu được ý nghĩa của các hình thức đó."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Vận dụng những hiểu biết về sự vận chuyển các chất qua màng để giải thích một số hiện tượng thực tiễn (bón phân quá mức làm cháy lá cây, bảo quản thực phẩm bằng ướp muối/đường)."
                }
            ],
            "Bài 11: Thực hành: Thí nghiệm co và phản co nguyên sinh": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Mô tả được nguyên lí co nguyên sinh và phản co nguyên sinh ở tế bào thực vật trong các dung dịch ưu trương, nhược trương, đẳng trương."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Làm được tiêu bản tế bào biểu bì lá tía tô, nhỏ dung dịch muối loãng để quan sát hiện tượng co nguyên sinh và nhỏ nước cất để quan sát phản co nguyên sinh."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Ghi chép diễn biến, vẽ hình minh họa và giải thích được nguyên nhân gây co nguyên sinh và phản co nguyên sinh."
                }
            ],
            "Bài 12: Truyền tin tế bào": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm về thông tin giữa các tế bào và các kiểu truyền tin tế bào (tiếp xúc trực tiếp, cận tiết, nội tiết, qua synapse)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được 3 giai đoạn của quá trình truyền tin tế bào: tiếp nhận, truyền tin và đáp ứng."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được cơ chế tương tác đặc hiệu giữa phân tử tín hiệu (phối tử) và thụ thể màng hoặc thụ thể nội bào."
                }
            ]
        },
        "Chương 4: Chuyển hóa năng lượng trong tế bào": {
            "Bài 13: Khái quát về chuyển hóa vật chất và năng lượng": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Phát biểu được khái niệm chuyển hóa năng lượng trong tế bào."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm, cấu trúc và cơ chế tác động của enzyme; trình bày vai trò của enzyme trong quá trình trao đổi chất và chuyển hóa năng lượng."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được cấu tạo và chức năng của ATP về giá trị năng lượng sinh học; trình bày được quá trình tổng hợp và phân giải ATP."
                },
                {
                    "code": "NT5",
                    "level": "Thông hiểu",
                    "text": "Phân biệt được các dạng năng lượng trong tế bào (động năng, thế năng); giải thích năng lượng tích luỹ và sử dụng cho các hoạt động sống là hóa năng."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được các yếu tố ảnh hưởng đến hoạt tính xúc tác của enzyme (nhiệt độ, pH, nồng độ cơ chất, nồng độ enzyme, chất ức chế/hoạt hoá)."
                }
            ],
            "Bài 14: Phân giải và tổng hợp các chất trong tế bào": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm phân giải và tổng hợp các chất trong tế bào."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được các giai đoạn của hô hấp tế bào: đường phân, chu trình Krebs, chuỗi truyền electron."
                },
                {
                    "code": "NT5",
                    "level": "Thông hiểu",
                    "text": "So sánh được hô hấp hiếu khí, phân giải kị khí và lên men về điều kiện, vị trí diễn ra và hiệu quả năng lượng."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được mối quan hệ hai chiều mật thiết giữa quá trình phân giải và quá trình tổng hợp trong tế bào."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Giải thích được hiện tượng mỏi cơ khi vận động mạnh và ứng dụng quá trình lên men trong đời sống (làm sữa chua, muối dưa, ủ rượu)."
                }
            ],
            "Bài 15: Thực hành: Thí nghiệm phân tích ảnh hưởng của một số yếu tố đến hoạt tính của enzyme và kiểm tra hoạt tính của enzyme amylase": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Mô tả được nguyên tắc thí nghiệm đo hoạt tính enzyme catalase hoặc amylase dưới ảnh hưởng của nhiệt độ, pH."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thực hiện được thí nghiệm phân tích ảnh hưởng của nhiệt độ/pH đến hoạt tính enzyme amylase (hoặc catalase từ củ khoai tây/gan)."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Quan sát mức độ đổi màu của iot hoặc tốc độ sủi bọt khí để rút ra kết luận về nhiệt độ/pH tối ưu của enzyme; viết báo cáo thu hoạch."
                }
            ]
        },
        "Chương 5: Chu kì tế bào và phân bào": {
            "Bài 16: Chu kì tế bào và nguyên phân": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm chu kì tế bào và các pha của kì trung gian (G1, S, G2)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Dựa vào sơ đồ, trình bày được các kì của nguyên phân (kì đầu, kì giữa, kì sau, kì cuối) và sự phân chia tế bào chất."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Dựa vào cơ chế nhân đôi và phân li đồng đều của NST để giải thích được quá trình nguyên phân là cơ chế sinh sản của tế bào."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được sự phân chia tế bào một cách không bình thường có thể dẫn đến bệnh ung thư; nêu được các điểm kiểm soát chu kì tế bào."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Trình bày được một số thông tin về bệnh ung thư ở Việt Nam; đề xuất được các biện pháp phòng tránh ung thư trong đời sống hằng ngày."
                }
            ],
            "Bài 17: Giảm phân": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm giảm phân và đối tượng tế bào thực hiện giảm phân (tế bào sinh dục chín)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được diễn biến chính của các kì trong giảm phân I và giảm phân II."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được cơ chế tiếp hợp và trao đổi chéo giữa các nhiễm sắc tử không chị em ở kì đầu I tạo biến dị tổ hợp."
                },
                {
                    "code": "NT5",
                    "level": "Thông hiểu",
                    "text": "So sánh được sự khác nhau giữa nguyên phân và giảm phân về diễn biến NST, số lần phân bào và kết quả tạo tế bào con."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được ý nghĩa sinh học và thực tiễn của quá trình giảm phân kết hợp thụ tinh và nguyên phân trong duy trì bộ NST của loài."
                }
            ],
            "Bài 18: Thực hành: Làm và quan sát tiêu bản quá trình nguyên phân và giảm phân": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Nắm vững quy trình làm tiêu bản tạm thời quan sát tế bào chóp rễ hành (nguyên phân) hoặc bao phấn hoa (giảm phân)."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thực hiện được các thao tác nhuộm mẫu vật, ép tiêu bản và quan sát dưới kính hiển vi quang học."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Nhận diện và vẽ được hình ảnh tế bào ở các kì nguyên phân/giảm phân quan sát được trên tiêu bản."
                }
            ],
            "Bài 19: Công nghệ tế bào": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm và nguyên lí của công nghệ tế bào thực vật và công nghệ tế bào động vật."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được một số thành tựu nổi bật của công nghệ tế bào thực vật (nuôi cấy mô, lai tế bào sinh dưỡng) và tế bào động vật (nhân bản vô tính, liệu pháp tế bào gốc)."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Đánh giá được vai trò, triển vọng và những quan ngại về mặt đạo đức sinh học của công nghệ tế bào trong y học và nông nghiệp."
                }
            ]
        },
        "Chương 6: Sinh học vi sinh vật": {
            "Bài 20: Sự đa dạng và phương pháp nghiên cứu vi sinh vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm vi sinh vật và kể tên các nhóm vi sinh vật chính."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Kể tên và mô tả được một số phương pháp nghiên cứu vi sinh vật thông dụng (quan sát, phân lập, nuôi cấy, định danh)."
                },
                {
                    "code": "NT3",
                    "level": "Thông hiểu",
                    "text": "Phân biệt được 4 kiểu dinh dưỡng ở vi sinh vật dựa vào nguồn năng lượng và nguồn carbon (quang tự dưỡng, hóa tự dưỡng, quang dị dưỡng, hóa dị dưỡng)."
                }
            ],
            "Bài 21: Trao đổi chất, sinh trưởng và sinh sản ở vi sinh vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm sinh trưởng ở quần thể vi sinh vật; trình bày đặc điểm 4 pha sinh trưởng trong nuôi cấy không liên tục (pha tiềm phát, luỹ thừa, cân bằng, suy vong)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được các yếu tố ảnh hưởng đến sinh trưởng của vi sinh vật (nhiệt độ, độ ẩm, pH, ánh sáng, chất hóa học)."
                },
                {
                    "code": "NT3",
                    "level": "Thông hiểu",
                    "text": "Phân biệt được các hình thức sinh sản ở vi sinh vật nhân sơ (phân đôi, nảy chồi, tạo bào tử) và vi sinh vật nhân thực."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Trình bày được ý nghĩa của việc sử dụng kháng sinh đúng cách để ức chế hoặc tiêu diệt vi sinh vật gây bệnh; phân tích tác hại của việc lạm dụng kháng sinh."
                }
            ],
            "Bài 22: Vai trò và ứng dụng của vi sinh vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được vai trò của vi sinh vật trong tự nhiên (phân giải chất hữu cơ, chu trình tuần hoàn vật chất) và đối với con người."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được một số ứng dụng nổi bật của vi sinh vật trong nông nghiệp, công nghiệp thực phẩm, y học và xử lí môi trường."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Vận dụng được hiểu biết về vi sinh vật để giải thích các hiện tượng thực tế và đề xuất biện pháp bảo quản thực phẩm, giữ gìn vệ sinh môi trường."
                }
            ],
            "Bài 23: Thực hành: Một số phương pháp nghiên cứu vi sinh vật thông dụng, tìm hiểu về các sản phẩm công nghệ vi sinh vật và làm một số sản phẩm lên men từ vi sinh vật": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Mô tả được quy trình làm sữa chua, dưa chua hoặc lên men rượu ethylic từ nấm men."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thực hiện đúng các bước trong quy trình lên men sữa chua, dưa chua đảm bảo chất lượng và vệ sinh an toàn thực phẩm."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Giải thích được cơ chế lên men lactic, lên men rượu; đánh giá chất lượng sản phẩm và viết báo cáo thu hoạch."
                }
            ]
        },
        "Chương 7: Virus": {
            "Bài 24: Khái quát về virus": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm, đặc điểm chung và cấu tạo cơ bản của virus (lõi nucleic acid, vỏ capsid, vỏ ngoài)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được 5 giai đoạn nhân lên của virus trong tế bào chủ: hấp phụ, xâm nhập, sinh tổng hợp, lắp ráp và giải phóng."
                },
                {
                    "code": "NT3",
                    "level": "Thông hiểu",
                    "text": "Phân biệt được chu trình sinh tan và chu trình tiềm tan ở bacteriophage (thực khuẩn thể)."
                }
            ],
            "Bài 25: Một số bệnh do virus và các thành tựu nghiên cứu ứng dụng virus": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Kể tên được một số bệnh nguy hiểm do virus gây ra ở người, động vật và thực vật (cúm, sởi, dại, viêm gan, HIV, Covid-19)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được các con đường lây truyền virus và các thành tựu ứng dụng virus trong y học (sản xuất vaccine, liệu pháp gene) và nông nghiệp (chế phẩm trừ sâu sinh học)."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Đề xuất và thực hiện được các biện pháp phòng chống bệnh do virus gây ra cho bản thân, gia đình và cộng đồng."
                }
            ],
            "Bài 26: Thực hành: Điều tra một số bệnh do virus và tuyên truyền phòng chống bệnh": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Xây dựng được kế hoạch và bảng câu hỏi điều tra tình hình mắc một số bệnh do virus phổ biến tại địa phương."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Tiến hành thu thập thông tin, phỏng vấn hoặc điều tra thực địa về tỉ lệ tiêm chủng và ý thức phòng bệnh virus."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Tổng hợp số liệu điều tra, thiết kế được ấn phẩm tuyên truyền (áp phích, tờ rơi, video ngắn) về phòng chống bệnh do virus."
                }
            ]
        }
    },
    "Chuyên đề Sinh học 10": {
        "Chuyên đề 1: Công nghệ tế bào và một số thành tựu": {
            "Bài 1: Công nghệ tế bào thực vật và thành tựu": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm và cơ sở khoa học của công nghệ tế bào thực vật (tính toàn năng, sự phân hoá và phản phân hoá tế bào)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được quy trình nuôi cấy mô tế bào thực vật và thành tựu trong nhân giống vô tính cây trồng quý hiếm."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Mô tả và phân tích được nguyên lí lai tế bào sinh dưỡng (dung hợp tế bào trần) và nuôi cấy hạt phấn/noãn chưa thụ tinh để tạo dòng đơn bội."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Đánh giá được hiệu quả kinh tế và triển vọng của công nghệ tế bào thực vật trong nông nghiệp, lâm nghiệp và bảo tồn nguồn gene."
                }
            ],
            "Bài 2: Tế bào gốc và một số thành tựu": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm tế bào gốc; phân biệt được tế bào gốc toàn năng, vạn năng, đa năng và đơn năng."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được một số thành tựu sử dụng tế bào gốc trong nghiên cứu y học và trị liệu (thay thế mô tổn thương, điều trị bỏng, ung thư máu)."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Thảo luận được triển vọng của liệu pháp tế bào gốc và những vấn đề đạo đức sinh học liên quan đến việc thu nhận tế bào gốc phôi."
                }
            ],
            "Bài 3: Công nghệ tế bào động vật và thành tựu": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Trình bày được các giai đoạn chung của công nghệ tế bào động vật."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Nêu được quy trình và thành tựu của kỹ thuật nhân bản vô tính động vật và thụ tinh trong ống nghiệm."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được triển vọng của liệu pháp gene và công nghệ tế bào động vật trong sản xuất protein sinh dược và y học tái tạo."
                }
            ],
            "Bài 4: Dự án: Tìm hiểu về một số thành tựu công nghệ tế bào": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Xây dựng được đề cương kế hoạch dự án tìm hiểu một thành tựu công nghệ tế bào thực vật hoặc động vật."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thu thập, tổng hợp tư liệu, hình ảnh, video từ internet, sách báo hoặc tham quan cơ sở nuôi cấy mô tại địa phương."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Viết báo cáo dự án và thuyết trình được về quy trình, thành tựu và triển vọng của công nghệ tế bào đã lựa chọn."
                }
            ]
        },
        "Chuyên đề 2: Công nghệ enzyme và ứng dụng": {
            "Bài 5: Khái quát về công nghệ enzyme": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm công nghệ enzyme và các nguồn thu nhận enzyme tự nhiên (thực vật, động vật, vi sinh vật)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được ưu thế của việc sản xuất enzyme từ nguồn vi sinh vật so với động vật và thực vật."
                }
            ],
            "Bài 6: Quy trình công nghệ sản xuất enzyme": [
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được các bước cơ bản trong quy trình sản xuất enzyme quy mô công nghiệp: chuẩn bị nguồn giống, lên men, tách chiết, tinh chế và tạo chế phẩm."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được các điều kiện công nghệ cần kiểm soát trong quá trình nuôi cấy vi sinh vật để thu nhận enzyme tối ưu."
                }
            ],
            "Bài 7: Ứng dụng của enzyme": [
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được một số ứng dụng điển hình của enzyme trong công nghiệp thực phẩm (chế biến tinh bột, làm bánh mì, phô mai), dệt, giấy, y dược và bảo vệ môi trường."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Phân tích được triển vọng của công nghệ enzyme bất động hoá và ứng dụng enzyme trong sản xuất sạch, bền vững."
                }
            ],
            "Bài 8: Dự án: Tìm hiểu về ứng dụng của enzyme": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Lập được kế hoạch tìm hiểu ứng dụng enzyme trong một ngành nghề sản xuất cụ thể hoặc đời sống gia đình."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thu thập mẫu vật, phỏng vấn, khảo sát thực tế quy trình sử dụng enzyme (ví dụ: làm nước tương, nước mắm, chế phẩm tẩy rửa sinh học)."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Báo cáo kết quả dự án dưới dạng poster hoặc bài trình chiếu đa phương tiện."
                }
            ]
        },
        "Chuyên đề 3: Công nghệ vi sinh vật trong xử lí ô nhiễm môi trường": {
            "Bài 9: Vai trò của vi sinh vật trong xử lí ô nhiễm môi trường": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được các tác nhân chính gây ô nhiễm môi trường đất, nước, không khí và vai trò của vi sinh vật như mắt xích phân giải trong sinh quyển."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được cơ chế phân giải các chất hữu cơ, chuyển hóa kim loại nặng và xử lí chất độc hại của hệ vi sinh vật bản địa."
                }
            ],
            "Bài 10: Vi sinh vật trong phân giải các hợp chất làm ô nhiễm môi trường": [
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Mô tả được ứng dụng quá trình phân giải hiếu khí của vi sinh vật trong xử lí ô nhiễm môi trường."
                },
                {
                    "code": "NT3",
                    "level": "Thông hiểu",
                    "text": "Mô tả được ứng dụng quá trình phân giải kị khí và lên men trong xử lí ô nhiễm môi trường (quá trình sinh biogas, ủ phân compost)."
                }
            ],
            "Bài 11: Công nghệ ứng dụng vi sinh vật trong xử lí ô nhiễm môi trường": [
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được một số công nghệ xử lí ô nhiễm bằng vi sinh vật: xử lí đất nhiễm dầu/hóa chất, xử lí nước thải bằng bùn hoạt tính, công nghệ thu nhận khí sinh học biogas, xử lí rác thải rắn hữu cơ."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Đề xuất được mô hình ứng dụng vi sinh vật phù hợp để xử lí rác thải sinh hoạt hữu cơ tại gia đình hoặc trường học."
                }
            ],
            "Bài 12: Dự án: Điều tra công nghệ ứng dụng vi sinh vật trong xử lí ô nhiễm môi trường tại địa phương": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Lập được kế hoạch điều tra thực trạng ô nhiễm và các biện pháp ứng dụng vi sinh vật đang được triển khai tại địa phương."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thực hiện điều tra khảo sát tại bãi rác, trạm xử lí nước thải hoặc các hầm biogas gia đình tại địa phương."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Phân tích số liệu, đánh giá hiệu quả và đề xuất giải pháp cải thiện công nghệ sinh học xử lí môi trường."
                }
            ]
        }
    },
    "Sinh học 11": {
        "Chương 1: Trao đổi chất và chuyển hóa năng lượng ở sinh vật": {
            "Bài 1: Khái quát về trao đổi chất và chuyển hóa năng lượng": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được các dấu hiệu đặc trưng của trao đổi chất và chuyển hóa năng lượng ở sinh vật."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Nêu được các phương thức trao đổi chất và chuyển hóa năng lượng (tự dưỡng và dị dưỡng); nêu được khái niệm tự dưỡng và dị dưỡng."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Dựa vào sơ đồ chuyển hóa năng lượng trong sinh giới, mô tả tóm tắt được ba giai đoạn chuyển hóa năng lượng (tổng hợp, phân giải và huy động năng lượng)."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được vai trò của trao đổi chất và chuyển hóa năng lượng đối với sinh vật; phân tích được vai trò của sinh vật tự dưỡng trong sinh giới."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Trình bày được mối quan hệ giữa trao đổi chất và chuyển hóa năng lượng ở cấp tế bào và cấp cơ thể."
                }
            ],
            "Bài 2: Trao đổi nước và khoáng ở thực vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Kể tên được các con đường hấp thụ nước và khoáng từ đất vào mạch gỗ của rễ (con đường qua gian bào và con đường qua tế bào chất)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Nêu được vai trò sinh lý của một số nguyên tố khoáng thiết yếu (đa lượng: N, P, K, Ca, Mg, S; vi lượng: Fe, Mn, B, Zn, Cu, Mo, Cl)."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được động lực vận chuyển dòng mạch gỗ (áp suất rễ, lực thoát hơi nước ở lá, lực liên kết phân tử nước) và dòng mạch rây trong cây."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được cơ chế đóng mở khí khổng và vai trò của thoát hơi nước đối với đời sống thực vật; phân tích các yếu tố ảnh hưởng đến trao đổi nước và khoáng."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Vận dụng kiến thức thoát hơi nước và dinh dưỡng khoáng để giải thích các hiện tượng thực tế và đề xuất các biện pháp kỹ thuật tưới tiêu, bón phân hợp lí cho cây trồng."
                }
            ],
            "Bài 3: Thực hành: Trao đổi nước và khoáng ở thực vật": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Mô tả được nguyên lí thí nghiệm chứng minh sự hút nước ở rễ, sự thoát hơi nước ở lá và con đường vận chuyển nước trong thân."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thực hiện được thí nghiệm chứng minh thoát hơi nước ở lá (phương pháp dán giấy tẩm CoCl2) và vận chuyển nước ở thân (cành hoa cúc/hoa huệ cắm vào nước màu)."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Quan sát được khí khổng dưới kính hiển vi; giải thích kết quả và hoàn thành báo cáo thực hành."
                }
            ],
            "Bài 4: Quang hợp ở thực vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Phát biểu được khái niệm và viết phương trình tổng quát của quang hợp; nêu vai trò của quang hợp đối với sinh vật và sinh quyển."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Mô tả được cấu tạo của lục lạp phù hợp với chức năng quang hợp; nêu được hệ sắc tố quang hợp (chlorophyll và carotenoid) và vai trò hấp thụ, truyền năng lượng ánh sáng."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được diễn biến của pha sáng (quang phân li nước, tổng hợp ATP, NADPH) và pha tối (chu trình Calvin)."
                },
                {
                    "code": "NT5",
                    "level": "Thông hiểu",
                    "text": "So sánh được đặc điểm quang hợp ở các nhóm thực vật C3, C4 và CAM thích nghi với các điều kiện môi trường sống khác nhau."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được ảnh hưởng của các nhân tố ngoại cảnh (ánh sáng, nồng độ CO2, nhiệt độ, nước, khoáng) đến cường độ quang hợp."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Vận dụng hiểu biết về quang hợp để đề xuất các biện pháp kỹ thuật nông học nhằm nâng cao năng suất cây trồng (tăng diện tích lá, tăng cường độ quang hợp, tuyển chọn giống, trồng xen canh)."
                }
            ],
            "Bài 5: Thực hành: Quang hợp ở thực vật": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Mô tả được nguyên tắc tách chiết sắc tố quang hợp bằng dung môi hữu cơ và nguyên tắc thí nghiệm chứng minh quang hợp giải phóng O2, tạo tinh bột."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thực hiện được thí nghiệm tách chiết và nhận biết các nhóm sắc tố (chlorophyll, carotenoid) từ lá xanh và củ quả có màu; thí nghiệm chứng minh quang hợp tạo O2 (nhánh rong đuôi chó) và tinh bột (lá khoai lang/khoai tây)."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Giải thích được kết quả thí nghiệm và viết báo cáo thu hoạch."
                }
            ],
            "Bài 6: Hô hấp ở thực vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm hô hấp ở thực vật và vai trò của hô hấp đối với cơ thể thực vật (cung cấp năng lượng ATP, nhiệt năng, sản phẩm trung gian)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Kể tên các con đường hô hấp ở thực vật (hô hấp hiếu khí và lên men)."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được các giai đoạn của hô hấp hiếu khí (đường phân, chu trình Krebs, chuỗi truyền electron ti thể); phân tích các yếu tố ảnh hưởng đến hô hấp (nhiệt độ, hàm lượng nước, nồng độ O2, CO2)."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Trình bày được mối quan hệ hai chiều giữa hô hấp và quang hợp ở thực vật."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Vận dụng hiểu biết về các yếu tố ảnh hưởng đến hô hấp để giải thích cơ sở khoa học của các phương pháp bảo quản nông sản (bảo quản khô, lạnh, kín/nồng độ CO2 cao)."
                }
            ],
            "Bài 7: Thực hành: Hô hấp ở thực vật": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Mô tả được nguyên lí thí nghiệm chứng minh hạt nảy mầm hô hấp thải CO2, hút O2 và toả nhiệt."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Lắp ráp được hệ thống dụng cụ và tiến hành thành công thí nghiệm đo nhiệt độ toả ra, phát hiện CO2 (vẩn đục nước vôi trong) trong quá trình hô hấp của hạt nảy mầm."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Ghi chép số liệu, giải thích hiện tượng và viết báo cáo thực hành."
                }
            ],
            "Bài 8: Dinh dưỡng và tiêu hoá ở động vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm tiêu hoá; phân biệt được tiêu hoá nội bào và tiêu hoá ngoại bào."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được các hình thức tiêu hoá ở động vật chưa có cơ quan tiêu hoá, động vật có túi tiêu hoá và động vật có ống tiêu hoá."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Mô tả được quá trình biến đổi cơ học và hóa học của thức ăn trong các cơ quan của ống tiêu hoá ở người."
                },
                {
                    "code": "NT5",
                    "level": "Thông hiểu",
                    "text": "So sánh được đặc điểm cấu tạo ống tiêu hoá và quá trình tiêu hoá thức ăn ở động vật ăn thực vật (động vật nhai lại, động vật có manh tràng phát triển) và động vật ăn thịt."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Vận dụng hiểu biết về tiêu hoá để xây dựng chế độ dinh dưỡng hợp lí, an toàn thực phẩm và phòng chống các bệnh về đường tiêu hoá."
                }
            ],
            "Bài 9: Hô hấp ở động vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm hô hấp ở động vật và các hình thức trao đổi khí (qua bề mặt cơ thể, hệ thống ống khí, mang, phổi)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Kể tên đại diện các nhóm động vật tương ứng với từng hình thức trao đổi khí."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được 4 đặc điểm của bề mặt trao đổi khí giúp khuếch tán khí hiệu quả (rộng, mỏng, ẩm ướt, mạng lưới mao mạch giàu máu)."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được vì sao trao đổi khí ở chim (phổi kèm hệ thống túi khí) và ở cá xương (dòng nước chảy một chiều, song song và ngược chiều dòng máu mao mạch mang) đạt hiệu quả cao nhất."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Phân tích được tác hại của khói thuốc lá, ô nhiễm không khí đối với hệ hô hấp và đề xuất các biện pháp bảo vệ hệ hô hấp, rèn luyện thở sâu."
                }
            ],
            "Bài 10: Tuần hoàn ở động vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Trình bày được khái quát hệ vận chuyển trong cơ thể động vật và nêu được một số dạng hệ vận chuyển ở các nhóm động vật khác nhau."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Dựa vào hình ảnh, sơ đồ, phân biệt được các dạng hệ tuần hoàn ở động vật (hệ tuần hoàn hở và kín; đơn và kép)."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Mô tả được cấu tạo và hoạt động của hệ mạch và quá trình vận chuyển máu trong hệ mạch (huyết áp, vận tốc máu và sự trao đổi chất giữa máu với các tế bào)."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Trình bày được cấu tạo, hoạt động của tim và sự phù hợp giữa cấu tạo và chức năng của tim; giải thích được khả năng tự phát nhịp gây nên tính tự động của tim; nêu được hoạt động tim mạch được điều hoà bằng cơ chế thần kinh và thể dịch."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Phân tích được tác hại của việc lạm dụng rượu, bia đối với sức khoẻ của con người, đặc biệt là hệ tim mạch; trình bày được vai trò của thể dục, thể thao đối với tuần hoàn; kể được các bệnh thường gặp về hệ tuần hoàn; trình bày được một số biện pháp phòng chống các bệnh tim mạch."
                }
            ],
            "Bài 11: Thực hành: Một số thí nghiệm về hệ tuần hoàn": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Nắm vững nguyên lí đo huyết áp bằng huyết áp kế điện tử và cách đếm nhịp tim ở người."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thực hiện được thao tác đo huyết áp, đếm nhịp tim ở người trong các trạng thái nghỉ ngơi và sau khi vận động thể lực; mổ quan sát tim và hệ mạch ở động vật (ếch hoặc cá)."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Nhận xét, so sánh số liệu huyết áp, nhịp tim trước và sau vận động; giải thích nguyên nhân biến đổi và viết báo cáo."
                }
            ],
            "Bài 12: Miễn dịch ở người và động vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được các nguyên nhân bên trong và bên ngoài gây ra các bệnh ở động vật và người; phát biểu được khái niệm miễn dịch và mô tả được khái quát hệ miễn dịch ở người."
                },
                {
                    "code": "NT3",
                    "level": "Thông hiểu",
                    "text": "Phân biệt được miễn dịch không đặc hiệu (hàng rào bảo vệ, đáp ứng viêm, sốt, peptide kháng khuẩn) và miễn dịch đặc hiệu (miễn dịch dịch thể và miễn dịch qua trung gian tế bào)."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được vì sao nguy cơ mắc bệnh ở người rất lớn nhưng xác suất bị bệnh rất nhỏ; trình bày được cơ chế mắc bệnh và cơ chế chống bệnh ở động vật; phân tích được vai trò của việc chủ động tiêm phòng vaccine; giải thích được cơ sở của hiện tượng dị ứng với chất kích thích, thức ăn, cơ sở khoa học của thử phản ứng khi tiêm kháng sinh."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Trình bày được quá trình phá vỡ chức năng của hệ miễn dịch trong cơ thể người bệnh (bệnh tự miễn, ung thư và hội chứng suy giảm miễn dịch mắc phải AIDS); điều tra được tiêm phòng bệnh, dịch bệnh trong trường học hoặc tại địa phương."
                }
            ],
            "Bài 13: Bài tiết và cân bằng nội môi": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Phát biểu được khái niệm bài tiết và trình bày vai trò của bài tiết; kể tên được một số cơ quan tham gia điều hoà cân bằng nội môi (thận, gan, phổi, da) và một số hằng số nội môi cơ thể (pH, huyết áp, áp suất thẩm thấu, đường huyết)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được vai trò của thận trong bài tiết và cân bằng nội môi (lọc máu tạo nước tiểu, điều hoà áp suất thẩm thấu)."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Nêu được khái niệm nội môi, cân bằng nội môi và giải thích được cơ chế chung điều hoà nội môi (bộ phận tiếp nhận kích thích, bộ phận điều khiển, bộ phận thực hiện và liên hệ ngược)."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Nêu được tầm quan trọng của việc xét nghiệm định kì các chỉ số sinh hóa liên quan đến cân bằng nội môi và giải thích được kết quả xét nghiệm (chỉ số đường huyết, creatinin, acid uric)."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Trình bày được các biện pháp bảo vệ thận và các biện pháp phòng tránh một số bệnh liên quan đến thận và bài tiết như suy thận, sỏi thận (uống đủ nước, ăn nhạt, không nhịn tiểu)."
                }
            ]
        },
        "Chương 2: Cảm ứng ở sinh vật": {
            "Bài 14: Khái quát về cảm ứng ở sinh vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm cảm ứng và vai trò của cảm ứng đối với sự tồn tại và phát triển của sinh vật."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Trình bày được cơ chế cảm ứng chung gồm các bộ phận: tiếp nhận kích thích, dẫn truyền thông tin, xử lí thông tin và đáp ứng."
                }
            ],
            "Bài 15: Cảm ứng ở thực vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm và các hình thức hướng động (hướng sáng, hướng trọng lực, hướng hoá, hướng nước, hướng tiếp xúc) và ứng động (ứng động sinh trưởng, ứng động không sinh trưởng) ở thực vật."
                },
                {
                    "code": "NT5",
                    "level": "Thông hiểu",
                    "text": "Phân biệt được hướng động và ứng động; phân biệt cơ chế phản ứng theo hướng kích thích và không theo hướng kích thích."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Giải thích được các hiện tượng cảm ứng trong tự nhiên và nêu được các ứng dụng cảm ứng thực vật trong trồng trọt (làm giàn, bấm ngọn, chiếu sáng nhân tạo kích thích/ức chế ra hoa)."
                }
            ],
            "Bài 16: Thực hành: Cảm ứng ở thực vật": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Mô tả được quy trình bố trí thí nghiệm tính hướng sáng (hộp bìa khoét lỗ) và tính hướng trọng lực (đặt chậu cây nằm ngang) ở hạt đậu nảy mầm."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Tiến hành được thí nghiệm chứng minh tính hướng sáng, hướng trọng lực hoặc quan sát hiện tượng cụp lá ở cây trinh nữ khi va chạm."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Ghi chép số liệu, chụp ảnh kết quả, giải thích hiện tượng cong của thân/rễ và viết báo cáo."
                }
            ],
            "Bài 17: Cảm ứng ở động vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được các hình thức cảm ứng ở các nhóm động vật khác nhau (động vật chưa có hệ thần kinh, hệ thần kinh lưới, chuỗi hạch, ống)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Mô tả được cấu tạo của neuron và synapse; trình bày được các loại thụ thể cảm giác (cơ học, hóa học, nhiệt, đau, điện từ)."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được cơ chế hình thành điện thế nghỉ, điện thế hoạt động và sự lan truyền xung thần kinh trên sợi thần kinh (có và không có bao myelin); cơ chế truyền tin qua synapse hóa học."
                },
                {
                    "code": "NT5",
                    "level": "Thông hiểu",
                    "text": "Phân biệt được phản xạ không điều kiện và phản xạ có điều kiện về nguồn gốc, tính chất, trung gian thần kinh và ý nghĩa thích nghi."
                }
            ],
            "Bài 18: Tập tính ở động vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm tập tính ở động vật; phân biệt được tập tính bẩm sinh và tập tính học được."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Kể tên và mô tả được một số dạng tập tính phổ biến: kiếm ăn, bảo vệ lãnh thổ, sinh sản, di cư, xã hội."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Trình bày được các hình thức học tập ở động vật: quen nhờn, in vết, điều kiện hoá đáp ứng, điều kiện hoá hành động, học ngầm, học giải quyết vấn đề."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Vận dụng hiểu biết về tập tính động vật để giải thích các hiện tượng trong tự nhiên và ứng dụng trong chăn nuôi (huấn luyện gia súc, thú nuôi), bảo vệ mùa màng và an ninh."
                }
            ]
        },
        "Chương 3: Sinh trưởng và phát triển ở sinh vật": {
            "Bài 19: Khái quát về sinh trưởng và phát triển ở sinh vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm và trình bày được các dấu hiệu đặc trưng của sinh trưởng và phát triển ở sinh vật."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được mối quan hệ mật thiết giữa sinh trưởng và phát triển trong đời sống cá thể sinh vật."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Nêu được khái niệm vòng đời và tuổi thọ của sinh vật; phân tích được các yếu tố bên trong (di truyền) và bên ngoài (dinh dưỡng, lối sống, môi trường) ảnh hưởng đến tuổi thọ con người."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Trình bày được một số ứng dụng hiểu biết về vòng đời của sinh vật trong thực tiễn (tiêu diệt côn trùng gây hại ở giai đoạn nhạy cảm) và đề xuất lối sống lành mạnh kéo dài tuổi thọ."
                }
            ],
            "Bài 20: Sinh trưởng và phát triển ở thực vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được đặc điểm sinh trưởng và phát triển ở thực vật; nêu khái niệm mô phân sinh và phân biệt các loại mô phân sinh (đỉnh, bên, lóng)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được quá trình sinh trưởng sơ cấp và sinh trưởng thứ cấp ở thực vật; nêu khái niệm và vai trò của hormone thực vật."
                },
                {
                    "code": "NT3",
                    "level": "Thông hiểu",
                    "text": "Phân biệt được các loại hormone kích thích sinh trưởng (auxin, gibberellin, cytokinin) và hormone ức chế sinh trưởng (abscisic acid, ethylene)."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Trình bày được sự tương quan giữa các hormone thực vật (tương quan chung và tương quan riêng); trình bày quá trình phát triển ở thực vật có hoa và các nhân tố chi phối (quang chu kì, phitocrom, xuân hoá)."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Vận dụng được hiểu biết về sinh trưởng và phát triển ở thực vật để giải thích một số ứng dụng trong thực tiễn (kỹ thuật chiết cành, giâm cành, tỉa cành, bấm ngọn, điều khiển ra hoa trái vụ)."
                }
            ],
            "Bài 21: Thực hành: Bấm ngọn, tỉa cành; xử lí kích thích tố và tính tuổi cây": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Mô tả được cơ sở sinh học của kỹ thuật bấm ngọn tỉa cành (loại bỏ ưu thế đỉnh) và cách đếm vòng gỗ hàng năm để tính tuổi cây thân gỗ."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thực hiện được thao tác bấm ngọn, tỉa cành trên cây trồng phù hợp hoặc đếm vòng gỗ hàng năm để xác định tuổi cây."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Ghi chép sự phát triển của chồi nách sau khi bấm ngọn và hoàn thành báo cáo thực hành."
                }
            ],
            "Bài 22: Sinh trưởng và phát triển ở động vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được đặc điểm sinh trưởng và phát triển ở động vật; trình bày được các giai đoạn chính (giai đoạn phôi và hậu phôi)."
                },
                {
                    "code": "NT3",
                    "level": "Thông hiểu",
                    "text": "Phân biệt được phát triển không qua biến thái và phát triển qua biến thái (biến thái hoàn toàn và biến thái không hoàn toàn); phân tích ý nghĩa thích nghi của biến thái hoàn toàn."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Trình bày được các giai đoạn phát triển của con người từ hợp tử đến cơ thể trưởng thành; phân tích đặc điểm tuổi dậy thì và ứng dụng hiểu biết về tuổi dậy thì trong chăm sóc bản thân."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Nêu được vai trò của một số hormone ảnh hưởng đến sinh trưởng và phát triển ở động vật có xương sống (hormone sinh trưởng GH, thyroxine, estrogen, testosterone) và côn trùng (juvenile, ecdysone)."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Vận dụng hiểu biết về các giai đoạn phát triển để áp dụng chế độ ăn uống, chăm sóc sức khoẻ hợp lí; đề xuất các biện pháp điều khiển sinh trưởng và phát triển ở vật nuôi (chọn giống, cải tạo môi trường sống, tiêm vaccine)."
                }
            ],
            "Bài 23: Thực hành: Quan sát quá trình biến thái ở động vật": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Mô tả được các giai đoạn phát triển qua biến thái của bướm (trứng, sâu non, nhộng, bướm) và của ếch (trứng, nòng nọc, ếch con, ếch trưởng thành)."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Quan sát tiêu bản, mô hình hoặc video về các giai đoạn biến thái ở côn trùng và lưỡng cư."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Vẽ sơ đồ vòng đời, so sánh sự khác biệt về hình thái, môi trường sống giữa các giai đoạn và viết báo cáo."
                }
            ]
        },
        "Chương 4: Sinh sản ở sinh vật": {
            "Bài 24: Khái quát về sinh sản ở sinh vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Phát biểu được khái niệm sinh sản, sinh sản vô tính, sinh sản hữu tính và nêu được các dấu hiệu đặc trưng của sinh sản ở sinh vật."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được vai trò của sinh sản đối với sự tồn tại của loài và phân biệt được các hình thức sinh sản ở sinh vật."
                }
            ],
            "Bài 25: Sinh sản ở thực vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm sinh sản vô tính ở thực vật; kể tên các hình thức sinh sản sinh dưỡng tự nhiên và nhân tạo (giâm cành, chiết cành, ghép cành, nuôi cấy mô)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Mô tả được cấu tạo của hoa; trình bày được quá trình hình thành hạt phấn và túi phôi ở thực vật có hoa."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được quá trình thụ phấn, thụ tinh kép ở thực vật có hoa và sự hình thành hạt, quả sau thụ tinh."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Vận dụng kiến thức sinh sản ở thực vật để giải thích các phương pháp nhân giống vô tính và đề xuất kỹ thuật thụ phấn nhân tạo nâng cao tỉ lệ đậu quả."
                }
            ],
            "Bài 26: Thực hành: Nhân giống vô tính và thụ phấn cho cây": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Nắm vững nguyên lí và quy trình kỹ thuật giâm cành, chiết cành, ghép chồi/ghép cành hoặc thụ phấn nhân tạo cho cây trồng."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thực hiện thành công thao tác giâm cành (rau ngót, mía), chiết cành hoặc thụ phấn chéo nhân tạo cho hoa (ngô, bí đỏ, mướp)."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Theo dõi tỉ lệ ra rễ, tỉ lệ đậu quả và viết báo cáo thu hoạch thực hành."
                }
            ],
            "Bài 27: Sinh sản ở động vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được các hình thức sinh sản vô tính ở động vật (phân đôi, nảy chồi, phân mảnh, trinh sinh)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được các giai đoạn của sinh sản hữu tính ở động vật: hình thành tinh trùng và trứng, thụ tinh, phát triển phôi thai; phân biệt thụ tinh ngoài và thụ tinh trong; đẻ trứng và đẻ con."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Trình bày được cơ chế điều hòa sinh sản bằng hormone ở người (trục dưới đồi - tuyến yên - tuyến sinh dục)."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Vận dụng hiểu biết về sinh sản để giải thích cơ sở khoa học của các biện pháp tránh thai, thụ tinh trong ống nghiệm (IVF) và nâng cao ý thức bảo vệ sức khoẻ sinh sản vị thành niên."
                }
            ]
        },
        "Chương 5: Mối quan hệ giữa các quá trình sinh lý trong cơ thể sinh vật và một số ngành nghề liên quan đến sinh học cơ thể": {
            "Bài 28: Mối quan hệ giữa các quá trình sinh lý trong cơ thể sinh vật": [
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Dựa vào sơ đồ, phân tích được mối quan hệ qua lại chặt chẽ giữa các hệ cơ quan (tiêu hoá, tuần hoàn, hô hấp, bài tiết, thần kinh, nội tiết) trong cơ thể thực vật và động vật."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Chứng minh được cơ thể sinh vật là một hệ thống mở, tự điều chỉnh và toàn vẹn sinh học."
                }
            ],
            "Bài 29: Một số ngành nghề liên quan đến sinh học cơ thể": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được một số ngành nghề liên quan đến sinh học cơ thể: Y học, Dược học, Thú y, Trồng trọt, Chăn nuôi, Thuỷ sản, Lâm nghiệp."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Đánh giá được triển vọng và yêu cầu năng lực của các ngành nghề liên quan đến sinh học cơ thể; liên hệ định hướng nghề nghiệp bản thân."
                }
            ]
        }
    },
    "Chuyên đề Sinh học 11": {
        "Chuyên đề 1: Dinh dưỡng khoáng - Tăng năng suất cây trồng và nông nghiệp sạch": {
            "Bài 1: Nguyên tắc và các biện pháp kỹ thuật sử dụng dinh dưỡng khoáng trong nền nông nghiệp sạch": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm nông nghiệp sạch và các tiêu chuẩn của sản phẩm nông nghiệp sạch."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được các nguyên tắc sử dụng dinh dưỡng khoáng trong việc tăng năng suất cây trồng (đúng loại, đúng liều lượng, đúng thời điểm, bón phối hợp cân đối các yếu tố dinh dưỡng và bón đúng cách)."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Phân tích được một số biện pháp kỹ thuật sử dụng dinh dưỡng khoáng nhằm tạo nên nền nông nghiệp sạch; lấy được ví dụ minh hoạ."
                }
            ],
            "Bài 2: Mô hình thuỷ canh theo hướng phát triển nông nghiệp sạch": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm thuỷ canh, khí canh và các dạng mô hình thuỷ canh phổ biến (thuỷ canh tĩnh, hồi lưu)."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Mô tả được cấu tạo hệ thống và quy trình pha chế dung dịch dinh dưỡng cho mô hình thuỷ canh."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Đánh giá được ưu điểm, hạn chế của phương pháp thuỷ canh và đề xuất mô hình trồng rau thuỷ canh quy mô hộ gia đình."
                }
            ],
            "Bài 3: Dự án: Điều tra tình hình sử dụng phân bón tại địa phương hoặc thực hành trồng cây với các kỹ thuật bón phân phù hợp": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Xây dựng được kế hoạch điều tra thực trạng sử dụng phân bón vô cơ, hữu cơ tại các nông hộ địa phương."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Tiến hành phỏng vấn, khảo sát thực địa tình hình bón phân cho một số loại cây trồng chủ lực."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Tổng hợp số liệu, đánh giá mức độ hợp lí và đề xuất khuyến nghị bón phân an toàn, hiệu quả cho bà con nông dân."
                }
            ],
            "Bài 4: Thực hành: Thí nghiệm chứng minh tác dụng của loại phân bón, cách bón và hàm lượng phân bón đối với cây trồng": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Mô tả được nguyên tắc bố trí thí nghiệm so sánh các công thức phân bón khác nhau trên cùng một loại cây chỉ thị."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Tiến hành gieo trồng, theo dõi chiều cao cây, số lá, khối lượng tươi của cây trồng thí nghiệm qua các tuần."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Xử lí số liệu thống kê sinh học, vẽ đồ thị sinh trưởng và viết báo cáo kết quả thí nghiệm."
                }
            ]
        },
        "Chuyên đề 2: Một số bệnh dịch ở người và cách phòng, chống": {
            "Bài 5: Bệnh dịch và tác nhân gây bệnh ở người": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Kể tên được một số bệnh dịch phổ biến ở người (cúm, tả, sốt rét, lao phổi, sởi, Covid-19, sốt xuất huyết, HIV/AIDS)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Nêu được một số tác nhân gây bệnh (vi khuẩn, virus, nấm, động vật nguyên sinh, giun sán)."
                }
            ],
            "Bài 6: Nguyên nhân lây nhiễm bệnh dịch ở người": [
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được các con đường lây nhiễm bệnh dịch ở người: qua đường hô hấp, tiêu hoá, tiếp xúc máu, đường tình dục và qua động vật trung gian truyền bệnh."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được cơ chế bùng phát ổ dịch trong cộng đồng khi gặp điều kiện thời tiết, mật độ dân cư và vệ sinh môi trường thuận lợi."
                }
            ],
            "Bài 7: Các biện pháp phòng, chống bệnh dịch ở người": [
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được các biện pháp phòng, chống bệnh dịch tương ứng với từng đường lây nhiễm (cách li, đeo khẩu trang, vệ sinh an toàn thực phẩm, diệt vector truyền bệnh, tiêm chủng vaccine)."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Thực hiện được các biện pháp bảo vệ bản thân và gia đình; đề xuất các giải pháp kiểm soát dịch bệnh truyền nhiễm trong trường học."
                }
            ],
            "Bài 8: Dự án: Điều tra một số bệnh dịch phổ biến ở người và tuyên truyền phòng, chống bệnh": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Lập được kế hoạch điều tra tình hình mắc một số bệnh dịch truyền nhiễm tại trạm y tế xã/phường hoặc địa bàn dân cư."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thu thập số liệu thống kê ca mắc, phân tích nguyên nhân và nguy cơ lây lan dịch bệnh."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Thiết kế sản phẩm truyền thông (infographic, video, cẩm nang) để tuyên truyền phòng chống dịch bệnh hiệu quả."
                }
            ]
        },
        "Chuyên đề 3: Vệ sinh an toàn thực phẩm": {
            "Bài 9: Khái quát về vệ sinh an toàn thực phẩm": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được định nghĩa vệ sinh an toàn thực phẩm theo Luật An toàn thực phẩm; phân biệt được thực phẩm sạch, thực phẩm an toàn và thực phẩm ô nhiễm."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Kể tên được các nhóm tác nhân gây ô nhiễm thực phẩm (vi sinh vật, hóa học, chất độc tự nhiên, ô nhiễm vật lý)."
                }
            ],
            "Bài 10: Ngộ độc thực phẩm": [
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được một số nguyên nhân gây ngộ độc thực phẩm (vi sinh vật sinh độc tố, dư lượng thuốc bảo vệ thực vật, kim loại nặng, độc tố tự nhiên trong cá nóc, nấm độc); lấy được ví dụ minh hoạ."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được một số biện pháp phòng ngộ độc thực phẩm (chọn lựa thực phẩm an toàn, bảo quản đúng nhiệt độ, chế biến chín kĩ, không để lẫn thực phẩm sống và chín)."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Vận dụng hiểu biết để xử lí sơ cứu ban đầu khi xảy ra tình huống ngộ độc thực phẩm trong gia đình hoặc trường học."
                }
            ],
            "Bài 11: Dự án: Điều tra về hiện trạng mất vệ sinh an toàn thực phẩm tại địa phương": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Xây dựng kế hoạch và tiêu chí khảo sát vệ sinh an toàn thực phẩm tại căng tin trường học, chợ truyền thống hoặc hàng quán đường phố."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thực hiện điều tra thực trạng bảo quản thực phẩm, nguồn gốc xuất xứ và quy trình chế biến."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Báo cáo kết quả điều tra và đề xuất kiến nghị tới nhà trường và chính quyền địa phương."
                }
            ]
        }
    },
    "Sinh học 12": {
        "Chương 1: Di truyền phân tử": {
            "Bài 1: DNA và cơ chế tái bản DNA": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Mô tả được cấu trúc hóa học và cấu trúc không gian của phân tử DNA theo mô hình Watson - Crick."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được nguyên tắc bán bảo toàn và nguyên tắc bổ sung trong cơ chế tái bản DNA."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được vai trò của các enzyme tham gia vào quá trình tái bản DNA (helicase, RNA polymerase/primase, DNA polymerase, ligase)."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được sự hình thành mạch liên tục và mạch gián đoạn (đoạn Okazaki) trong quá trình tái bản DNA."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Vận dụng nguyên tắc bổ sung và bán bảo toàn để giải các bài tập xác định số lượng, tỉ lệ nucleotide, số liên kết hydrogen và số phân tử DNA con tạo thành."
                }
            ],
            "Bài 2: Gene, hệ gene và quá trình truyền đạt thông tin di truyền": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm gene, hệ gene và cấu trúc chung của một gene cấu trúc (vùng điều hòa, vùng mã hóa, vùng kết thúc); nêu được khái niệm mã di truyền và các đặc điểm của mã di truyền."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được diễn biến các giai đoạn của quá trình phiên mã và dịch mã."
                },
                {
                    "code": "NT3",
                    "level": "Thông hiểu",
                    "text": "Phân biệt được các loại RNA: mRNA, tRNA, rRNA về cấu trúc và chức năng."
                },
                {
                    "code": "NT5",
                    "level": "Thông hiểu",
                    "text": "So sánh được quá trình phiên mã và dịch mã ở sinh vật nhân sơ và sinh vật nhân thực."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được mối quan hệ DNA - RNA - protein - tính trạng; giải thích cơ chế phiên mã ngược ở retrovirus."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Dựa vào bảng mã di truyền, xác định được trình tự amino acid của chuỗi polypeptide tương ứng từ trình tự nucleotide trên mạch khuôn của gene."
                }
            ],
            "Bài 3: Điều hoà biểu hiện gene": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm điều hoà biểu hiện gene và ý nghĩa sinh học của điều hoà biểu hiện gene đối với sinh vật."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Mô tả được cấu trúc của operon Lac ở vi khuẩn E. coli (vùng khởi động P, vùng vận hành O, nhóm gene cấu trúc Z, Y, A và gene điều hoà R)."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Dựa vào sơ đồ, phân tích được cơ chế điều hoà hoạt động của operon Lac khi môi trường có lactose và khi không có lactose."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được các mức độ điều hoà biểu hiện gene ở sinh vật nhân thực (trước phiên mã, phiên mã, sau phiên mã, dịch mã và sau dịch mã)."
                }
            ],
            "Bài 4: Đột biến gene": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm đột biến gene và đột biến điểm."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được vai trò và ý nghĩa của đột biến gene trong tiến hóa, chọn giống và nghiên cứu di truyền."
                },
                {
                    "code": "NT3",
                    "level": "Thông hiểu",
                    "text": "Phân loại và phân biệt được các dạng đột biến điểm: thay thế, thêm, mất một cặp nucleotide."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được nguyên nhân và cơ chế phát sinh đột biến gene qua các thế hệ nhân đôi DNA do tác nhân vật lý (tia UV), hóa học (5-BU, EMS) và sinh học."
                },
                {
                    "code": "NT7",
                    "level": "Vận dụng",
                    "text": "Nhận ra và chỉnh sửa được điểm sai trong sơ đồ cơ chế phát sinh đột biến gene."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Giải thích được hậu quả của các dạng đột biến điểm lên cấu trúc và chức năng của chuỗi polypeptide tương ứng (đột biến đồng nghĩa, sai nghĩa, vô nghĩa, dịch khung)."
                }
            ],
            "Bài 5: Công nghệ gene": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm công nghệ gene, DNA tái tổ hợp và sinh vật biến đổi gene (GMO)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được các bước cơ bản trong kỹ thuật chuyển gene sử dụng plasmid làm thể truyền (tách chiết, cắt nối tạo DNA tái tổ hợp, chuyển vào tế bào nhận, phân lập dòng tế bào mang DNA tái tổ hợp)."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được vai trò của các enzyme giới hạn (restrictase) và enzyme nối (ligase) trong công nghệ tạo DNA tái tổ hợp."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Nêu được các thành tựu ứng dụng của công nghệ gene trong sản xuất dược phẩm (insulin, kháng thể), nông nghiệp (cây trồng chuyển gene kháng sâu, tăng dinh dưỡng) và bảo vệ môi trường; thảo luận về đạo đức sinh học và an toàn sinh học của GMO."
                }
            ],
            "Bài 6: Thực hành: Tách chiết DNA": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Mô tả được nguyên lí tách chiết DNA từ tế bào: phá vỡ màng bằng chất tẩy rửa, kết tủa protein bằng muối và kết tủa DNA bằng cồn lạnh."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thực hiện đúng quy trình tách chiết DNA từ mẫu vật thực vật (chuối, dâu tây) hoặc gan gà/lợn."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Quan sát được búi sợi DNA kết tủa màu trắng đục trong cồn lạnh và viết báo cáo thu hoạch."
                }
            ]
        },
        "Chương 2: Di truyền nhiễm sắc thể": {
            "Bài 7: Cấu trúc và chức năng của nhiễm sắc thể": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Dựa vào sơ đồ (hoặc hình ảnh), trình bày được cấu trúc siêu hiển vi của NST (nucleosome, sợi cơ bản, sợi nhiễm sắc, vùng xếp cuộn, chromatid)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được NST là vật chất di truyền; mô tả được cách sắp xếp các gene trên NST, mỗi gene định vị tại một vị trí xác định gọi là locus."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Trình bày được ý nghĩa của nguyên phân, giảm phân và thụ tinh trong nghiên cứu di truyền; từ đó giải thích sự vận động và truyền thông tin di truyền của các gene qua các thế hệ tế bào và cơ thể."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được sự vận động của NST (tự nhân đôi, phân li, tổ hợp, tái tổ hợp) là cơ sở của sự vận động của gene trong các quy luật di truyền, biến dị tổ hợp và biến dị số lượng NST."
                }
            ],
            "Bài 8: Học thuyết di truyền của Mendel": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được phương pháp nghiên cứu di truyền của Mendel (phương pháp phân tích cơ thể lai) và lí do Mendel chọn đậu Hà Lan làm đối tượng nghiên cứu."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Phát biểu được nội dung quy luật phân li và quy luật phân li độc lập của Mendel."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được cơ sở tế bào học của quy luật phân li và quy luật phân li độc lập dựa trên sự phân li và tổ hợp tự do của các cặp NST tương đồng trong giảm phân và thụ tinh."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Vận dụng quy luật Mendel để giải quyết các bài toán di truyền xác định kiểu gene, kiểu hình, tỉ lệ phân li ở đời con và giải thích các hiện tượng thực tế trong chọn giống."
                }
            ],
            "Bài 9: Mở rộng học thuyết Mendel": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm trội hoàn toàn, trội không hoàn toàn, đồng trội (ví dụ nhóm máu ABO) và gene đa hiệu."
                },
                {
                    "code": "NT3",
                    "level": "Thông hiểu",
                    "text": "Phân biệt được các kiểu tương tác gene không allele: tương tác bổ sung, tương tác cộng gộp và tương tác át chế."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được cơ chế sinh hóa của tương tác gene trong việc hình thành tính trạng và giải thích tính trạng số lượng chịu ảnh hưởng của tương tác cộng gộp."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Vận dụng kiến thức tương tác gene để giải các bài tập xác định tỉ lệ phân li kiểu hình biến đổi so với tỉ lệ Mendel."
                }
            ],
            "Bài 10: Di truyền giới tính và di truyền liên kết với giới tính": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm NST giới tính; di truyền giới tính và các cơ chế xác định giới tính (kiểu XX-XY, XX-XO, ZZ-ZW và do môi trường)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Nêu được bối cảnh ra đời thí nghiệm của Morgan trên ruồi giấm; trình bày được cách bố trí thí nghiệm, qua đó nêu được khái niệm di truyền liên kết với giới tính."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được cơ chế di truyền xác định giới tính; giải thích được tỉ lệ lí thuyết giới tính trong tự nhiên thường là 1:1."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được đặc điểm di truyền của gene nằm trên vùng không tương đồng của NST X (di truyền chéo) và gene trên NST Y (di truyền thẳng); trình bày được quan điểm của bản thân về việc điều khiển giới tính ở người theo ý muốn."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Vận dụng những hiểu biết về di truyền giới tính và liên kết giới tính để giải thích các vấn đề trong thực tiễn (điều khiển giới tính trong chăn nuôi, phát hiện bệnh do rối loạn cơ chế phân li, tổ hợp NST giới tính ở người như mù màu, máu khó đông)."
                }
            ],
            "Bài 11: Liên kết gene và hoán vị gene": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Trình bày được cách bố trí và tiến hành thí nghiệm của Morgan về liên kết gene và hoán vị gene trên ruồi giấm; phát biểu được khái niệm liên kết gene, hoán vị gene và tần số hoán vị gene (f)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Nêu được ý nghĩa của việc lập bản đồ di truyền."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được cơ sở tế bào học và ý nghĩa của liên kết gene (duy trì nhóm tính trạng tốt luôn đi cùng nhau) và hoán vị gene (tiếp hợp và trao đổi chéo giữa các chromatid khác nguồn gốc ở kì đầu giảm phân I tạo biến dị tổ hợp phong phú)."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Nêu được quan điểm của Mendel và Morgan về tính quy luật của hiện tượng di truyền."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Tính được tần số hoán vị gene và xác định trật tự, khoảng cách giữa các gene trên nhiễm sắc thể từ kết quả phép lai phân tích."
                }
            ],
            "Bài 12: Đột biến nhiễm sắc thể": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Phát biểu được khái niệm đột biến NST; phân biệt được đột biến cấu trúc NST và đột biến số lượng NST."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được nguyên nhân và cơ chế phát sinh đột biến cấu trúc NST (mất đoạn, lặp đoạn, đảo đoạn, chuyển đoạn); lấy được ví dụ minh hoạ."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được nguyên nhân và cơ chế phát sinh đột biến số lượng NST (lệch bội và đa bội: tự đa bội, dị đa bội); lấy được ví dụ minh hoạ."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được tác hại của một số dạng đột biến NST đối với sinh vật và con người; phân tích được mối quan hệ giữa di truyền và biến dị."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Trình bày được vai trò của đột biến NST trong tiến hóa, trong chọn giống và trong nghiên cứu di truyền (tạo giống cây trồng tam bội, tứ bội không hạt, năng suất cao)."
                }
            ],
            "Bài 13: Di truyền học người và di truyền y học": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm và vai trò của di truyền học người, di truyền y học."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Nêu được một số phương pháp nghiên cứu di truyền người (nghiên cứu phả hệ, di truyền phân tử, di truyền tế bào); xây dựng được phả hệ để xác định sự di truyền tính trạng trong gia đình."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm y học tư vấn; trình bày được cơ sở của y học tư vấn; trình bày được một số thành tựu và ứng dụng của liệu pháp gene."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được vì sao cần đến cơ sở tư vấn hôn nhân gia đình trước khi kết hôn và sàng lọc trước sinh."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Nêu được khái niệm liệu pháp gene; vận dụng hiểu biết về liệu pháp gene để giải thích việc chữa trị các bệnh di truyền phân tử."
                }
            ],
            "Bài 14: Thực hành: Quan sát một số dạng đột biến nhiễm sắc thể": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Mô tả được các đặc điểm nhận dạng bộ NST bình thường và bộ NST bị đột biến lệch bội (Down, Turner, Klinefelter) hoặc đột biến cấu trúc NST trên ảnh chụp hiển vi."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Quan sát tiêu bản cố định hoặc ảnh chụp tiêu bản hiển vi, đếm số lượng và xác định dạng đột biến NST."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Vẽ hoặc ghi chép đặc điểm kiểu nhân và hoàn thành báo cáo thực hành."
                }
            ]
        },
        "Chương 3: Mở rộng học thuyết di truyền nhân": {
            "Bài 15: Di truyền gene ngoài nhân": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm di truyền gene ngoài nhân (di truyền tế bào chất, di truyền qua ti thể và lục lạp)."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được thí nghiệm của Correns trên cây hoa phấn (Mirabilis jalapa) chứng minh di truyền theo dòng mẹ."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được đặc điểm của di truyền gene ngoài nhân (kết quả phép lai thuận nghịch khác nhau, con luôn mang kiểu hình của mẹ, không tuân theo các tỉ lệ Mendel)."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Giải thích được một số ứng dụng của di truyền tế bào chất trong tạo giống bất thụ đực phục vụ sản xuất hạt lai F1."
                }
            ],
            "Bài 16: Tương tác giữa kiểu gene với môi trường và thành tựu chọn giống": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm mức phản ứng, thường biến và sự mềm dẻo kiểu hình."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được mối quan hệ giữa kiểu gene, môi trường và kiểu hình (kiểu gene quy định khả năng phản ứng, môi trường quy định kiểu hình cụ thể); phân biệt được thường biến và đột biến."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Vận dụng khái niệm mức phản ứng để giải thích vai trò của giống (kiểu gene) và biện pháp kỹ thuật canh tác (môi trường) đối với năng suất vật nuôi, cây trồng."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Trình bày được một số thành tựu chọn giống vật nuôi và cây trồng nổi bật tại Việt Nam."
                }
            ],
            "Bài 17: Thực hành: Thí nghiệm về thường biến ở cây trồng": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Mô tả được nguyên tắc thí nghiệm chứng minh cùng một kiểu gene nhưng nuôi cấy ở các điều kiện ánh sáng, dinh dưỡng khác nhau sẽ cho kiểu hình khác nhau (thường biến)."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Trồng các cành cây có cùng kiểu gene (rau muống, khoai lang) trong các môi trường khác nhau (dưới nước, trên cạn, bóng râm, ngoài sáng)."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Theo dõi, ghi chép kích thước lá, màu sắc lá, đường kính thân và rút ra kết luận về tính mềm dẻo kiểu hình."
                }
            ]
        },
        "Chương 4: Di truyền quần thể": {
            "Bài 18: Di truyền học quần thể": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Phát biểu được khái niệm quần thể (từ góc độ di truyền học) và lấy được ví dụ minh hoạ; phát biểu được khái niệm vốn gene của quần thể."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được các đặc trưng di truyền của quần thể: tần số của các allele và tần số của các kiểu gene."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Nêu được cấu trúc di truyền của quần thể ngẫu phối; mô tả được trạng thái cân bằng di truyền của quần thể; trình bày được định luật Hardy – Weinberg và các điều kiện nghiệm đúng."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được sự biến đổi cấu trúc di truyền của quần thể tự thụ phấn và quần thể giao phối gần qua các thế hệ (tăng đồng hợp, giảm dị hợp)."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Trình bày được ảnh hưởng của tự thụ phấn, giao phối gần, ngẫu phối chi phối tần số các allele và thành phần kiểu gene của một quần thể."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Giải thích một số vấn đề thực tiễn: vấn đề hôn nhân cận huyết trong gia đình; vấn đề thoái hoá giống do tự thụ phấn ở cây giao phấn và giao phối gần ở động vật."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Tính toán được tần số allele, tần số kiểu gene và kiểm tra trạng thái cân bằng Hardy - Weinberg của một quần thể ngẫu phối."
                }
            ]
        },
        "Chương 5: Bằng chứng và các học thuyết tiến hóa": {
            "Bài 19: Các bằng chứng tiến hóa": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm bằng chứng trực tiếp (hoá thạch) và bằng chứng gián tiếp (giải phẫu so sánh, tế bào học, sinh học phân tử)."
                },
                {
                    "code": "NT3",
                    "level": "Thông hiểu",
                    "text": "Phân biệt được cơ quan tương đồng, cơ quan tương tự và cơ quan thoái hoá; phân tích được ý nghĩa của các cơ quan đó trong việc xác định mối quan hệ họ hàng giữa các loài."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được các bằng chứng tế bào học và sinh học phân tử (tính thống nhất về mã di truyền, cấu tạo DNA, protein) chứng minh sinh giới có chung nguồn gốc."
                }
            ],
            "Bài 20: Quan niệm của Darwin về chọn lọc tự nhiên và hình thành loài": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được những nội dung cốt lõi trong học thuyết tiến hóa của Darwin: biến dị cá thể, đấu tranh sinh tồn, chọn lọc tự nhiên và chọn lọc nhân tạo."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Trình bày được cơ chế tác động của chọn lọc tự nhiên theo Darwin (chọn lọc những biến dị có lợi, đào thải biến dị có hại, tích luỹ các đặc điểm thích nghi qua nhiều thế hệ)."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được cơ chế hình thành loài mới theo quan niệm phân li tính trạng của Darwin từ một loài tổ tiên chung."
                }
            ],
            "Bài 21: Học thuyết tiến hóa tổng hợp hiện đại": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Phát biểu được khái niệm tiến hóa nhỏ và tiến hóa lớn; nêu đơn vị tiến hóa cơ sở là quần thể."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được vai trò của 5 nhân tố tiến hóa đối với sự biến đổi tần số allele và thành phần kiểu gene của quần thể: đột biến, di - nhập gene, chọn lọc tự nhiên, các yếu tố ngẫu nhiên (phiêu bạt di truyền), giao phối không ngẫu nhiên."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được cơ chế hình thành quần thể thích nghi và vai trò sáng tạo của chọn lọc tự nhiên."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được các cơ chế cách li sinh sản (cách li trước hợp tử và cách li sau hợp tử) và các con đường hình thành loài mới (khác khu vực địa lí, cùng khu vực địa lí: cách li sinh thái, cách li tập tính, lai xa và đa bội hoá)."
                }
            ],
            "Bài 22: Tiến hóa lớn và quá trình phát sinh chủng loại": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Phát biểu được khái niệm tiến hóa lớn; phân biệt được tiến hóa lớn và tiến hóa nhỏ; nêu được một số minh chứng về tiến hóa lớn."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Vẽ được sơ đồ ba giai đoạn phát sinh sự sống trên Trái Đất (tiến hóa hóa học, tiến hóa tiền sinh học, tiến hóa sinh học)."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Dựa vào sơ đồ, trình bày được các đại địa chất (Thái cổ, Nguyên sinh, Cổ sinh, Trung sinh, Tân sinh) và các biến cố lớn thể hiện sự phát triển của sinh vật trong các đại đó."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Dựa vào sơ đồ cây sự sống, trình bày được sinh giới có nguồn gốc chung và phân tích được sự phát sinh chủng loại là kết quả của tiến hóa."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Vẽ được sơ đồ các giai đoạn chính trong quá trình phát sinh loài người; nêu được loài người hiện nay (H. sapiens) đã tiến hóa từ loài vượn người (Australopithecus) qua các giai đoạn trung gian."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Làm được bài tập sưu tầm tài liệu về sự phát sinh và phát triển của sinh giới hoặc của loài người."
                }
            ]
        },
        "Chương 6: Môi trường và sinh thái học quần thể": {
            "Bài 23: Môi trường và các nhân tố sinh thái": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Phát biểu được khái niệm môi trường sống của sinh vật; nêu được khái niệm nhân tố sinh thái."
                },
                {
                    "code": "NT3",
                    "level": "Thông hiểu",
                    "text": "Phân biệt được các nhân tố sinh thái vô sinh và hữu sinh; lấy được ví dụ về tác động của các nhân tố sinh thái lên đời sống sinh vật và sự thích nghi của sinh vật với các nhân tố đó."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Trình bày được các quy luật về tác động của các nhân tố sinh thái lên đời sống sinh vật: quy luật giới hạn sinh thái, quy luật tác động tổng hợp, quy luật tác động không đồng đều."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phát biểu được khái niệm nhịp sinh học; giải thích được nhịp sinh học chính là sự thích nghi của sinh vật với những thay đổi có tính chu kì của môi trường; tìm hiểu được nhịp sinh học của cơ thể mình."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Phân tích được những hoạt động của sinh vật có thể tác động làm thay đổi môi trường sống của chúng."
                }
            ],
            "Bài 24: Sinh thái học quần thể": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Phát biểu được khái niệm quần thể sinh vật (dưới góc độ sinh thái học); nêu được các mối quan hệ hỗ trợ và cạnh tranh giữa các cá thể trong quần thể."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được các đặc trưng cơ bản của quần thể: kích thước quần thể, mật độ cá thể, tỉ lệ giới tính, nhóm tuổi, kiểu phân bố cá thể (ngẫu nhiên, đồng đều, theo nhóm)."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được các nhân tố ảnh hưởng đến kích thước quần thể (mức sinh sản, tử vong, xuất cư, nhập cư); phân biệt tăng trưởng theo tiềm năng sinh học (đường cong chữ J) và tăng trưởng thực tế (đường cong chữ S)."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân biệt được biến động số lượng cá thể theo chu kì và không theo chu kì; giải thích cơ chế điều hoà mật độ quần thể duy trì trạng thái cân bằng."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Vận dụng các hiểu biết về đặc trưng quần thể trong trồng trọt, chăn nuôi (xác định mật độ nuôi trồng hợp lí, khai thác tài nguyên sinh vật đúng mức, kiểm soát dịch hại)."
                }
            ],
            "Bài 25: Thực hành: Xác định một số đặc trưng của quần thể": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Mô tả được phương pháp ô tiêu chuẩn và phương pháp đánh bắt - thả lại để ước tính kích thước hoặc mật độ cá thể của quần thể."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thực hiện được phương pháp đếm cá thể trong ô tiêu chuẩn (quần thể thực vật/cỏ) hoặc bài tập mô phỏng đánh bắt - thả lại (quần thể động vật)."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Xử lí số liệu, tính mật độ trung bình và viết báo cáo thu hoạch."
                }
            ]
        },
        "Chương 7: Sinh thái học quần xã": {
            "Bài 26: Quần xã sinh vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Phát biểu được khái niệm quần xã sinh vật và lấy được ví dụ minh hoạ."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được các đặc trưng cơ bản về thành phần loài trong quần xã: độ phong phú, loài ưu thế, loài đặc trưng, loài chủ chốt."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Mô tả được cấu trúc không gian của quần xã (phân tầng thẳng đứng và phân bố theo mặt phẳng nằm ngang) và ý nghĩa thích nghi giảm cạnh tranh, tận dụng nguồn sống."
                },
                {
                    "code": "NT5",
                    "level": "Thông hiểu",
                    "text": "Phân biệt được các mối quan hệ sinh thái giữa các loài trong quần xã: quan hệ hỗ trợ (cộng sinh, hợp tác, hội sinh) và quan hệ đối kháng (cạnh tranh, kí sinh, ức chế - cảm nhiễm, sinh vật ăn sinh vật)."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được hiện tượng khống chế sinh học và cân bằng sinh học trong quần xã; nêu ứng dụng của khống chế sinh học trong nông nghiệp."
                }
            ],
            "Bài 27: Thực hành: Tìm hiểu cấu trúc dinh dưỡng của quần xã trong tự nhiên": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Nắm vững cách quan sát và lập sơ đồ các chuỗi thức ăn, lưới thức ăn trong một hệ sinh thái cụ thể (vườn trường, ao hồ, rừng)."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Điều tra, ghi chép thành phần sinh vật sản xuất, sinh vật tiêu thụ các bậc và sinh vật phân giải trong khu vực khảo sát."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Xây dựng được lưới thức ăn hoàn chỉnh và viết báo cáo đánh giá tính ổn định của quần xã."
                }
            ],
            "Bài 28: Hệ sinh thái": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Phát biểu được khái niệm hệ sinh thái; kể tên các thành phần cấu trúc của hệ sinh thái (thành phần vô sinh và thành phần hữu sinh: sinh vật sản xuất, tiêu thụ, phân giải)."
                },
                {
                    "code": "NT3",
                    "level": "Thông hiểu",
                    "text": "Phân biệt được các kiểu hệ sinh thái chủ yếu trên Trái Đất: hệ sinh thái tự nhiên (trên cạn, dưới nước) và hệ sinh thái nhân tạo."
                },
                {
                    "code": "NT5",
                    "level": "Thông hiểu",
                    "text": "So sánh được đặc điểm cấu trúc và độ ổn định giữa hệ sinh thái tự nhiên và hệ sinh thái nhân tạo."
                }
            ],
            "Bài 29: Trao đổi vật chất và chuyển hóa năng lượng trong hệ sinh thái": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm chuỗi thức ăn, lưới thức ăn và bậc dinh dưỡng trong hệ sinh thái."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được các loại hình tháp sinh thái: hình tháp số lượng, hình tháp sinh khối và hình tháp năng lượng; nêu ý nghĩa của từng loại hình tháp."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được dòng năng lượng trong hệ sinh thái (nguồn năng lượng mặt trời đi vào chuỗi thức ăn, thất thoát qua hô hấp, nhiệt, chất thải và hiệu suất sinh thái giữa các bậc dinh dưỡng khoảng 10%)."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Tính được hiệu suất sinh thái và sản lượng sinh vật sơ cấp/thứ cấp trong một chuỗi thức ăn đơn giản."
                }
            ],
            "Bài 30: Diễn thế sinh thái": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Phát biểu được khái niệm diễn thế sinh thái; phân biệt được diễn thế nguyên sinh và diễn thế thứ sinh về khởi đầu, diễn biến và kết quả."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được các nguyên nhân gây ra diễn thế sinh thái: nguyên nhân bên ngoài (thay đổi khí hậu, thiên tai) và nguyên nhân bên trong (tác động của chính quần xã làm biến đổi môi trường)."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được tầm quan trọng của việc nghiên cứu diễn thế sinh thái trong khai thác tài nguyên và bảo vệ môi trường."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Đề xuất các biện pháp phục hồi và bảo vệ hệ sinh thái rừng ngập mặn hoặc đất sau canh tác nương rẫy."
                }
            ],
            "Bài 31: Sinh quyển, khu sinh học và chu trình sinh - địa - hoá": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Phát biểu được khái niệm sinh quyển; giải thích được sinh quyển là một cấp độ tổ chức sống lớn nhất hành tinh; trình bày được một số biện pháp bảo vệ sinh quyển."
                },
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Phát biểu được khái niệm khu sinh học; phát biểu được khái niệm chu trình sinh – địa – hoá các chất."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được đặc điểm của các khu sinh học trên cạn chủ yếu (rừng nhiệt đới, xa van, sa mạc, thảo nguyên, rừng ôn đới, rừng taiga, đài nguyên) và các khu sinh học nước ngọt, khu sinh học nước mặn trên Trái Đất; trình bày được các biện pháp bảo vệ tài nguyên sinh học của các khu sinh học đó."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Vẽ được sơ đồ khái quát chu trình trao đổi chất trong tự nhiên; trình bày được chu trình sinh – địa – hoá của một số chất: nước, carbon, nitrogen và ý nghĩa sinh học của các chu trình đó."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Vận dụng kiến thức về các chu trình trao đổi chất vào giải thích các vấn đề của thực tiễn (hiệu ứng nhà kính, phú dưỡng nguồn nước, biến đổi khí hậu toàn cầu)."
                }
            ],
            "Bài 32: Thực hành: Thiết kế một hệ sinh thái nhân tạo": [
                {
                    "code": "TH1",
                    "level": "Thông hiểu",
                    "text": "Thiết kế được bản vẽ hoặc sơ đồ mô hình một hệ sinh thái nhân tạo (bể cá cảnh thuỷ sinh kín hoặc mô hình vườn - ao - chuồng VAC)."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Lựa chọn các sinh vật sản xuất, tiêu thụ và chất nền phù hợp để lắp ráp bể nuôi cá cảnh cân bằng sinh học."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Theo dõi sự duy trì chất lượng nước, sức khoẻ của cá và viết báo cáo đánh giá sự cân bằng của hệ sinh thái nhân tạo."
                }
            ]
        },
        "Chương 8: Sinh thái học phục hồi, bảo tồn và phát triển bền vững": {
            "Bài 33: Sinh thái học phục hồi và bảo tồn đa dạng sinh vật": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm sinh thái học phục hồi; đa dạng sinh vật (đa dạng di truyền, loài, hệ sinh thái) và các mối đe dọa làm suy thoái đa dạng sinh vật."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được các nguyên lí phục hồi hệ sinh thái bị suy thoái và các biện pháp bảo tồn nguyên vị (in-situ) và chuyển vị (ex-situ)."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Đề xuất các hành động cụ thể để bảo vệ các loài nguy cấp, quý hiếm và gìn giữ đa dạng sinh học tại địa phương."
                }
            ],
            "Bài 34: Phát triển bền vững": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm phát triển bền vững và 3 trụ cột của phát triển bền vững: kinh tế, xã hội và môi trường."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được vai trò của giáo dục môi trường và kinh tế tuần hoàn trong việc đảm bảo phát triển bền vững; giải thích vì sao khai thác tài nguyên tái tạo phải ở mức dưới ngưỡng phục hồi."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Xây dựng kế hoạch hành động giảm phát thải carbon, tiết kiệm năng lượng và tiêu dùng xanh cho cá nhân và gia đình."
                }
            ],
            "Bài 35: Dự án: Tìm hiểu thực trạng bảo tồn hệ sinh thái tại địa phương và đề xuất giải pháp bảo tồn": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Lập được kế hoạch điều tra thực trạng một hệ sinh thái tự nhiên tại địa phương (vườn quốc gia, khu bảo tồn, rừng phòng hộ hoặc lưu vực sông)."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thu thập số liệu về diện tích, mức độ ô nhiễm, áp lực khai thác của con người đối với hệ sinh thái."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Viết báo cáo đánh giá tác động môi trường và đề xuất giải pháp bảo tồn khả thi tới cơ quan quản lý địa phương."
                }
            ]
        }
    },
    "Chuyên đề Sinh học 12": {
        "Chuyên đề 1: Sinh học phân tử": {
            "Bài 1: Khái quát sinh học phân tử và các thành tựu": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được đối tượng, phạm vi nghiên cứu và khái niệm sinh học phân tử."
                },
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được một số thành tựu hiện đại về lí thuyết và ứng dụng của sinh học phân tử trong y học, nông nghiệp, công nghệ sinh học và giám định pháp y."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Phân tích được các nguyên tắc ứng dụng sinh học phân tử trong thực tiễn: nguyên tắc về an toàn sinh học và nguyên tắc về đạo đức sinh học."
                }
            ],
            "Bài 2: Phương pháp tách chiết DNA": [
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Mô tả được nguyên tắc của phương pháp tách chiết và tinh sạch DNA mẫu vật."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được vai trò của phương pháp điện di trên gel agarose và kỹ thuật nhân bản gene PCR (Polymerase Chain Reaction) trong phân tích phân tử DNA."
                }
            ],
            "Bài 3: Công nghệ gene": [
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được các bước cơ bản trong công nghệ tạo DNA tái tổ hợp và chuyển gene vào tế bào vi khuẩn hoặc thực vật."
                },
                {
                    "code": "NT6",
                    "level": "Thông hiểu",
                    "text": "Giải thích được cơ chế hoạt động của hệ thống chỉnh sửa hệ gene CRISPR/Cas9 và kỹ thuật giải trình tự gene thế hệ mới (NGS)."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Đánh giá được triển vọng và những thách thức về an toàn sinh học, đạo đức của việc ứng dụng công nghệ chỉnh sửa gene trên phôi người."
                }
            ],
            "Bài 4: Dự án: Tìm hiểu về các sản phẩm chuyển gene và thu thập các thông tin đánh giá về triển vọng của công nghệ gene trong tương lai": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Xây dựng đề cương dự án tìm hiểu về các giống cây trồng hoặc động vật chuyển gene đang được lưu hành thương mại."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thu thập thông tin, nhãn mác thực phẩm biến đổi gene trên thị trường và quan điểm của người tiêu dùng."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Trình bày báo cáo dự án và thảo luận về hành lang pháp lí quản lý sinh vật biến đổi gene tại Việt Nam."
                }
            ]
        },
        "Chuyên đề 2: Kiểm soát sinh học": {
            "Bài 5: Khái niệm, cơ sở khoa học và vai trò của kiểm soát sinh học": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Nêu được khái niệm kiểm soát sinh học và vai trò của kiểm soát sinh học trong nông nghiệp bền vững."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được cơ sở sinh thái học (quan hệ vật ăn thịt - con mồi, kí sinh - vật chủ, cạnh tranh sinh học) và cơ sở di truyền học của kiểm soát sinh học (kỹ thuật tiệt sinh côn trùng SIT)."
                }
            ],
            "Bài 6: Biện pháp kiểm soát sinh học": [
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được các biện pháp kiểm soát sinh học chủ yếu: bảo vệ và tăng cường thiên địch bản địa, nhập nội thiên địch ngoại lai, sử dụng chế phẩm vi sinh vật trừ sâu hại, bẫy pheromone."
                },
                {
                    "code": "VD1",
                    "level": "Vận dụng",
                    "text": "Phân tích được ưu điểm và nhược điểm của biện pháp kiểm soát sinh học so với biện pháp hóa học; đề xuất quy trình kiểm soát sâu hại cho một loại cây trồng cụ thể."
                }
            ],
            "Bài 7: Dự án: Sưu tầm/điều tra ứng dụng kiểm soát sinh học": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Lập kế hoạch sưu tầm mẫu vật, hình ảnh hoặc điều tra thực tế việc ứng dụng kiểm soát sinh học tại địa phương."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Thu thập tư liệu về các mô hình nuôi thả ong mắt đỏ kí sinh sâu đục thân, thả cá diệt bọ gậy hoặc sử dụng chế phẩm nấm đối kháng Trichoderma."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Viết bài thu hoạch và trình bày triển lãm ảnh/poster về thiên địch và kiểm soát sinh học."
                }
            ]
        },
        "Chuyên đề 3: Sinh thái nhân văn": {
            "Bài 8: Khái niệm và giá trị của sinh thái nhân văn trong phát triển bền vững": [
                {
                    "code": "NT1",
                    "level": "Nhận biết",
                    "text": "Phát biểu được khái niệm sinh thái nhân văn và đối tượng nghiên cứu (mối quan hệ tương hỗ giữa con người và hệ thống xã hội với môi trường tự nhiên)."
                },
                {
                    "code": "NT4",
                    "level": "Thông hiểu",
                    "text": "Phân tích được giá trị của sinh thái nhân văn trong việc duy trì cân bằng sinh thái và định hướng phát triển bền vững."
                }
            ],
            "Bài 9: Giá trị của sinh thái nhân văn trong một số lĩnh vực": [
                {
                    "code": "NT2",
                    "level": "Nhận biết",
                    "text": "Trình bày được các ứng dụng của sinh thái nhân văn trong quy hoạch đô thị sinh thái, nông nghiệp sinh thái và du lịch sinh thái."
                },
                {
                    "code": "VD2",
                    "level": "Vận dụng",
                    "text": "Vận dụng các nguyên tắc sinh thái nhân văn để đánh giá một mô hình nông nghiệp sinh thái hoặc làng nghề truyền thống tại địa phương."
                }
            ],
            "Bài 10: Dự án: Điều tra, tìm hiểu về một trong các lĩnh vực sinh thái nhân văn tại địa phương": [
                {
                    "code": "TH3",
                    "level": "Thông hiểu",
                    "text": "Lập được kế hoạch dự án điều tra hiện trạng sinh thái nhân văn tại một khu đô thị, làng sinh thái hoặc điểm du lịch cộng đồng."
                },
                {
                    "code": "TH4",
                    "level": "Vận dụng",
                    "text": "Tiến hành khảo sát thực địa, phỏng vấn người dân và chính quyền về sự hài hoà giữa cảnh quan tự nhiên và đời sống nhân sinh."
                },
                {
                    "code": "TH5",
                    "level": "Vận dụng",
                    "text": "Viết báo cáo đánh giá tác động và đề xuất giải pháp kiến trúc cảnh quan hoặc phát triển du lịch sinh thái bền vững."
                }
            ]
        }
    }
}




# --- HÀM TẠO HÌNH ẢNH VECTOR SVG TRỰC TIẾP CHO CÂU HỎI ---
import base64

def _svg_to_b64(svg_str):
    b64 = base64.b64encode(svg_str.strip().encode('utf-8')).decode('utf-8')
    return f"data:image/svg+xml;base64,{b64}"




# --- TẠO HÌNH ẢNH VECTOR SVG TRỰC TIẾP CHO CÂU HỎI & ĐÁP ÁN ---
def _svg_to_b64(svg_str):
    b64 = base64.b64encode(svg_str.strip().encode('utf-8')).decode('utf-8')
    return f"data:image/svg+xml;base64,{b64}"

SVG_DNA_REP = _svg_to_b64("""<svg width="520" height="230" xmlns="http://www.w3.org/2000/svg">
  <rect width="100%" height="100%" fill="#f8f9fa" rx="8" stroke="#dee2e6"/>
  <text x="260" y="25" font-size="14" font-weight="bold" text-anchor="middle" fill="#0d6efd">SƠ ĐỒ CHẠC TÁI BẢN DNA VÀ SỰ TỔNG HỢP MẠCH MỚI</text>
  <path d="M 50,70 L 200,70 L 440,30" stroke="#dc3545" stroke-width="4" fill="none"/>
  <path d="M 50,150 L 200,150 L 440,190" stroke="#0d6efd" stroke-width="4" fill="none"/>
  <path d="M 200,70 L 420,35" stroke="#198754" stroke-width="3" stroke-dasharray="5,3" fill="none"/>
  <circle cx="200" cy="110" r="18" fill="#6f42c1"/>
  <text x="200" y="114" font-size="10" fill="#ffffff" font-weight="bold" text-anchor="middle">Helicase</text>
  <text x="445" y="30" font-size="12" fill="#dc3545" font-weight="bold">3'</text>
  <text x="445" y="190" font-size="12" fill="#0d6efd" font-weight="bold">5'</text>
</svg>""")

SVG_MUTATION_5BU = _svg_to_b64("""<svg width="500" height="180" xmlns="http://www.w3.org/2000/svg">
  <rect width="100%" height="100%" fill="#f8f9fa" rx="8" stroke="#dee2e6"/>
  <text x="250" y="22" font-size="13" font-weight="bold" text-anchor="middle" fill="#dc3545">SƠ ĐỒ CƠ CHẾ GÂY ĐỘT BIẾN CỦA 5-BROMOURACIL (5-BU)</text>
  <rect x="30" y="50" width="80" height="30" fill="#e2e3e5" rx="4"/>
  <text x="70" y="70" font-size="12" font-weight="bold" text-anchor="middle">A - T</text>
  <text x="130" y="70" font-size="16" fill="#6c757d">→</text>
  <rect x="150" y="50" width="90" height="30" fill="#fff3cd" rx="4" stroke="#ffc107"/>
  <text x="195" y="70" font-size="12" font-weight="bold" fill="#856404" text-anchor="middle">A - 5-BU</text>
  <text x="255" y="70" font-size="16" fill="#6c757d">→</text>
  <rect x="275" y="50" width="90" height="30" fill="#ffe5d0" rx="4" stroke="#fd7e14"/>
  <text x="320" y="70" font-size="12" font-weight="bold" fill="#d9480f" text-anchor="middle">G - 5-BU</text>
  <text x="380" y="70" font-size="16" fill="#6c757d">→</text>
  <rect x="400" y="50" width="80" height="30" fill="#f8d7da" rx="4" stroke="#dc3545"/>
  <text x="440" y="70" font-size="12" font-weight="bold" fill="#721c24" text-anchor="middle">G - C</text>
  <text x="250" y="125" font-size="11" fill="#495057" text-anchor="middle">Cặp ban đầu (A-T) → Nhân đôi 1 → Nhân đôi 2 → Nhân đôi 3 (Cặp đột biến G-C)</text>
</svg>""")

SVG_PHOTOSYNTHESIS_GRAPH = _svg_to_b64("""<svg width="480" height="220" xmlns="http://www.w3.org/2000/svg">
  <rect width="100%" height="100%" fill="#f8f9fa" rx="8" stroke="#dee2e6"/>
  <text x="240" y="22" font-size="13" font-weight="bold" text-anchor="middle" fill="#198754">ĐỒ THỊ ẢNH HƯỞNG CỦA CƯỜNG ĐỘ ÁNH SÁNG ĐẾN QUANG HỢP</text>
  <line x1="50" y1="180" x2="440" y2="180" stroke="#212529" stroke-width="2"/>
  <line x1="50" y1="180" x2="50" y2="40" stroke="#212529" stroke-width="2"/>
  <path d="M 50,180 Q 150,80 280,70 L 430,70" stroke="#198754" stroke-width="3" fill="none"/>
  <circle cx="110" cy="140" r="4" fill="#dc3545"/>
  <text x="110" y="160" font-size="10" fill="#dc3545" font-weight="bold" text-anchor="middle">Điểm bù ánh sáng (Ic)</text>
  <circle cx="280" cy="70" r="4" fill="#0d6efd"/>
  <text x="280" y="55" font-size="10" fill="#0d6efd" font-weight="bold" text-anchor="middle">Điểm bão hòa ánh sáng (Is)</text>
  <text x="240" y="202" font-size="11" fill="#212529">Cường độ ánh sáng (Lux)</text>
</svg>""")

SVG_CELL_CULTURE = _svg_to_b64("""<svg width="500" height="170" xmlns="http://www.w3.org/2000/svg">
  <rect width="100%" height="100%" fill="#f8f9fa" rx="8" stroke="#dee2e6"/>
  <text x="250" y="22" font-size="13" font-weight="bold" text-anchor="middle" fill="#0d6efd">SƠ ĐỒ QUY TRÌNH NUÔI CẤY MÔ TẾ BÀO THỰC VẬT</text>
  <rect x="20" y="50" width="90" height="50" rx="5" fill="#e2e3e5" stroke="#6c757d"/>
  <text x="65" y="72" font-size="10" font-weight="bold" text-anchor="middle">1. Mẫu mô nguyên</text>
  <text x="65" y="88" font-size="9" text-anchor="middle">(Cây mẹ)</text>
  <text x="118" y="80" font-size="14" fill="#6c757d">→</text>
  <rect x="132" y="50" width="100" height="50" rx="5" fill="#fff3cd" stroke="#ffc107"/>
  <text x="182" y="72" font-size="10" font-weight="bold" text-anchor="middle">2. Tạo mô sẹo (Callus)</text>
  <text x="182" y="88" font-size="9" text-anchor="middle">(Phản phân hóa)</text>
  <text x="240" y="80" font-size="14" fill="#6c757d">→</text>
  <rect x="254" y="50" width="100" height="50" rx="5" fill="#cff4fc" stroke="#0dcaf0"/>
  <text x="304" y="72" font-size="10" font-weight="bold" text-anchor="middle">3. Tái sinh cơ quan</text>
  <text x="304" y="88" font-size="9" text-anchor="middle">(Tạo chồi/rễ)</text>
  <text x="362" y="80" font-size="14" fill="#6c757d">→</text>
  <rect x="376" y="50" width="105" height="50" rx="5" fill="#d1e7dd" stroke="#198754"/>
  <text x="428" y="72" font-size="10" font-weight="bold" text-anchor="middle">4. Cây hoàn chỉnh</text>
  <text x="428" y="88" font-size="9" text-anchor="middle">(Vườn ươm)</text>
</svg>""")

SVG_OPT_NUCLEOTIDE_A = _svg_to_b64("""<svg width="180" height="100" xmlns="http://www.w3.org/2000/svg">
  <rect width="100%" height="100%" fill="#ffffff" rx="4" stroke="#0d6efd"/>
  <text x="90" y="18" font-size="11" font-weight="bold" fill="#0d6efd" text-anchor="middle">Phương án A: Deoxyribonucleotide</text>
  <circle cx="40" cy="55" r="14" fill="#ffc107"/>
  <text x="40" y="59" font-size="10" font-weight="bold" text-anchor="middle">P</text>
  <polygon points="80,40 100,55 90,80 70,80 60,55" fill="#0dcaf0"/>
  <text x="80" y="65" font-size="9" font-weight="bold" text-anchor="middle">Đường dR</text>
  <rect x="115" y="45" width="45" height="25" rx="3" fill="#dc3545"/>
  <text x="137" y="61" font-size="10" font-weight="bold" fill="#ffffff" text-anchor="middle">Base A</text>
</svg>""")

SVG_OPT_NUCLEOTIDE_B = _svg_to_b64("""<svg width="180" height="100" xmlns="http://www.w3.org/2000/svg">
  <rect width="100%" height="100%" fill="#ffffff" rx="4" stroke="#6c757d"/>
  <text x="90" y="18" font-size="11" font-weight="bold" fill="#6c757d" text-anchor="middle">Phương án B: Ribonucleotide (Uracil)</text>
  <circle cx="40" cy="55" r="14" fill="#ffc107"/>
  <text x="40" y="59" font-size="10" font-weight="bold" text-anchor="middle">P</text>
  <polygon points="80,40 100,55 90,80 70,80 60,55" fill="#198754"/>
  <text x="80" y="65" font-size="9" font-weight="bold" fill="#ffffff" text-anchor="middle">Đường Ribose</text>
  <rect x="115" y="45" width="45" height="25" rx="3" fill="#6f42c1"/>
  <text x="137" y="61" font-size="10" font-weight="bold" fill="#ffffff" text-anchor="middle">Base U</text>
</svg>""")

SVG_OPT_NUCLEOTIDE_C = _svg_to_b64("""<svg width="180" height="100" xmlns="http://www.w3.org/2000/svg">
  <rect width="100%" height="100%" fill="#ffffff" rx="4" stroke="#6c757d"/>
  <text x="90" y="18" font-size="11" font-weight="bold" fill="#6c757d" text-anchor="middle">Phương án C: Amino acid</text>
  <rect x="30" y="40" width="120" height="35" rx="4" fill="#e2e3e5"/>
  <text x="90" y="62" font-size="10" font-weight="bold" text-anchor="middle">H2N - CH(R) - COOH</text>
</svg>""")

SVG_OPT_NUCLEOTIDE_D = _svg_to_b64("""<svg width="180" height="100" xmlns="http://www.w3.org/2000/svg">
  <rect width="100%" height="100%" fill="#ffffff" rx="4" stroke="#6c757d"/>
  <text x="90" y="18" font-size="11" font-weight="bold" fill="#6c757d" text-anchor="middle">Phương án D: Glucose</text>
  <polygon points="90,35 125,55 125,80 90,95 55,80 55,55" fill="#ffe5d0" stroke="#fd7e14"/>
  <text x="90" y="68" font-size="10" font-weight="bold" fill="#d9480f" text-anchor="middle">C6H12O6</text>
</svg>""")



# --- CƠ SỞ DỮ LIỆU NGÂN HÀNG CÂU HỎI MẪU NÂNG CẤP (TÍCH HỢP ẢNH CÂU HỎI & ẢNH ĐÁP ÁN A, B, C, D) ---
DEFAULT_QUESTIONS_DATABASE = [
    {
        "id": 1,
        "book": "Sinh học 10",
        "topic": "Phần Mở đầu",
        "lesson": "Bài 1: Giới thiệu khái quát môn Sinh học",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Đối tượng nghiên cứu của Sinh học là",
        "opt_a": "các sinh vật sống và các cấp độ tổ chức của thế giới sống.",
        "opt_a_img": "",
        "opt_b": "các vật thể vô sinh trong tự nhiên.",
        "opt_b_img": "",
        "opt_c": "các hiện tượng thời tiết, khí hậu.",
        "opt_c_img": "",
        "opt_d": "các chất hóa học và phản ứng hạt nhân.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Sinh học là ngành khoa học nghiên cứu về sự sống, đối tượng là các sinh vật sống và các cấp độ tổ chức sống."
    },
    {
        "id": 2,
        "book": "Sinh học 10",
        "topic": "Phần Mở đầu",
        "lesson": "Bài 1: Giới thiệu khái quát môn Sinh học",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Lĩnh vực nào sau đây KHÔNG thuộc ngành Sinh học?",
        "opt_a": "Thiên văn học.",
        "opt_a_img": "",
        "opt_b": "Giải phẫu học.",
        "opt_b_img": "",
        "opt_c": "Thực vật học.",
        "opt_c_img": "",
        "opt_d": "Vi sinh vật học.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Thiên văn học thuộc lĩnh vực Vật lý/Vũ trụ học, không thuộc ngành Sinh học."
    },
    {
        "id": 3,
        "book": "Sinh học 10",
        "topic": "Phần Mở đầu",
        "lesson": "Bài 1: Giới thiệu khái quát môn Sinh học",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Ngành nghề nào sau đây ứng dụng trực tiếp kiến thức Sinh học trong chăm sóc sức khỏe con người?",
        "opt_a": "Y học.",
        "opt_a_img": "",
        "opt_b": "Nông nghiệp.",
        "opt_b_img": "",
        "opt_c": "Lâm nghiệp.",
        "opt_c_img": "",
        "opt_d": "Thủy sản.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Y học là ngành nghề ứng dụng trực tiếp thành tựu sinh học trong chẩn đoán và điều trị bệnh bảo vệ sức khỏe."
    },
    {
        "id": 4,
        "book": "Sinh học 10",
        "topic": "Phần Mở đầu",
        "lesson": "Bài 1: Giới thiệu khái quát môn Sinh học",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Sự phát triển đáp ứng nhu cầu của thế hệ hiện tại mà không làm tổn hại đến khả năng đáp ứng nhu cầu đó của các thế hệ tương lai gọi là",
        "opt_a": "phát triển bền vững.",
        "opt_a_img": "",
        "opt_b": "phát triển kinh tế nhanh.",
        "opt_b_img": "",
        "opt_c": "phát triển công nghiệp hóa.",
        "opt_c_img": "",
        "opt_d": "phát triển tự do.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Phát triển bền vững là sự phát triển kết hợp hài hòa giữa tăng trưởng kinh tế, bảo vệ môi trường và công bằng xã hội."
    },
    {
        "id": 5,
        "book": "Sinh học 10",
        "topic": "Phần Mở đầu",
        "lesson": "Bài 1: Giới thiệu khái quát môn Sinh học",
        "qtype": "TNKQ đúng - sai",
        "level": "Thông hiểu",
        "indicator": "NT6",
        "content": "Khi nói về các lĩnh vực nghiên cứu và ứng dụng của Sinh học, mỗi phát biểu sau đây là Đúng hay Sai?",
        "opt_a": "",
        "opt_a_img": "",
        "opt_b": "",
        "opt_b_img": "",
        "opt_c": "",
        "opt_c_img": "",
        "opt_d": "",
        "opt_d_img": "",
        "tf_a": "Y học và Dược học là hai lĩnh vực ứng dụng sinh học trọng tâm trong bảo vệ sức khỏe con người.",
        "tf_ans_a": "Đúng",
        "tf_b": "Đạo đức sinh học chỉ áp dụng đối với các nghiên cứu trên thực vật.",
        "tf_ans_b": "Sai",
        "tf_c": "Sinh học cung cấp cơ sở khoa học cho việc khai thác hợp lý và bảo vệ tài nguyên thiên nhiên.",
        "tf_ans_c": "Đúng",
        "tf_d": "Công nghệ sinh học đóng vai trò quan trọng trong việc tạo ra các giống cây trồng biến đổi gen kháng sâu bệnh.",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "a) Đúng, b) Sai, c) Đúng, d) Đúng",
        "a_image": "",
        "explanation": "Ý b sai vì Đạo đức sinh học áp dụng cho nghiên cứu trên con người, động vật và toàn bộ thế giới sống."
    },
    {
        "id": 6,
        "book": "Sinh học 10",
        "topic": "Phần Mở đầu",
        "lesson": "Bài 1: Giới thiệu khái quát môn Sinh học",
        "qtype": "TNKQ trả lời ngắn",
        "level": "Thông hiểu",
        "indicator": "NT6",
        "content": "Thuật ngữ nào chỉ tập hợp các nguyên tắc, chuẩn mực đạo đức áp dụng trong các nghiên cứu và ứng dụng sinh học đối với con người và sinh vật?",
        "opt_a": "",
        "opt_a_img": "",
        "opt_b": "",
        "opt_b_img": "",
        "opt_c": "",
        "opt_c_img": "",
        "opt_d": "",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "Đạo đức sinh học",
        "a_image": "",
        "explanation": "Đạo đức sinh học là quy chuẩn nguyên tắc đạo đức trong nghiên cứu ứng dụng sinh học."
    },
    {
        "id": 7,
        "book": "Sinh học 10",
        "topic": "Phần Mở đầu",
        "lesson": "Bài 2: Phương pháp nghiên cứu và học tập môn Sinh học",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Phương pháp nghiên cứu sinh học nào sử dụng các giác quan hoặc thiết bị hỗ trợ để thu thập thông tin về đối tượng mà không làm thay đổi trạng thái tự nhiên của chúng?",
        "opt_a": "Phương pháp quan sát.",
        "opt_a_img": "",
        "opt_b": "Phương pháp làm việc trong phòng thí nghiệm.",
        "opt_b_img": "",
        "opt_c": "Phương pháp thực địa.",
        "opt_c_img": "",
        "opt_d": "Phương pháp giải phẫu.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Phương pháp quan sát sử dụng giác quan/kính hiển vi thu thập số liệu tự nhiên mà không can thiệp tác động."
    },
    {
        "id": 8,
        "book": "Sinh học 10",
        "topic": "Phần Mở đầu",
        "lesson": "Bài 2: Phương pháp nghiên cứu và học tập môn Sinh học",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Phương pháp nào sau đây giúp nhà nghiên cứu chủ động tác động vào đối tượng và điều khiển các yếu tố môi trường nhằm kiểm tra một giả thuyết?",
        "opt_a": "Phương pháp thực nghiệm khoa học.",
        "opt_a_img": "",
        "opt_b": "Phương pháp quan sát.",
        "opt_b_img": "",
        "opt_c": "Phương pháp điều tra thực địa.",
        "opt_c_img": "",
        "opt_d": "Phương pháp phân tích định tính.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Thực nghiệm khoa học là phương pháp chủ động thay đổi yếu tố môi trường để kiểm chứng giả thuyết."
    },
    {
        "id": 9,
        "book": "Sinh học 10",
        "topic": "Phần Mở đầu",
        "lesson": "Bài 2: Phương pháp nghiên cứu và học tập môn Sinh học",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Thông hiểu",
        "indicator": "TH1",
        "content": "Môn học/lĩnh vực liên ngành sử dụng công nghệ thông tin để thu thập, lưu trữ, phân tích và xử lý dữ liệu sinh học gọi là",
        "opt_a": "Tin sinh học.",
        "opt_a_img": "",
        "opt_b": "Sinh học thực nghiệm.",
        "opt_b_img": "",
        "opt_c": "Sinh học giải phẫu.",
        "opt_c_img": "",
        "opt_d": "Sinh học sinh thái.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Tin sinh học (Bioinformatics) kết hợp công nghệ thông tin để quản lý và xử lý dữ liệu sinh học lớn."
    },
    {
        "id": 10,
        "book": "Sinh học 10",
        "topic": "Phần Mở đầu",
        "lesson": "Bài 2: Phương pháp nghiên cứu và học tập môn Sinh học",
        "qtype": "TNKQ đúng - sai",
        "level": "Thông hiểu",
        "indicator": "TH1",
        "content": "Khi nói về tiến trình nghiên cứu khoa học và Tin sinh học, mỗi phát biểu sau đây là Đúng hay Sai?",
        "opt_a": "",
        "opt_a_img": "",
        "opt_b": "",
        "opt_b_img": "",
        "opt_c": "",
        "opt_c_img": "",
        "opt_d": "",
        "opt_d_img": "",
        "tf_a": "Nếu kết quả thực nghiệm không ủng hộ giả thuyết, nhà khoa học phải bác bỏ hoặc điều chỉnh giả thuyết đó.",
        "tf_ans_a": "Đúng",
        "tf_b": "Kính hiển vi điện tử có độ phóng đại nhỏ hơn kính hiển vi quang học.",
        "tf_ans_b": "Sai",
        "tf_c": "Tin sinh học giúp truy tìm trình tự gen và dự đoán cấu trúc không gian của protein.",
        "tf_ans_c": "Đúng",
        "tf_d": "Báo cáo kết quả nghiên cứu chỉ cần chứa hình ảnh, không cần số liệu cụ thể.",
        "tf_ans_d": "Sai",
        "q_image": "",
        "answer": "a) Đúng, b) Sai, c) Đúng, d) Sai",
        "a_image": "",
        "explanation": "Ý b sai vì kính hiển vi điện tử độ phóng đại lớn hơn nhiều. Ý d sai vì báo cáo bắt buộc phải chứa số liệu cụ thể."
    },
    {
        "id": 11,
        "book": "Sinh học 10",
        "topic": "Phần Mở đầu",
        "lesson": "Bài 2: Phương pháp nghiên cứu và học tập môn Sinh học",
        "qtype": "TNKQ trả lời ngắn",
        "level": "Thông hiểu",
        "indicator": "TH1",
        "content": "Lĩnh vực liên ngành kết hợp sinh học và công nghệ thông tin để lưu trữ, phân tích dữ liệu sinh học gọi là gì?",
        "opt_a": "",
        "opt_a_img": "",
        "opt_b": "",
        "opt_b_img": "",
        "opt_c": "",
        "opt_c_img": "",
        "opt_d": "",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "Tin sinh học",
        "a_image": "",
        "explanation": "Tin sinh học kết hợp CNTT xử lý dữ liệu sinh học."
    },
    {
        "id": 12,
        "book": "Sinh học 10",
        "topic": "Phần Mở đầu",
        "lesson": "Bài 3: Các cấp độ tổ chức của thế giới sống",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Cấp độ tổ chức sống cơ bản nhất của mọi sinh vật là",
        "opt_a": "Tế bào.",
        "opt_a_img": "",
        "opt_b": "Phân tử.",
        "opt_b_img": "",
        "opt_c": "Bào quan.",
        "opt_c_img": "",
        "opt_d": "Mô.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Tế bào là đơn vị cấu trúc và chức năng cơ bản nhất của mọi cơ thể sống."
    },
    {
        "id": 13,
        "book": "Sinh học 10",
        "topic": "Phần Mở đầu",
        "lesson": "Bài 3: Các cấp độ tổ chức của thế giới sống",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Đặc điểm nào sau đây KHÔNG phải là đặc điểm chung của các cấp độ tổ chức sống?",
        "opt_a": "Là hệ thống kín, không giao lưu với môi trường.",
        "opt_a_img": "",
        "opt_b": "Tổ chức theo nguyên tắc thứ bậc.",
        "opt_b_img": "",
        "opt_c": "Hệ thống tự điều chỉnh.",
        "opt_c_img": "",
        "opt_d": "Liên tục tiến hóa.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Tất cả các cấp độ tổ chức sống đều là hệ thống mở, luôn trao đổi chất và năng lượng với môi trường."
    },
    {
        "id": 14,
        "book": "Sinh học 10",
        "topic": "Phần Mở đầu",
        "lesson": "Bài 3: Các cấp độ tổ chức của thế giới sống",
        "qtype": "TNKQ đúng - sai",
        "level": "Thông hiểu",
        "indicator": "NT5",
        "content": "Khi nói về các đặc điểm chung của thế giới sống, mỗi phát biểu sau đây là Đúng hay Sai?",
        "opt_a": "",
        "opt_a_img": "",
        "opt_b": "",
        "opt_b_img": "",
        "opt_c": "",
        "opt_c_img": "",
        "opt_d": "",
        "opt_d_img": "",
        "tf_a": "Cơ thể sống trao đổi chất và năng lượng liên tục với môi trường xung quanh.",
        "tf_ans_a": "Đúng",
        "tf_b": "Khả năng tự điều chỉnh giúp cơ thể sinh vật chống lại mọi sự thay đổi của môi trường mà không bị bệnh.",
        "tf_ans_b": "Sai",
        "tf_c": "Đột biến gen là một trong những nguyên nhân tạo nên sự đa dạng của thế giới sống qua quá trình tiến hóa.",
        "tf_ans_c": "Đúng",
        "tf_d": "Đặc tính nổi trội chỉ xuất hiện ở cấp độ tổ chức phân tử.",
        "tf_ans_d": "Sai",
        "q_image": "",
        "answer": "a) Đúng, b) Sai, c) Đúng, d) Sai",
        "a_image": "",
        "explanation": "Ý b sai vì nếu môi trường biến đổi quá giới hạn tự điều chỉnh, cơ thể sẽ bị bệnh hoặc tử vong. Ý d sai vì đặc tính nổi trội xuất hiện ở các cấp tổ chức cao hơn."
    },
    {
        "id": 15,
        "book": "Sinh học 10",
        "topic": "Chương 1: Thành phần hóa học của tế bào",
        "lesson": "Bài 4: Các nguyên tố hóa học và nước",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Bốn nguyên tố chính chiếm khoảng 96% khối lượng cơ thể sống của sinh vật là",
        "opt_a": "C, H, O, N.",
        "opt_a_img": "",
        "opt_b": "C, H, O, P.",
        "opt_b_img": "",
        "opt_c": "C, H, O, Ca.",
        "opt_c_img": "",
        "opt_d": "C, N, P, S.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Bốn nguyên tố C, H, O, N chiếm khoảng 96% khối lượng cơ thể sống."
    },
    {
        "id": 16,
        "book": "Sinh học 10",
        "topic": "Chương 1: Thành phần hóa học của tế bào",
        "lesson": "Bài 4: Các nguyên tố hóa học và nước",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Các liên kết được hình thành giữa các phân tử nước với nhau là",
        "opt_a": "liên kết hydrogen.",
        "opt_a_img": "",
        "opt_b": "liên kết cộng hóa trị.",
        "opt_b_img": "",
        "opt_c": "liên kết ionic.",
        "opt_c_img": "",
        "opt_d": "liên kết peptide.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Giữa các phân tử nước hình thành liên kết hydrogen giữa nguyên tử O của phân tử này với H của phân tử khác."
    },
    {
        "id": 17,
        "book": "Sinh học 10",
        "topic": "Chương 1: Thành phần hóa học của tế bào",
        "lesson": "Bài 4: Các nguyên tố hóa học và nước",
        "qtype": "TNKQ trả lời ngắn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Tên loại liên kết yếu hình thành giữa nguyên tử Oxygen (mang điện tích âm nhẹ) của phân tử nước này với nguyên tử Hydrogen (mang điện tích dương nhẹ) của phân tử nước khác là gì?",
        "opt_a": "",
        "opt_a_img": "",
        "opt_b": "",
        "opt_b_img": "",
        "opt_c": "",
        "opt_c_img": "",
        "opt_d": "",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "Liên kết hydrogen",
        "a_image": "",
        "explanation": "Liên kết hydrogen nối giữa các phân tử nước."
    },
    {
        "id": 18,
        "book": "Sinh học 10",
        "topic": "Chương 1: Thành phần hóa học của tế bào",
        "lesson": "Bài 5: Các phân tử sinh học",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Bốn nhóm phân tử sinh học chính trong tế bào gồm",
        "opt_a": "Carbohydrate, Lipid, Protein, Nucleic acid.",
        "opt_a_img": "",
        "opt_b": "Đường, Mỡ, Vitamin, Khoáng chất.",
        "opt_b_img": "",
        "opt_c": "Protein, DNA, RNA, Amino acid.",
        "opt_c_img": "",
        "opt_d": "Carbohydrate, Lipid, Enzyme, Hormone.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "4 nhóm phân tử sinh học chính gồm Carbohydrate, Lipid, Protein và Nucleic acid."
    },
    {
        "id": 19,
        "book": "Sinh học 10",
        "topic": "Chương 1: Thành phần hóa học của tế bào",
        "lesson": "Bài 5: Các phân tử sinh học",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Thông hiểu",
        "indicator": "NT3",
        "content": "Polysaccharide nào sau đây đóng vai trò cấu tạo nên thành tế bào thực vật?",
        "opt_a": "Cellulose.",
        "opt_a_img": "",
        "opt_b": "Glycogen.",
        "opt_b_img": "",
        "opt_c": "Chitin.",
        "opt_c_img": "",
        "opt_d": "Tinh bột.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Cellulose gồm các chuỗi glucose liên kết song song tạo thành vi sợi cấu tạo nên thành tế bào thực vật."
    },
    {
        "id": 20,
        "book": "Sinh học 10",
        "topic": "Chương 1: Thành phần hóa học của tế bào",
        "lesson": "Bài 5: Các phân tử sinh học",
        "qtype": "TNKQ đúng - sai",
        "level": "Thông hiểu",
        "indicator": "NT3",
        "content": "Khi nói về Protein và Nucleic acid, mỗi phát biểu sau đây là Đúng hay Sai?",
        "opt_a": "",
        "opt_a_img": "",
        "opt_b": "",
        "opt_b_img": "",
        "opt_c": "",
        "opt_c_img": "",
        "opt_d": "",
        "opt_d_img": "",
        "tf_a": "Hiện tượng biến tính protein là hiện tượng protein bị mất cấu trúc không gian ba chiều và mất chức năng sinh học.",
        "tf_ans_a": "Đúng",
        "tf_b": "Phân tử DNA gồm 2 chuỗi xoắn kép ngược chiều nhau, liên kết theo nguyên tắc bổ sung (A-T, G-C).",
        "tf_ans_b": "Đúng",
        "tf_c": "Liên kết giữa A với T trên hai mạch của DNA là 3 liên kết hydrogen, giữa G với C là 2 liên kết hydrogen.",
        "tf_ans_c": "Sai",
        "tf_d": "mRNA có vai trò làm khuôn cho quá trình dịch mã tổng hợp protein.",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "a) Đúng, b) Đúng, c) Sai, d) Đúng",
        "a_image": "",
        "explanation": "Ý c sai vì A liên kết T bằng 2 liên kết hydrogen, G liên kết C bằng 3 liên kết hydrogen."
    },
    {
        "id": 21,
        "book": "Sinh học 10",
        "topic": "Chương 1: Thành phần hóa học của tế bào",
        "lesson": "Bài 6: Thực hành: Nhận biết một số phân tử sinh học",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Thông hiểu",
        "indicator": "TH3",
        "content": "Để nhận biết sự có mặt của đường khử (như glucose) trong mẫu thí nghiệm, người ta sử dụng thuốc thử nào sau đây?",
        "opt_a": "Dung dịch Benedict (hoặc Fehling).",
        "opt_a_img": "",
        "opt_b": "Dung dịch Iodine.",
        "opt_b_img": "",
        "opt_c": "Thuốc thử Biuret.",
        "opt_c_img": "",
        "opt_d": "Dung dịch Sudan III.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Dung dịch Benedict đun nóng với đường khử sẽ tạo kết tủa đỏ gạch Cu2O."
    },
    {
        "id": 22,
        "book": "Sinh học 10",
        "topic": "Chương 1: Thành phần hóa học của tế bào",
        "lesson": "Bài 6: Thực hành: Nhận biết một số phân tử sinh học",
        "qtype": "TNKQ trả lời ngắn",
        "level": "Thông hiểu",
        "indicator": "TH3",
        "content": "Thuốc thử Biuret được sử dụng trong phòng thí nghiệm để nhận biết nhóm phân tử sinh học nào?",
        "opt_a": "",
        "opt_a_img": "",
        "opt_b": "",
        "opt_b_img": "",
        "opt_c": "",
        "opt_c_img": "",
        "opt_d": "",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "Protein",
        "a_image": "",
        "explanation": "Thuốc thử Biuret tạo phức màu tím với liên kết peptide trong protein."
    },
    {
        "id": 23,
        "book": "Sinh học 10",
        "topic": "Chương 2: Cấu trúc tế bào",
        "lesson": "Bài 7: Tế bào nhân sơ",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Thành tế bào vi khuẩn được cấu tạo chủ yếu từ hợp chất nào?",
        "opt_a": "Peptidoglycan.",
        "opt_a_img": "",
        "opt_b": "Cellulose.",
        "opt_b_img": "",
        "opt_c": "Chitin.",
        "opt_c_img": "",
        "opt_d": "Hemicellulose.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Thành tế bào vi khuẩn cấu tạo từ Peptidoglycan."
    },
    {
        "id": 24,
        "book": "Sinh học 10",
        "topic": "Chương 2: Cấu trúc tế bào",
        "lesson": "Bài 7: Tế bào nhân sơ",
        "qtype": "TNKQ trả lời ngắn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Hợp chất hóa học chính cấu tạo nên thành tế bào vi khuẩn là gì?",
        "opt_a": "",
        "opt_a_img": "",
        "opt_b": "",
        "opt_b_img": "",
        "opt_c": "",
        "opt_c_img": "",
        "opt_d": "",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "Peptidoglycan",
        "a_image": "",
        "explanation": "Thành tế bào vi khuẩn cấu tạo chủ yếu từ peptidoglycan."
    },
    {
        "id": 25,
        "book": "Sinh học 10",
        "topic": "Chương 2: Cấu trúc tế bào",
        "lesson": "Bài 8: Tế bào nhân thực",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Bào quan nào thực hiện chức năng hô hấp tế bào, chuyển hóa năng lượng từ chất hữu cơ thành ATP?",
        "opt_a": "Ti thể.",
        "opt_a_img": "",
        "opt_b": "Lục lạp.",
        "opt_b_img": "",
        "opt_c": "Peroxisome.",
        "opt_c_img": "",
        "opt_d": "Lysosome.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Ti thể là nhà máy năng lượng hô hấp tạo ATP cho tế bào."
    },
    {
        "id": 26,
        "book": "Sinh học 10",
        "topic": "Chương 2: Cấu trúc tế bào",
        "lesson": "Bài 8: Tế bào nhân thực",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Bào quan chỉ có ở tế bào thực vật mà KHÔNG có ở tế bào động vật là",
        "opt_a": "Lục lạp.",
        "opt_a_img": "",
        "opt_b": "Ti thể.",
        "opt_b_img": "",
        "opt_c": "Ribosome.",
        "opt_c_img": "",
        "opt_d": "Bộ máy Golgi.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Lục lạp chứa diệp lục quang hợp có ở tế bào thực vật, không có ở tế bào động vật."
    },
    {
        "id": 27,
        "book": "Sinh học 10",
        "topic": "Chương 3: Trao đổi chất qua màng và truyền tin tế bào",
        "lesson": "Bài 10: Trao đổi chất qua màng tế bào",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Môi trường có nồng độ chất tan bên ngoài tế bào LỚN HƠN nồng độ chất tan bên trong tế bào được gọi là",
        "opt_a": "Môi trường ưu trương.",
        "opt_a_img": "",
        "opt_b": "Môi trường nhược trương.",
        "opt_b_img": "",
        "opt_c": "Môi trường đẳng trương.",
        "opt_c_img": "",
        "opt_d": "Môi trường bão hòa.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Môi trường ưu trương có nồng độ chất tan ngoại bào cao hơn nội bào."
    },
    {
        "id": 28,
        "book": "Sinh học 10",
        "topic": "Chương 3: Trao đổi chất qua màng và truyền tin tế bào",
        "lesson": "Bài 10: Trao đổi chất qua màng tế bào",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Thông hiểu",
        "indicator": "NT3",
        "content": "Hiện tượng khối chất tế bào co lại và tách khỏi thành tế bào ở tế bào thực vật đặt trong môi trường ưu trương gọi là",
        "opt_a": "Hiện tượng co nguyên sinh.",
        "opt_a_img": "",
        "opt_b": "Hiện tượng phản co nguyên sinh.",
        "opt_b_img": "",
        "opt_c": "Hiện tượng nhập bào.",
        "opt_c_img": "",
        "opt_d": "Hiện tượng trương nước.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Co nguyên sinh là hiện tượng màng sinh chất co tách khỏi thành tế bào do tế bào mất nước trong môi trường ưu trương."
    },
    {
        "id": 29,
        "book": "Sinh học 10",
        "topic": "Chương 4: Chuyển hóa năng lượng trong tế bào",
        "lesson": "Bài 13: Khái quát về chuyển hóa vật chất và năng lượng",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Dạng năng lượng chủ yếu lưu trữ và sử dụng trực tiếp trong các hoạt động sống của tế bào là",
        "opt_a": "hóa năng trong ATP.",
        "opt_a_img": "",
        "opt_b": "nhiệt năng.",
        "opt_b_img": "",
        "opt_c": "cơ năng.",
        "opt_c_img": "",
        "opt_d": "quang năng.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "ATP là đồng tiền năng lượng hóa năng duy trì hoạt động sống tế bào."
    },
    {
        "id": 30,
        "book": "Sinh học 10",
        "topic": "Chương 4: Chuyển hóa năng lượng trong tế bào",
        "lesson": "Bài 13: Khái quát về chuyển hóa vật chất và năng lượng",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Vùng cấu trúc không gian đặc biệt của enzyme liên kết khớp với cơ chất được gọi là",
        "opt_a": "trung tâm hoạt động.",
        "opt_a_img": "",
        "opt_b": "trung tâm điều hòa.",
        "opt_b_img": "",
        "opt_c": "chuỗi polypeptide.",
        "opt_c_img": "",
        "opt_d": "coenzyme.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Trung tâm hoạt động là vùng cấu trúc liên kết đặc hiệu với cơ chất."
    },
    {
        "id": 31,
        "book": "Sinh học 10",
        "topic": "Chương 4: Chuyển hóa năng lượng trong tế bào",
        "lesson": "Bài 14: Phân giải và tổng hợp các chất trong tế bào",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Giai đoạn đường phân của quá trình hô hấp tế bào xảy ra tại vị trí nào?",
        "opt_a": "Tế bào chất.",
        "opt_a_img": "",
        "opt_b": "Mào ti thể.",
        "opt_b_img": "",
        "opt_c": "Chất nền ti thể.",
        "opt_c_img": "",
        "opt_d": "Màng ngoài ti thể.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Đường phân phân giải glucose diễn ra ở tế bào chất."
    },
    {
        "id": 32,
        "book": "Sinh học 10",
        "topic": "Chương 5: Chu kì tế bào và phân bào",
        "lesson": "Bài 16: Chu kì tế bào và nguyên phân",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Sự nhân đôi của DNA và nhiễm sắc thể (NST) diễn ra ở pha nào của kỳ trung gian?",
        "opt_a": "Pha S.",
        "opt_a_img": "",
        "opt_b": "Pha G1.",
        "opt_b_img": "",
        "opt_c": "Pha G2.",
        "opt_c_img": "",
        "opt_d": "Pha M.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Pha S (Synthesis) là pha nhân đôi DNA và NST."
    },
    {
        "id": 33,
        "book": "Sinh học 10",
        "topic": "Chương 5: Chu kì tế bào và phân bào",
        "lesson": "Bài 16: Chu kì tế bào và nguyên phân",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Thông hiểu",
        "indicator": "NT4",
        "content": "Các NST co xoắn cực đại và tập trung thành 1 hàng trên mặt phẳng xích đạo ở kỳ nào của nguyên phân?",
        "opt_a": "kỳ giữa.",
        "opt_a_img": "",
        "opt_b": "kỳ đầu.",
        "opt_b_img": "",
        "opt_c": "kỳ sau.",
        "opt_c_img": "",
        "opt_d": "kỳ cuối.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Ở kỳ giữa nguyên phân, các NST kép co xoắn cực đại xếp 1 hàng trên mặt phẳng xích đạo."
    },
    {
        "id": 34,
        "book": "Sinh học 10",
        "topic": "Chương 5: Chu kì tế bào và phân bào",
        "lesson": "Bài 17: Giảm phân",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Hiện tượng tiếp hợp và trao đổi chéo giữa các chromatid trong cặp NST tương đồng diễn ra ở kỳ nào của giảm phân?",
        "opt_a": "Kỳ đầu I.",
        "opt_a_img": "",
        "opt_b": "Kỳ giữa I.",
        "opt_b_img": "",
        "opt_c": "Kỳ đầu II.",
        "opt_c_img": "",
        "opt_d": "Kỳ giữa II.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Kỳ đầu I xảy ra sự tiếp hợp và trao đổi chéo giữa 2 chromatid khác nguồn gốc."
    },
    {
        "id": 35,
        "book": "Sinh học 10",
        "topic": "Chương 5: Chu kì tế bào và phân bào",
        "lesson": "Bài 17: Giảm phân",
        "qtype": "TNKQ trả lời ngắn",
        "level": "Thông hiểu",
        "indicator": "NT4",
        "content": "Từ 5 tế bào sinh tinh thực hiện giảm phân bình thường sẽ tạo ra tổng cộng bao nhiêu tinh trùng?",
        "opt_a": "",
        "opt_a_img": "",
        "opt_b": "",
        "opt_b_img": "",
        "opt_c": "",
        "opt_c_img": "",
        "opt_d": "",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "20",
        "a_image": "",
        "explanation": "Mỗi tế bào sinh tinh qua giảm phân tạo 4 tinh trùng. 5 x 4 = 20 tinh trùng."
    },
    {
        "id": 36,
        "book": "Sinh học 10",
        "topic": "Chương 5: Chu kì tế bào và phân bào",
        "lesson": "Bài 19: Công nghệ tế bào",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Cơ sở khoa học chính của công nghệ tế bào dựa trên đặc tính di truyền nào của tế bào?",
        "opt_a": "Tính toàn năng, sự biệt hóa và phản biệt hóa.",
        "opt_a_img": "",
        "opt_b": "Tính cảm ứng và vận động.",
        "opt_b_img": "",
        "opt_c": "Khả năng đột biến và biến dị tổ hợp.",
        "opt_c_img": "",
        "opt_d": "Khả năng trao đổi chất và năng lượng.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Công nghệ tế bào dựa trên tính toàn năng, khả năng biệt hóa và phản biệt hóa của tế bào."
    },
    {
        "id": 37,
        "book": "Sinh học 10",
        "topic": "Chương 6: Sinh học vi sinh vật",
        "lesson": "Bài 20: Sự đa dạng và phương pháp nghiên cứu vi sinh vật",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Vi sinh vật sử dụng nguồn năng lượng là ánh sáng và nguồn carbon là CO2 thuộc kiểu dinh dưỡng nào?",
        "opt_a": "Quang tự dưỡng.",
        "opt_a_img": "",
        "opt_b": "Quang dị dưỡng.",
        "opt_b_img": "",
        "opt_c": "Hóa tự dưỡng.",
        "opt_c_img": "",
        "opt_d": "Hóa dị dưỡng.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Nguồn năng lượng ánh sáng + nguồn carbon CO2 = Quang tự dưỡng."
    },
    {
        "id": 38,
        "book": "Sinh học 10",
        "topic": "Chương 6: Sinh học vi sinh vật",
        "lesson": "Bài 21: Trao đổi chất, sinh trưởng và sinh sản ở vi sinh vật",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Thời gian tính từ khi một tế bào vi sinh vật sinh ra cho đến khi tế bào đó phân chia được gọi là gì?",
        "opt_a": "Thời gian thế hệ.",
        "opt_a_img": "",
        "opt_b": "Chu kì tế bào.",
        "opt_b_img": "",
        "opt_c": "Pha tiềm phát.",
        "opt_c_img": "",
        "opt_d": "Pha lũy thừa.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Thời gian thế hệ (g) là khoảng thời gian từ khi 1 tế bào sinh ra đến khi nó phân chia."
    },
    {
        "id": 39,
        "book": "Sinh học 10",
        "topic": "Chương 7: Virus",
        "lesson": "Bài 24: Khái quát về virus",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Đặc điểm cấu tạo cơ bản nào khiến virus khác biệt hoàn toàn với tất cả các nhóm sinh vật khác?",
        "opt_a": "Chưa có cấu trúc tế bào.",
        "opt_a_img": "",
        "opt_b": "Kích thước nhỏ.",
        "opt_b_img": "",
        "opt_c": "Có chứa vật chất di truyền.",
        "opt_c_img": "",
        "opt_d": "Có vỏ protein.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Virus là dạng sống chưa có cấu trúc tế bào, chỉ gồm lõi nucleic acid và vỏ protein."
    },
    {
        "id": 40,
        "book": "Sinh học 10",
        "topic": "Chương 7: Virus",
        "lesson": "Bài 24: Khái quát về virus",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Hai thành phần cấu tạo bắt buộc phải có ở tất cả các loại virus là",
        "opt_a": "Lõi Nucleic acid và vỏ Capsid.",
        "opt_a_img": "",
        "opt_b": "Màng bọc và gai Glycoprotein.",
        "opt_b_img": "",
        "opt_c": "Nhân tế bào và Tế bào chất.",
        "opt_c_img": "",
        "opt_d": "Ti thể và Ribosome.",
        "opt_d_img": "",
        "tf_a": "",
        "tf_ans_a": "Đúng",
        "tf_b": "",
        "tf_ans_b": "Đúng",
        "tf_c": "",
        "tf_ans_c": "Đúng",
        "tf_d": "",
        "tf_ans_d": "Đúng",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Mọi virus đều gồm 2 thành phần cơ bản: lõi nucleic acid (ADN hoặc ARN) và vỏ capsid bằng protein."
    },
    {
        "id": 41,
        "book": "Sinh học 11",
        "topic": "Chương 1: Trao đổi chất và chuyển hóa năng lượng ở sinh vật",
        "lesson": "Bài 1: Khái quát về trao đổi chất và chuyển hóa năng lượng",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Nhóm sinh vật nào sau đây có khả năng tự dưỡng (tự tổng hợp chất hữu cơ từ chất vô cơ)?",
        "opt_a": "Thực vật có hoa và tảo xanh",
        "opt_a_img": "",
        "opt_b": "Động vật ăn thực vật",
        "opt_b_img": "",
        "opt_c": "Nấm men và nấm mốc",
        "opt_c_img": "",
        "opt_d": "Vi khuẩn phân giải trong đất",
        "opt_d_img": "",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Thực vật có hoa và tảo xanh chứa diệp lục có khả năng quang hợp tự dưỡng."
    },
    {
        "id": 42,
        "book": "Sinh học 11",
        "topic": "Chương 1: Trao đổi chất và chuyển hóa năng lượng ở sinh vật",
        "lesson": "Bài 2: Trao đổi nước và khoáng ở thực vật",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Cơ quan chính thực hiện chức năng hấp thụ nước và khoáng ở thực vật trên cạn là:",
        "opt_a": "Rễ (vùng lông hút)",
        "opt_a_img": "",
        "opt_b": "Thân cây (mạch gỗ)",
        "opt_b_img": "",
        "opt_c": "Lá cây (khí khổng)",
        "opt_c_img": "",
        "opt_d": "Hoa và quả",
        "opt_d_img": "",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Miền lông hút của rễ cung cấp diện tích bề mặt lớn giúp rễ chủ động hút nước và khoáng từ dung dịch đất."
    },
    {
        "id": 43,
        "book": "Sinh học 11",
        "topic": "Chương 1: Trao đổi chất và chuyển hóa năng lượng ở sinh vật",
        "lesson": "Bài 4: Quang hợp ở thực vật",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Thông hiểu",
        "indicator": "NT6",
        "content": "Dựa vào đồ thị ánh sáng quang hợp bên dưới, tại điểm bão hòa ánh sáng (Is), nếu tiếp tục tăng cường độ ánh sáng thì cường độ quang hợp sẽ:",
        "opt_a": "Không tăng thêm nữa (đạt ngưỡng tối đa)",
        "opt_a_img": "",
        "opt_b": "Tăng lên nhanh chóng theo tỉ lệ thuận",
        "opt_b_img": "",
        "opt_c": "Giảm ngay xuống mức 0",
        "opt_c_img": "",
        "opt_d": "Tăng gấp đôi",
        "opt_d_img": "",
        "q_image": "data:image/svg+xml;base64,PHN2ZyB3aWR0aD0iNDgwIiBoZWlnaHQ9IjIyMCIgeG1sbnM9Imh0dHA6Ly93d3cudzMub3JnLzIwMDAvc3ZnIj4KICA8cmVjdCB3aWR0aD0iMTAwJSIgaGVpZ2h0PSIxMDAlIiBmaWxsPSIjZjhmOWZhIiByeD0iOCIgc3Ryb2tlPSIjZGVlMmU2Ii8+CiAgPHRleHQgeD0iMjQwIiB5PSIyMiIgZm9udC1zaXplPSIxMyIgZm9udC13ZWlnaHQ9ImJvbGQiIHRleHQtYW5jaG9yPSJtaWRkbGUiIGZpbGw9IiMxOTg3NTQiPsSQ4buSIFRI4buKIOG6ok5IIEjGr+G7nk5HIEPhu6ZBIEPGr+G7nE5HIMSQ4buYIMOBTkggU8OBTkcgxJDhur5OIFFVQU5HIEjhu6JQPC90ZXh0PgogIDxsaW5lIHgxPSI1MCIgeTE9IjE4MCIgeDI9IjQ0MCIgeTI9IjE4MCIgc3Ryb2tlPSIjMjEyNTI5IiBzdHJva2Utd2lkdGg9IjIiLz4KICA8bGluZSB4MT0iNTAiIHkxPSIxODAiIHgyPSI1MCIgeTI9IjQwIiBzdHJva2U9IiMyMTI1MjkiIHN0cm9rZS13aWR0aD0iMiIvPgogIDxwYXRoIGQ9Ik0gNTAsMTgwIFEgMTUwLDgwIDI4MCw3MCBMIDQzMCw3MCIgc3Ryb2tlPSIjMTk4NzU0IiBzdHJva2Utd2lkdGg9IjMiIGZpbGw9Im5vbmUiLz4KICA8Y2lyY2xlIGN4PSIxMTAiIGN5PSIxNDAiIHI9IjQiIGZpbGw9IiNkYzM1NDUiLz4KICA8dGV4dCB4PSIxMTAiIHk9IjE2MCIgZm9udC1zaXplPSIxMCIgZmlsbD0iI2RjMzU0NSIgZm9udC13ZWlnaHQ9ImJvbGQiIHRleHQtYW5jaG9yPSJtaWRkbGUiPsSQaeG7g20gYsO5IMOhbmggc8OhbmcgKEljKTwvdGV4dD4KICA8Y2lyY2xlIGN4PSIyODAiIGN5PSI3MCIgcj0iNCIgZmlsbD0iIzBkNmVmZCIvPgogIDx0ZXh0IHg9IjI4MCIgeT0iNTUiIGZvbnQtc2l6ZT0iMTAiIGZpbGw9IiMwZDZlZmQiIGZvbnQtd2VpZ2h0PSJib2xkIiB0ZXh0LWFuY2hvcj0ibWlkZGxlIj7EkGnhu4NtIGLDo28gaMOyYSDDoW5oIHPDoW5nIChJcyk8L3RleHQ+CiAgPHRleHQgeD0iMjQwIiB5PSIyMDIiIGZvbnQtc2l6ZT0iMTEiIGZpbGw9IiMyMTI1MjkiPkPGsOG7nW5nIMSR4buZIMOhbmggc8OhbmcgKEx1eCk8L3RleHQ+Cjwvc3ZnPg==",
        "answer": "A",
        "a_image": "",
        "explanation": "Tại điểm bão hòa ánh sáng Is, bộ máy quang hợp đã hoạt động tối đa, nếu tăng ánh sáng quá mức còn có thể gây ức chế quang hợp."
    },
    {
        "id": 44,
        "book": "Sinh học 11",
        "topic": "Chương 1: Trao đổi chất và chuyển hóa năng lượng ở sinh vật",
        "lesson": "Bài 4: Quang hợp ở thực vật",
        "qtype": "TNKQ đúng - sai",
        "level": "Thông hiểu",
        "indicator": "NT5",
        "content": "So sánh đặc điểm quang hợp giữa các nhóm thực vật C3, C4 và CAM, các phát biểu sau đây Đúng hay Sai?",
        "tf_a": "Thực vật C4 (ngô, mía, rau dền) có năng suất sinh học cao hơn thực vật C3.",
        "tf_ans_a": "Đúng",
        "tf_b": "Thực vật CAM (xương rồng, dứa) cố định CO2 vào ban đêm để tiết kiệm nước.",
        "tf_ans_b": "Đúng",
        "tf_c": "Chất nhận CO2 đầu tiên ở thực vật C3 là PEP (Phosphoenolpyruvate).",
        "tf_ans_c": "Sai",
        "tf_d": "Thực vật C3 không có hiện tượng hô hấp sáng (quang hô hấp).",
        "tf_ans_d": "Sai",
        "q_image": "",
        "answer": "a) Đúng, b) Đúng, c) Sai, d) Sai",
        "a_image": "",
        "explanation": "Ý c sai vì chất nhận CO2 ở C3 là RuBP. Ý d sai vì thực vật C3 có hiện tượng hô hấp sáng gây lãng phí sản phẩm quang hợp."
    },
    {
        "id": 45,
        "book": "Sinh học 12",
        "topic": "Chương 1: Di truyền phân tử",
        "lesson": "Bài 1: DNA và cơ chế tái bản DNA",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Nhận biết",
        "indicator": "NT1",
        "content": "Trong phân tử DNA mạch kép, các nucleotide giữa hai mạch liên kết với nhau bằng loại liên kết nào?",
        "opt_a": "Liên kết hydrogengen theo nguyên tắc bổ sung (A-T, G-C)",
        "opt_a_img": "",
        "opt_b": "Liên kết phosphodiester bền vững",
        "opt_b_img": "",
        "opt_c": "Liên kết peptide",
        "opt_c_img": "",
        "opt_d": "Liên kết glycoside",
        "opt_d_img": "",
        "q_image": "",
        "answer": "A",
        "a_image": "",
        "explanation": "Các base nitơ giữa hai mạch liên kết với nhau bằng liên kết hydrogengen: A liên kết T bằng 2 liên kết, G liên kết C bằng 3 liên kết."
    },
    {
        "id": 46,
        "book": "Sinh học 12",
        "topic": "Chương 1: Di truyền phân tử",
        "lesson": "Bài 1: DNA và cơ chế tái bản DNA",
        "qtype": "TNKQ nhiều lựa chọn",
        "level": "Thông hiểu",
        "indicator": "NT4",
        "content": "Quan sát sơ đồ chạc tái bản DNA bên dưới và cho biết phát biểu nào sau đây là ĐÚNG về chiều tổng hợp mạch mới:",
        "opt_a": "Enzyme DNA polymerase tổng hợp mạch mới theo chiều từ 5' -> 3' trên cả hai mạch khuôn",
        "opt_a_img": "",
        "opt_b": "Mạch mới luôn được tổng hợp theo chiều từ 3' -> 5'",
        "opt_b_img": "",
        "opt_c": "Đoạn Okazaki được tổng hợp kéo dài từ 3' -> 5'",
        "opt_d": "Mạch khuôn 5'->3' tổng hợp liên tục",
        "opt_d_img": "",
        "q_image": "data:image/svg+xml;base64,PHN2ZyB3aWR0aD0iNTIwIiBoZWlnaHQ9IjIzMCIgeG1sbnM9Imh0dHA6Ly93d3cudzMub3JnLzIwMDAvc3ZnIj4KICA8cmVjdCB3aWR0aD0iMTAwJSIgaGVpZ2h0PSIxMDAlIiBmaWxsPSIjZjhmOWZhIiByeD0iOCIgc3Ryb2tlPSIjZGVlMmU2Ii8+CiAgPHRleHQgeD0iMjYwIiB5PSIyNSIgZm9udC1zaXplPSIxNCIgZm9udC13ZWlnaHQ9ImJvbGQiIHRleHQtYW5jaG9yPSJtaWRkbGUiIGZpbGw9IiMwZDZlZmQiPlPGoCDEkOG7kiBDSOG6oEMgVMOBSSBC4bqiTiBETkEgVsOAIFPhu7AgVOG7lE5HIEjhu6JQIE3huqBDSCBN4buaSTwvdGV4dD4KICA8cGF0aCBkPSJNIDUwLDcwIEwgMjAwLDcwIEwgNDQwLDMwIiBzdHJva2U9IiNkYzM1NDUiIHN0cm9rZS13aWR0aD0iNCIgZmlsbD0ibm9uZSIvPgogIDxwYXRoIGQ9Ik0gNTAsMTUwIEwgMjAwLDE1MCBMIDQ0MCwxOTAiIHN0cm9rZT0iIzBkNmVmZCIgc3Ryb2tlLXdpZHRoPSI0IiBmaWxsPSJub25lIi8+CiAgPHBhdGggZD0iTSAyMDAsNzAgTCA0MjAsMzUiIHN0cm9rZT0iIzE5ODc1NCIgc3Ryb2tlLXdpZHRoPSIzIiBzdHJva2UtZGFzaGFycmF5PSI1LDMiIGZpbGw9Im5vbmUiLz4KICA8Y2lyY2xlIGN4PSIyMDAiIGN5PSIxMTAiIHI9IjE4IiBmaWxsPSIjNmY0MmMxIi8+CiAgPHRleHQgeD0iMjAwIiB5PSIxMTQiIGZvbnQtc2l6ZT0iMTAiIGZpbGw9IiNmZmZmZmYiIGZvbnQtd2VpZ2h0PSJib2xkIiB0ZXh0LWFuY2hvcj0ibWlkZGxlIj5IZWxpY2FzZTwvdGV4dD4KICA8dGV4dCB4PSI0NDUiIHk9IjMwIiBmb250LXNpemU9IjEyIiBmaWxsPSIjZGMzNTQ1IiBmb250LXdlaWdodD0iYm9sZCI+Myc8L3RleHQ+CiAgPHRleHQgeD0iNDQ1IiB5PSIxOTAiIGZvbnQtc2l6ZT0iMTIiIGZpbGw9IiMwZDZlZmQiIGZvbnQtd2VpZ2h0PSJib2xkIj41JzwvdGV4dD4KPC9zdmc+",
        "answer": "A",
        "a_image": "",
        "explanation": "Enzyme DNA polymerase chỉ bổ sung nucleotide vào đầu 3'-OH của mạch polynucleotide đang kéo dài, do đó mạch mới luôn được tổng hợp theo chiều 5' -> 3'."
    },
    {
        "id": 47,
        "book": "Sinh học 12",
        "topic": "Chương 1: Di truyền phân tử",
        "lesson": "Bài 4: Đột biến gene",
        "qtype": "TNKQ đúng - sai",
        "level": "Vận dụng",
        "indicator": "NT7",
        "content": "Sơ đồ dưới đây mô tả cơ chế tác động của hóa chất 5-Bromouracil (5-BU) gây đột biến gene qua các lần nhân đôi DNA:",
        "tf_a": "5-BU là chất tương tự base nitơ thymine, có thể bắt đôi với A hoặc G.",
        "tf_ans_a": "Đúng",
        "tf_b": "Qua 1 lần nhân đôi, 5-BU làm xuất hiện ngay cặp đột biến G-C.",
        "tf_ans_b": "Sai",
        "tf_c": "Cần 3 lần nhân đôi liên tiếp để tạo ra 1 phân tử DNA đột biến G-C từ A-T.",
        "tf_ans_c": "Đúng",
        "tf_d": "Đột biến do 5-BU gây ra thuộc dạng đột biến thay thế cặp nucleotide.",
        "tf_ans_d": "Đúng",
        "q_image": "data:image/svg+xml;base64,PHN2ZyB3aWR0aD0iNTAwIiBoZWlnaHQ9IjE4MCIgeG1sbnM9Imh0dHA6Ly93d3cudzMub3JnLzIwMDAvc3ZnIj4KICA8cmVjdCB3aWR0aD0iMTAwJSIgaGVpZ2h0PSIxMDAlIiBmaWxsPSIjZjhmOWZhIiByeD0iOCIgc3Ryb2tlPSIjZGVlMmU2Ii8+CiAgPHRleHQgeD0iMjUwIiB5PSIyMiIgZm9udC1zaXplPSIxMyIgZm9udC13ZWlnaHQ9ImJvbGQiIHRleHQtYW5jaG9yPSJtaWRkbGUiIGZpbGw9IiNkYzM1NDUiPlPGoCDEkOG7kiBDxqAgQ0jhur4gR8OCWSDEkOG7mFQgQknhur5OIEPhu6ZBIDUtQlJPTU9VUkFDSUwgKDUtQlUpPC90ZXh0PgogIDxyZWN0IHg9IjMwIiB5PSI1MCIgd2lkdGg9IjgwIiBoZWlnaHQ9IjMwIiBmaWxsPSIjZTJlM2U1IiByeD0iNCIvPgogIDx0ZXh0IHg9IjcwIiB5PSI3MCIgZm9udC1zaXplPSIxMiIgZm9udC13ZWlnaHQ9ImJvbGQiIHRleHQtYW5jaG9yPSJtaWRkbGUiPkEgLSBUPC90ZXh0PgogIDx0ZXh0IHg9IjEzMCIgeT0iNzAiIGZvbnQtc2l6ZT0iMTYiIGZpbGw9IiM2Yzc1N2QiPuKGkjwvdGV4dD4KICA8cmVjdCB4PSIxNTAiIHk9IjUwIiB3aWR0aD0iOTAiIGhlaWdodD0iMzAiIGZpbGw9IiNmZmYzY2QiIHJ4PSI0IiBzdHJva2U9IiNmZmMxMDciLz4KICA8dGV4dCB4PSIxOTUiIHk9IjcwIiBmb250LXNpemU9IjEyIiBmb250LXdlaWdodD0iYm9sZCIgZmlsbD0iIzg1NjQwNCIgdGV4dC1hbmNob3I9Im1pZGRsZSI+QSAtIDUtQlU8L3RleHQ+CiAgPHRleHQgeD0iMjU1IiB5PSI3MCIgZm9udC1zaXplPSIxNiIgZmlsbD0iIzZjNzU3ZCI+4oaSPC90ZXh0PgogIDxyZWN0IHg9IjI3NSIgeT0iNTAiIHdpZHRoPSI5MCIgaGVpZ2h0PSIzMCIgZmlsbD0iI2ZmZTVkMCIgcng9IjQiIHN0cm9rZT0iI2ZkN2UxNCIvPgogIDx0ZXh0IHg9IjMyMCIgeT0iNzAiIGZvbnQtc2l6ZT0iMTIiIGZvbnQtd2VpZ2h0PSJib2xkIiBmaWxsPSIjZDk0ODBmIiB0ZXh0LWFuY2hvcj0ibWlkZGxlIj5HIC0gNS1CVTwvdGV4dD4KICA8dGV4dCB4PSIzODAiIHk9IjcwIiBmb250LXNpemU9IjE2IiBmaWxsPSIjNmM3NTdkIj7ihpI8L3RleHQ+CiAgPHJlY3QgeD0iNDAwIiB5PSI1MCIgd2lkdGg9IjgwIiBoZWlnaHQ9IjMwIiBmaWxsPSIjZjhkN2RhIiByeD0iNCIgc3Ryb2tlPSIjZGMzNTQ1Ii8+CiAgPHRleHQgeD0iNDQwIiB5PSI3MCIgZm9udC1zaXplPSIxMiIgZm9udC13ZWlnaHQ9ImJvbGQiIGZpbGw9IiM3MjFjMjQiIHRleHQtYW5jaG9yPSJtaWRkbGUiPkcgLSBDPC90ZXh0PgogIDx0ZXh0IHg9IjI1MCIgeT0iMTI1IiBmb250LXNpemU9IjExIiBmaWxsPSIjNDk1MDU3IiB0ZXh0LWFuY2hvcj0ibWlkZGxlIj5D4bq3cCBiYW4gxJHhuqd1IChBLVQpIOKGkiBOaMOibiDEkcO0aSAxIOKGkiBOaMOibiDEkcO0aSAyIOKGkiBOaMOibiDEkcO0aSAzIChD4bq3cCDEkeG7mXQgYmnhur9uIEctQyk8L3RleHQ+Cjwvc3ZnPg==",
        "answer": "a) Đúng, b) Sai, c) Đúng, d) Đúng",
        "a_image": "",
        "explanation": "Sơ đồ 5-BU: A-T -> A-5BU -> G-5BU -> G-C. Cần 3 lần nhân đôi để hình thành cặp G-C."
    },
    {
        "id": 48,
        "book": "Sinh học 12",
        "topic": "Chương 2: Di truyền nhiễm sắc thể",
        "lesson": "Bài 8: Học thuyết di truyền của Mendel",
        "qtype": "TNKQ trả lời ngắn",
        "level": "Thông hiểu",
        "indicator": "VD1",
        "content": "Cho phép lai P: AaBb × AaBb. Tỉ lệ kiểu hình mang hai tính trạng trội (A-B-) ở đời con F1 là bao nhiêu phần mười sáu (x/16)? Điền số nguyên.",
        "opt_a": "",
        "opt_a_img": "",
        "opt_b": "",
        "opt_b_img": "",
        "opt_c": "",
        "opt_c_img": "",
        "opt_d": "",
        "opt_d_img": "",
        "q_image": "",
        "answer": "9",
        "a_image": "",
        "explanation": "Tỉ lệ kiểu hình F1 là 9 A-B- : 3 A-bb : 3 aaB- : 1 aabb. Số cần điền là 9."
    },
    {
        "id": 49,
        "book": "Sinh học 12",
        "topic": "Chương 1: Di truyền phân tử",
        "lesson": "Bài 1: DNA và cơ chế tái bản DNA",
        "qtype": "Tự luận",
        "level": "Thông hiểu",
        "indicator": "NT6",
        "content": "Trình bày nguyên tắc bán bảo toàn và nguyên tắc bổ sung trong cơ chế tái bản DNA. Vì sao trên một mạch khuôn, mạch mới lại được tổng hợp gián đoạn thành các đoạn Okazaki?",
        "opt_a": "",
        "opt_a_img": "",
        "opt_b": "",
        "opt_b_img": "",
        "opt_c": "",
        "opt_c_img": "",
        "opt_d": "",
        "opt_d_img": "",
        "q_image": "",
        "answer": "Đáp án: - Nguyên tắc bổ sung: A-T, G-C.\n- Nguyên tắc bán bảo toàn: Trong mỗi DNA con có 1 mạch cũ của DNA mẹ và 1 mạch mới tổng hợp.\n- Mạch gián đoạn: Do DNA polymerase chỉ tổng hợp chiều 5'->3' mà 2 mạch khuôn ngược chiều nhau.",
        "a_image": "",
        "explanation": "Thang điểm: NTBS & NBT (1.0đ), Giải thích chiều 5'->3' của DNA polymerase (1.0đ)."
    }
]

# --- HÀM TÍNH TOÁN MA TRẬN & ĐIỂM SỐ ---
def calculate_matrix_stats(rows):
    total_mcq_b = sum(r.get("mcq_b", 0) for r in rows)
    total_mcq_h = sum(r.get("mcq_h", 0) for r in rows)
    total_mcq_vd = sum(r.get("mcq_vd", 0) for r in rows)

    total_tf_b = sum(r.get("tf_b", 0) for r in rows)
    total_tf_h = sum(r.get("tf_h", 0) for r in rows)
    total_tf_vd = sum(r.get("tf_vd", 0) for r in rows)

    total_sa_b = sum(r.get("sa_b", 0) for r in rows)
    total_sa_h = sum(r.get("sa_h", 0) for r in rows)
    total_sa_vd = sum(r.get("sa_vd", 0) for r in rows)

    total_es_b = sum(r.get("es_b", 0) for r in rows)
    total_es_h = sum(r.get("es_h", 0) for r in rows)
    total_es_vd = sum(r.get("es_vd", 0) for r in rows)

    count_b = total_mcq_b + (total_tf_b * 4) + total_sa_b + total_es_b
    count_h = total_mcq_h + (total_tf_h * 4) + total_sa_h + total_es_h
    count_vd = total_mcq_vd + (total_tf_vd * 4) + total_sa_vd + total_es_vd
    total_items = count_b + count_h + count_vd

    score_b = (total_mcq_b * 0.25) + (total_tf_b * 1.0) + (total_sa_b * 0.25) + (total_es_b * 1.0)
    score_h = (total_mcq_h * 0.25) + (total_tf_h * 1.0) + (total_sa_h * 0.25) + (total_es_h * 1.0)
    score_vd = (total_mcq_vd * 0.25) + (total_tf_vd * 1.0) + (total_sa_vd * 0.25) + (total_es_vd * 1.0)
    total_score = score_b + score_h + score_vd

    pct_b = (score_b / total_score * 100) if total_score > 0 else 0
    pct_h = (score_h / total_score * 100) if total_score > 0 else 0
    pct_vd = (score_vd / total_score * 100) if total_score > 0 else 0

    return {
        "total_mcq_b": total_mcq_b, "total_mcq_h": total_mcq_h, "total_mcq_vd": total_mcq_vd,
        "total_tf_b": total_tf_b, "total_tf_h": total_tf_h, "total_tf_vd": total_tf_vd,
        "total_sa_b": total_sa_b, "total_sa_h": total_sa_h, "total_sa_vd": total_sa_vd,
        "total_es_b": total_es_b, "total_es_h": total_es_h, "total_es_vd": total_es_vd,
        "count_b": count_b, "count_h": count_h, "count_vd": count_vd, "total_items": total_items,
        "score_b": round(score_b, 2), "score_h": round(score_h, 2), "score_vd": round(score_vd, 2),
        "total_score": round(total_score, 2),
        "pct_b": round(pct_b, 1), "pct_h": round(pct_h, 1), "pct_vd": round(pct_vd, 1)
    }

# --- HÀM XUẤT EXCEL MA TRẬN ĐỀ CHUẨN ---
def export_matrix_to_excel(matrix_rows, school_name="TRƯỜNG THPT CHUYÊN/CHUẨN", exam_title="ĐỀ KIỂM TRA ĐỊNH KỲ MÔN SINH HỌC"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MaTranDeKiemTra"

    font_title = Font(name="Times New Roman", size=14, bold=True, color="002060")
    font_hdr = Font(name="Times New Roman", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Times New Roman", size=10, bold=True)
    font_regular = Font(name="Times New Roman", size=10)

    fill_hdr = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fill_subhdr = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    fill_total = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='A6A6A6'), right=Side(style='thin', color='A6A6A6'),
        top=Side(style='thin', color='A6A6A6'), bottom=Side(style='thin', color='A6A6A6')
    )
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.cell(row=1, column=1, value=f"{school_name.upper()} - MA TRẬN ĐỀ KIỂM TRA MÔN SINH HỌC").font = font_title
    ws.cell(row=2, column=1, value=f"Dành cho: {exam_title} (Khung GDPT 2018)").font = Font(name="Times New Roman", size=11, italic=True)

    headers_l1 = ["TT", "Chủ đề / Chương", "Nội dung / Đơn vị kiến thức", 
                  "TNKQ nhiều lựa chọn (0.25đ/câu)", "", "", 
                  "TNKQ Đúng - Sai (1.0đ/câu 4 ý)", "", "", 
                  "TNKQ trả lời ngắn (0.25đ/câu)", "", "", 
                  "Tự luận (1.0đ/câu)", "", "", 
                  "Tổng số câu/lệnh", "Tổng điểm", "Tỉ lệ %"]
    ws.append(headers_l1)

    headers_l2 = ["", "", "", "Biết", "Hiểu", "VD", "Biết", "Hiểu", "VD", "Biết", "Hiểu", "VD", "Biết", "Hiểu", "VD", "", "", ""]
    ws.append(headers_l2)

    for r in [3, 4]:
        for c in range(1, 18):
            cell = ws.cell(row=r, column=c)
            cell.font = font_hdr
            cell.fill = fill_hdr if r==3 else fill_subhdr
            cell.alignment = align_center
            cell.border = thin_border

    ws.merge_cells("A3:A4")
    ws.merge_cells("B3:B4")
    ws.merge_cells("C3:C4")
    ws.merge_cells("D3:F3")
    ws.merge_cells("G3:I3")
    ws.merge_cells("J3:L3")
    ws.merge_cells("M3:O3")
    ws.merge_cells("P3:P4")
    ws.merge_cells("Q3:Q4")

    for idx, row in enumerate(matrix_rows, start=1):
        mcq_b, mcq_h, mcq_vd = row.get("mcq_b", 0), row.get("mcq_h", 0), row.get("mcq_vd", 0)
        tf_b, tf_h, tf_vd = row.get("tf_b", 0), row.get("tf_h", 0), row.get("tf_vd", 0)
        sa_b, sa_h, sa_vd = row.get("sa_b", 0), row.get("sa_h", 0), row.get("sa_vd", 0)
        es_b, es_h, es_vd = row.get("es_b", 0), row.get("es_h", 0), row.get("es_vd", 0)

        items_cnt = (mcq_b + mcq_h + mcq_vd) + (tf_b + tf_h + tf_vd)*4 + (sa_b + sa_h + sa_vd) + (es_b + es_h + es_vd)
        row_score = (mcq_b + mcq_h + mcq_vd)*0.25 + (tf_b + tf_h + tf_vd)*1.0 + (sa_b + sa_h + sa_vd)*0.25 + (es_b + es_h + es_vd)*1.0

        r_data = [idx, row.get("topic", ""), row.get("lesson", ""),
                  mcq_b, mcq_h, mcq_vd, tf_b, tf_h, tf_vd, sa_b, sa_h, sa_vd, es_b, es_h, es_vd,
                  items_cnt, row_score, ""]
        ws.append(r_data)
        curr_r = ws.max_row
        for c_idx in range(1, 18):
            cell = ws.cell(row=curr_r, column=c_idx)
            cell.font = font_regular
            cell.border = thin_border
            if c_idx not in [2, 3]:
                cell.alignment = align_center

    stats = calculate_matrix_stats(matrix_rows)
    tot_row = ["", "TỔNG CỘNG", "",
               stats["total_mcq_b"], stats["total_mcq_h"], stats["total_mcq_vd"],
               stats["total_tf_b"], stats["total_tf_h"], stats["total_tf_vd"],
               stats["total_sa_b"], stats["total_sa_h"], stats["total_sa_vd"],
               stats["total_es_b"], stats["total_es_h"], stats["total_es_vd"],
               stats["total_items"], stats["total_score"], "100%"]
    ws.append(tot_row)
    last_r = ws.max_row
    for c_idx in range(1, 18):
        cell = ws.cell(row=last_r, column=c_idx)
        cell.font = font_bold
        cell.fill = fill_total
        cell.border = thin_border
        cell.alignment = align_center

    col_widths = [5, 25, 32, 7, 7, 7, 7, 7, 7, 7, 7, 7, 7, 7, 7, 14, 12, 10]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    out_buf = io.BytesIO()
    wb.save(out_buf)
    out_buf.seek(0)
    return out_buf

# --- HÀM XUẤT EXCEL BẢN ĐẶC TẢ ---
def export_spec_to_excel(matrix_rows, curriculum_data, school_name="TRƯỜNG THPT CHUYÊN/CHUẨN", exam_title="BẢN ĐẶC TẢ ĐỀ KIỂM TRA"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BanDacTaDeKiemTra"

    font_title = Font(name="Times New Roman", size=14, bold=True, color="002060")
    font_hdr = Font(name="Times New Roman", size=10, bold=True, color="FFFFFF")
    font_regular = Font(name="Times New Roman", size=10)

    fill_hdr = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='A6A6A6'), right=Side(style='thin', color='A6A6A6'),
        top=Side(style='thin', color='A6A6A6'), bottom=Side(style='thin', color='A6A6A6')
    )
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.cell(row=1, column=1, value=f"{school_name.upper()} - BẢN ĐẶC TẢ ĐỀ KIỂM TRA MÔN SINH HỌC").font = font_title
    ws.cell(row=2, column=1, value=f"Dành cho: {exam_title} (Khung GDPT 2018)").font = Font(name="Times New Roman", size=11, italic=True)

    headers = ["TT", "Chủ đề / Chương", "Nội dung / Bài học", "Yêu cầu cần đạt (Gắn mã chỉ báo)", "TNKQ nhiều lựa chọn", "TNKQ Đúng - Sai", "TNKQ trả lời ngắn", "Tự luận"]
    ws.append(headers)

    for c in range(1, 9):
        cell = ws.cell(row=3, column=c)
        cell.font = font_hdr
        cell.fill = fill_hdr
        cell.alignment = align_center
        cell.border = thin_border

    for idx, row in enumerate(matrix_rows, start=1):
        book, topic, lesson = row.get("book", ""), row.get("topic", ""), row.get("lesson", "")
        yccd_list = curriculum_data.get(book, {}).get(topic, {}).get(lesson, [])

        nb_text = "\n".join([f"• [{y['code']}] {y['text']}" for y in yccd_list if y.get('level') == "Nhận biết"]) or "• Nhận biết các khái niệm cơ bản."
        th_text = "\n".join([f"• [{y['code']}] {y['text']}" for y in yccd_list if y.get('level') == "Thông hiểu"]) or "• Giải thích cơ chế, vai trò."
        vd_text = "\n".join([f"• [{y['code']}] {y['text']}" for y in yccd_list if y.get('level') == "Vận dụng"]) or "• Vận dụng kiến thức thực tiễn."

        full_yccd = f"NHẬN BIẾT:\n{nb_text}\n\nTHÔNG HIỂU:\n{th_text}\n\nVẬN DỤNG:\n{vd_text}"

        mcq_str = f"Biết: {row.get('mcq_b',0)}, Hiểu: {row.get('mcq_h',0)}, VD: {row.get('mcq_vd',0)}"
        tf_str = f"Biết: {row.get('tf_b',0)}, Hiểu: {row.get('tf_h',0)}, VD: {row.get('tf_vd',0)}"
        sa_str = f"Biết: {row.get('sa_b',0)}, Hiểu: {row.get('sa_h',0)}, VD: {row.get('sa_vd',0)}"
        es_str = f"Biết: {row.get('es_b',0)}, Hiểu: {row.get('es_h',0)}, VD: {row.get('es_vd',0)}"

        r_data = [idx, topic, lesson, full_yccd, mcq_str, tf_str, sa_str, es_str]
        ws.append(r_data)
        curr_r = ws.max_row
        for c_idx in range(1, 9):
            cell = ws.cell(row=curr_r, column=c_idx)
            cell.font = font_regular
            cell.border = thin_border
            if c_idx not in [2, 3, 4]:
                cell.alignment = align_center

    col_widths = [5, 22, 28, 50, 18, 18, 18, 18]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    out_buf = io.BytesIO()
    wb.save(out_buf)
    out_buf.seek(0)
    return out_buf

# --- HÀM TẠO EXCEL MẪU NHẬP CÂU HỎI TÍCH HỢP ẢNH ĐÁP ÁN A, B, C, D ---
def generate_question_excel_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MauSoanCauHoi"

    header_font = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'), bottom=Side(style='thin', color='CCCCCC')
    )
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = [
        "Mã_ID", "Khối_Sách", "Chương_Chủ_đề", "Bài_học", "Dạng_thức", "Mức_độ", "Mã_chỉ_báo",
        "Nội_dung_câu_hỏi", "Phương_án_A", "Phương_án_B", "Phương_án_C", "Phương_án_D",
        "Ý_a", "Ý_b", "Ý_c", "Ý_d", "Đáp_án_Chính_Xác", "Lời_giải_chi_tiết", "Ảnh_câu_hỏi_Base64", "Ảnh_đáp_án_Base64",
        "Ảnh_Phương_án_A_Base64", "Ảnh_Phương_án_B_Base64", "Ảnh_Phương_án_C_Base64", "Ảnh_Phương_án_D_Base64"
    ]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border

    samples = [
        [1, "Sinh học 12", "Chương 1: Di truyền phân tử", "Bài 1: DNA và cơ chế tái bản DNA", "TNKQ nhiều lựa chọn", "Nhận biết", "NT1",
         "Trong phân tử DNA mạch kép, liên kết hydrogengen được hình thành giữa các base nitơ theo nguyên tắc bổ sung nào?",
         "A liên kết với T (2 liên kết), G liên kết với C (3 liên kết)", "A liên kết với C (2 liên kết), G liên kết với T (3 liên kết)",
         "A liên kết với U (2 liên kết), G liên kết với C (3 liên kết)", "A liên kết với T (3 liên kết), G liên kết với C (2 liên kết)",
         "", "", "", "", "A", "Theo mô hình Watson-Crick, A liên kết với T bằng 2 liên kết, G liên kết với C bằng 3 liên kết.", "", "", "", "", "", ""],
        [2, "Sinh học 10", "Chương 1: Thành phần hóa học của tế bào", "Bài 5: Các phân tử sinh học", "TNKQ nhiều lựa chọn", "Thông hiểu", "NT5",
         "Sơ đồ nào sau đây mô tả ĐÚNG cấu trúc Deoxyribonucleotide cấu tạo nên DNA?",
         "Sơ đồ Deoxyribonucleotide", "Sơ đồ Ribonucleotide Uracil", "Sơ đồ Amino acid", "Sơ đồ Glucose",
         "", "", "", "", "A", "Deoxyribonucleotide gồm Phosphate, đường Deoxyribose và Base Adenine.", "", "", "data:image/svg+xml;base64,...", "data:image/svg+xml;base64,...", "", ""]
    ]

    for row in samples:
        ws.append(row)

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = 20
    ws.column_dimensions['H'].width = 45
    ws.column_dimensions['Q'].width = 25

    out_buf = io.BytesIO()
    wb.save(out_buf)
    out_buf.seek(0)
    return out_buf

# --- HÀM XUẤT CÂU HỎI RA EXCEL ---
def export_questions_to_excel(questions):
    df_q = pd.DataFrame(questions)
    buf_q = io.BytesIO()
    with pd.ExcelWriter(buf_q, engine='openpyxl') as writer:
        df_q.to_excel(writer, index=False, sheet_name="NganHangCauHoi")
    buf_q.seek(0)
    return buf_q

# --- HÀM XUẤT NGÂN HÀNG CÂU HỎI RA WORD (.DOCX) ---


def append_image_to_docx(doc_or_para, img_str, width_inches=3.5):
    """Tự động chuyển đổi ảnh (Data URI, Base64, SVG XML) và chèn vào tệp Word docx."""
    if not img_str or not isinstance(img_str, str) or not img_str.strip():
        return False
    try:
        img_str = img_str.strip()
        img_bytes = None
        
        if img_str.startswith('data:image/'):
            header, b64_data = img_str.split(',', 1)
            raw_bytes = base64.b64decode(b64_data)
            if 'svg' in header:
                img_bytes = cairosvg.svg2png(bytestring=raw_bytes)
            else:
                img_bytes = raw_bytes
        elif img_str.startswith('<svg') or '</svg>' in img_str:
            img_bytes = cairosvg.svg2png(bytestring=img_str.encode('utf-8'))
        elif len(img_str) > 50:
            try:
                raw_bytes = base64.b64decode(img_str)
                if raw_bytes.startswith(b'<svg') or b'xmlns' in raw_bytes[:100]:
                    img_bytes = cairosvg.svg2png(bytestring=raw_bytes)
                else:
                    img_bytes = raw_bytes
            except Exception:
                pass
                
        if img_bytes:
            if hasattr(doc_or_para, 'add_paragraph'):
                p = doc_or_para.add_paragraph()
            else:
                p = doc_or_para
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(io.BytesIO(img_bytes), width=Inches(width_inches))
            return True
    except Exception as e:
        pass
    return False

# --- HÀM XUẤT NGÂN HÀNG CÂU HỎI RA WORD (.DOCX) ---
def export_questions_to_word(questions):
    doc_q = docx.Document()
    
    p_title = doc_q.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("NGÂN HÀNG CÂU HỎI MÔN SINH HỌC GDPT 2018")
    r_t.bold = True
    r_t.font.size = Pt(16)
    
    for idx, q_item in enumerate(questions, start=1):
        qtype_str = q_item.get('qtype', '')
        level_str = q_item.get('level', '')
        ind_str = q_item.get('indicator', '')
        lesson_str = q_item.get('lesson', '')
        
        p_q = doc_q.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)
        
        r_head = p_q.add_run(f"Câu {idx} [{qtype_str}] [{level_str}] [{ind_str}] - {lesson_str}:\n")
        r_head.bold = True
        
        p_q.add_run(f"{q_item.get('content', '')}\n")
        
        # Ảnh câu hỏi
        if q_item.get('q_image'):
            append_image_to_docx(doc_q, q_item['q_image'], width_inches=4.2)
            
        # Phương án trắc nghiệm
        if qtype_str == "TNKQ nhiều lựa chọn":
            opts = [('A', 'opt_a', 'opt_a_img'), ('B', 'opt_b', 'opt_b_img'), 
                    ('C', 'opt_c', 'opt_c_img'), ('D', 'opt_d', 'opt_d_img')]
            for label, key_txt, key_img in opts:
                txt = q_item.get(key_txt, '').strip()
                img = q_item.get(key_img, '')
                if txt or img:
                    p_opt = doc_q.add_paragraph()
                    p_opt.paragraph_format.space_after = Pt(2)
                    p_opt.add_run(f"{label}. ").bold = True
                    if txt:
                        p_opt.add_run(txt)
                    if img:
                        append_image_to_docx(p_opt, img, width_inches=2.8)
                        
        elif qtype_str == "TNKQ đúng - sai":
            tf_items = [('a', 'tf_a', 'tf_ans_a'), ('b', 'tf_b', 'tf_ans_b'), 
                        ('c', 'tf_c', 'tf_ans_c'), ('d', 'tf_d', 'tf_ans_d')]
            for label, key_txt, key_ans in tf_items:
                txt = q_item.get(key_txt, '').strip()
                ans = q_item.get(key_ans, '').strip()
                if txt:
                    p_tf = doc_q.add_paragraph()
                    p_tf.paragraph_format.space_after = Pt(2)
                    p_tf.add_run(f"{label}) {txt} ")
                    if ans:
                        r_a = p_tf.add_run(f"[{ans}]")
                        r_a.bold = True
                        
        # Đáp án
        p_ans = doc_q.add_paragraph()
        p_ans.paragraph_format.space_after = Pt(2)
        ans_str = q_item.get('answer', '')
        if not ans_str and qtype_str == "TNKQ đúng - sai":
            ans_str = f"a) {q_item.get('tf_ans_a','Đúng')}, b) {q_item.get('tf_ans_b','Đúng')}, c) {q_item.get('tf_ans_c','Đúng')}, d) {q_item.get('tf_ans_d','Đúng')}"
        r_ans = p_ans.add_run(f"Đáp án: {ans_str}\n")
        r_ans.italic = True
        r_ans.bold = True
        
        # Ảnh lời giải
        if q_item.get('a_image'):
            append_image_to_docx(doc_q, q_item['a_image'], width_inches=4.2)
            
        # Lời giải chi tiết
        if q_item.get('explanation'):
            p_exp = doc_q.add_paragraph()
            p_exp.paragraph_format.space_after = Pt(4)
            r_exp = p_exp.add_run(f"Lời giải chi tiết: {q_item.get('explanation')}")
            r_exp.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            
        doc_q.add_paragraph("-" * 55)
        
    doc_q_buf = io.BytesIO()
    doc_q.save(doc_q_buf)
    doc_q_buf.seek(0)
    return doc_q_buf

# --- HÀM XUẤT ĐỀ THI & ĐÁP ÁN RA WORD (.DOCX) ---
def export_exam_to_word(exam_parts, school_name="TRƯỜNG THPT CHUYÊN/CHUẨN", exam_title="ĐỀ KIỂM TRA ĐỊNH KỲ MÔN SINH HỌC", exam_duration="45 phút", exam_code="101"):
    doc = docx.Document()
    
    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sch = p_hdr.add_run(f"{school_name.upper()}\n")
    r_sch.bold = True
    r_sch.font.size = Pt(12)
    
    r_title = p_hdr.add_run(f"{exam_title.upper()}\n")
    r_title.bold = True
    r_title.font.size = Pt(14)
    
    r_sub = p_hdr.add_run(f"Môn: SINH HỌC - GDPT 2018 | Thời gian làm bài: {exam_duration} | Mã đề thi: {exam_code}\n")
    r_sub.italic = True
    r_sub.font.size = Pt(10)
    
    doc.add_paragraph("=" * 60)

    q_counter = 1

    if exam_parts.get("Part1_MCQ"):
        doc.add_heading("PHẦN I. Câu hỏi trắc nghiệm nhiều lựa chọn (Mỗi câu trả lời đúng được 0,25 điểm)", level=2)
        for q in exam_parts["Part1_MCQ"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            p.add_run(f"Câu {q_counter} [{q.get('level','')}]: ").bold = True
            p.add_run(f"{q.get('content','')}\n")
            
            if q.get("q_image"):
                append_image_to_docx(doc, q["q_image"], width_inches=4.2)
                
            opts = [('A', 'opt_a', 'opt_a_img'), ('B', 'opt_b', 'opt_b_img'), 
                    ('C', 'opt_c', 'opt_c_img'), ('D', 'opt_d', 'opt_d_img')]
            for label, key_txt, key_img in opts:
                txt = q.get(key_txt, '').strip()
                img = q.get(key_img, '')
                if txt or img:
                    p_opt = doc.add_paragraph()
                    p_opt.paragraph_format.space_after = Pt(2)
                    p_opt.add_run(f"{label}. ").bold = True
                    if txt:
                        p_opt.add_run(txt)
                    if img:
                        append_image_to_docx(p_opt, img, width_inches=2.8)
            q_counter += 1

    if exam_parts.get("Part2_TF"):
        doc.add_heading("PHẦN II. Câu hỏi trắc nghiệm Đúng/Sai (Mỗi câu chọn đúng 1 ý được 0,1đ; 2 ý 0,25đ; 3 ý 0,5đ; 4 ý 1,0đ)", level=2)
        for q in exam_parts.get("Part2_TF", []):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            p.add_run(f"Câu {q_counter} [{q.get('level','')}]: ").bold = True
            p.add_run(f"{q.get('content','')}\n")
            
            if q.get("q_image"):
                append_image_to_docx(doc, q["q_image"], width_inches=4.2)
                
            tf_items = [('a', 'tf_a'), ('b', 'tf_b'), ('c', 'tf_c'), ('d', 'tf_d')]
            for label, key_txt in tf_items:
                txt = q.get(key_txt, '').strip()
                if txt:
                    p_tf = doc.add_paragraph()
                    p_tf.paragraph_format.space_after = Pt(2)
                    p_tf.add_run(f"{label}) {txt}")
            q_counter += 1

    if exam_parts.get("Part3_SA"):
        doc.add_heading("PHẦN III. Câu hỏi trắc nghiệm trả lời ngắn (Mỗi câu trả lời đúng được 0,25 điểm hoặc 0,5 điểm)", level=2)
        for q in exam_parts.get("Part3_SA", []):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            p.add_run(f"Câu {q_counter} [{q.get('level','')}]: ").bold = True
            p.add_run(f"{q.get('content','')}\n")
            if q.get("q_image"):
                append_image_to_docx(doc, q["q_image"], width_inches=4.2)
            q_counter += 1

    if exam_parts.get("Part4_ES"):
        doc.add_heading("PHẦN IV. Câu hỏi tự luận", level=2)
        for q in exam_parts.get("Part4_ES", []):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            p.add_run(f"Câu {q_counter} [{q.get('level','')}]: ").bold = True
            p.add_run(f"{q.get('content','')}\n")
            if q.get("q_image"):
                append_image_to_docx(doc, q["q_image"], width_inches=4.2)
            q_counter += 1

    doc.add_page_break()
    p_ans_hdr = doc.add_paragraph()
    p_ans_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ah = p_ans_hdr.add_run(f"ĐÁP ÁN VÀ HƯỚNG DẪN CHẤM CHI TIẾT - MÃ ĐỀ {exam_code}\n")
    r_ah.bold = True
    r_ah.font.size = Pt(14)

    all_questions = exam_parts.get("Part1_MCQ", []) + exam_parts.get("Part2_TF", []) + exam_parts.get("Part3_SA", []) + exam_parts.get("Part4_ES", [])
    for idx, q in enumerate(all_questions, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(2)
        p.add_run(f"Câu {idx}: ").bold = True
        
        ans_str = q.get('answer', '')
        if not ans_str and q.get('qtype') == "TNKQ đúng - sai":
            ans_str = f"a) {q.get('tf_ans_a','Đúng')}, b) {q.get('tf_ans_b','Đúng')}, c) {q.get('tf_ans_c','Đúng')}, d) {q.get('tf_ans_d','Đúng')}"
        
        r_ans = p.add_run(f"Đáp án: {ans_str}")
        r_ans.italic = True
        r_ans.bold = True
        
        if q.get("a_image"):
            append_image_to_docx(doc, q["a_image"], width_inches=4.2)
            
        if q.get("explanation"):
            p_exp = doc.add_paragraph()
            p_exp.paragraph_format.space_after = Pt(4)
            p_exp.add_run(f"Lời giải: {q.get('explanation')}")

    out_buf = io.BytesIO()
    doc.save(out_buf)
    out_buf.seek(0)
    return out_buf



# --- KHỞI TẠO SESSION STATE ---
if "rename_history" not in st.session_state:
    st.session_state.rename_history = []

if "curriculum" not in st.session_state:
    st.session_state.curriculum = DEFAULT_CURRICULUM

if "matrix_rows" not in st.session_state:
    st.session_state.matrix_rows = [
        {
            "id": 1,
            "book": "Sinh học 12",
            "topic": "Chương 1: Di truyền phân tử",
            "lesson": "Bài 1: DNA và cơ chế tái bản DNA",
            "mcq_b": 2, "mcq_h": 1, "mcq_vd": 0,
            "tf_b": 0, "tf_h": 1, "tf_vd": 0,
            "sa_b": 0, "sa_h": 0, "sa_vd": 1,
            "es_b": 0, "es_h": 0, "es_vd": 0
        },
        {
            "id": 2,
            "book": "Sinh học 12",
            "topic": "Chương 1: Di truyền phân tử",
            "lesson": "Bài 4: Đột biến gene",
            "mcq_b": 1, "mcq_h": 1, "mcq_vd": 0,
            "tf_b": 0, "tf_h": 1, "tf_vd": 0,
            "sa_b": 0, "sa_h": 0, "sa_vd": 1,
            "es_b": 0, "es_h": 0, "es_vd": 0
        }
    ]

if "questions" not in st.session_state:
    st.session_state.questions = DEFAULT_QUESTIONS_DATABASE

# --- GIAO DIỆN CHÍNH APP ---
st.title("🧬 Ứng Dụng Ma Trận, Bản Đặc Tả & Tạo Đề Thi Sinh Học GDPT 2018")
st.caption("Công cụ chuẩn hóa cho Giáo viên Sinh học THPT - Tự động liên kết Ma trận, Bản đặc tả, Ngân hàng câu hỏi & Xuất Đề thi Word/Excel")

tab1, tab2, tab3, tab4, tab5 = st.tabs([
    "📚 1. Chương Trình & YCCĐ",
    "📊 2. Thiết Kế Ma Trận Đề",
    "📋 3. Thiết Kế Bản Đặc Tả",
    "📝 4. Ngân Hàng Câu Hỏi & Tạo Đề",
    "❓ 5. Hướng Dẫn & Tra Cứu"
])

# -------------------------------------------------------------
# TAB 1: KHUNG CHƯƠNG TRÌNH & YCCĐ
# -------------------------------------------------------------
with tab1:
    st.header("1. Khung Chương Trình & Yêu Cầu Cần Đạt (YCCĐ) GDPT 2018")
    st.caption("Tích hợp 100% Khung chương trình Sinh học 10, 11, 12 và Chuyên đề học tập (123 bài học, 437 YCCĐ)")

    # 1. BỘ LỌC CHỌN SÁCH - CHƯƠNG - BÀI
    st.subheader("🎯 1. Tra cứu & Chọn Bài học")
    c_bk, c_tp = st.columns(2)
    with c_bk:
        selected_book = st.selectbox("Chọn Cấp sách / Bộ sách:", list(st.session_state.curriculum.keys()), key="t1_book_select")
    
    topics_in_book = list(st.session_state.curriculum[selected_book].keys())
    with c_tp:
        selected_topic = st.selectbox("Chọn Chương / Chủ đề:", topics_in_book, key="t1_topic_select")

    lessons_in_topic = list(st.session_state.curriculum[selected_book][selected_topic].keys())
    selected_lesson = st.selectbox("Chọn Nội dung / Đơn vị kiến thức (Bài học):", lessons_in_topic, key="t1_lesson_select")

    # HIỂN THỊ THẺ THÔNG TIN BÀI HỌC ĐANG CHỌN RÕ RÀNG (KHẮC PHỤC HIỂN THỊ CHÍNH XÁC TÊN BÀI HỌC)
    st.markdown(f"""
    <div style="background-color: #f8f9fa; padding: 16px; border-radius: 8px; border-left: 5px solid #0d6efd; margin-top: 10px; margin-bottom: 20px; box-shadow: 0 1px 3px rgba(0,0,0,0.1);">
        <h3 style="margin: 0; color: #0d6efd; font-size: 1.25rem;">📖 BÀI HỌC ĐANG CHỌN: {selected_lesson}</h3>
        <p style="margin: 6px 0 0 0; color: #495057; font-size: 0.95rem;">
            <b>📂 Chương / Chủ đề:</b> {selected_topic} &nbsp;|&nbsp; 
            <b>📚 Cấp sách / Bộ sách:</b> {selected_book}
        </p>
    </div>
    """, unsafe_allow_html=True)

    # 2. KHU VỰC CHỈNH SỬA TÊN BÀI HỌC & TÊN CHƯƠNG / CHỦ ĐỀ
    st.subheader("🛠️ 2. Chỉnh Sửa Thông Tin Tên Bài Học & Tên Chương")
    c_edit_ls, c_edit_tp = st.columns(2)
    
    with c_edit_ls:
        with st.expander("✏️ Sửa Tên Bài Học Đang Chọn"):
            new_ls_input = st.text_input("Nhập tên bài học mới:", value=selected_lesson, key="inp_new_ls")
            if st.button("💾 Lưu Tên Bài Học Mới", key="btn_save_ls"):
                if new_ls_input.strip() and new_ls_input.strip() != selected_lesson:
                    old_ls = selected_lesson
                    new_ls = new_ls_input.strip()
                    
                    # Cập nhật trong session_state.curriculum
                    curr_topic_dict = st.session_state.curriculum[selected_book][selected_topic]
                    new_topic_dict = {}
                    for k, v in curr_topic_dict.items():
                        if k == old_ls:
                            new_topic_dict[new_ls] = v
                        else:
                            new_topic_dict[k] = v
                    st.session_state.curriculum[selected_book][selected_topic] = new_topic_dict
                    
                    # Ghi lịch sử đổi tên để đồng bộ sang các Tab khác
                    if "rename_history" not in st.session_state:
                        st.session_state.rename_history = []
                    st.session_state.rename_history.append({
                        "type": "lesson",
                        "book": selected_book,
                        "topic": selected_topic,
                        "old_name": old_ls,
                        "new_name": new_ls
                    })
                    st.success(f"✅ Đã đổi tên bài học thành: **{new_ls}**")
                    st.rerun()

    with c_edit_tp:
        with st.expander("✏️ Sửa Tên Chương / Chủ Đề Đang Chọn"):
            new_tp_input = st.text_input("Nhập tên chương / chủ đề mới:", value=selected_topic, key="inp_new_tp")
            if st.button("💾 Lưu Tên Chương Mới", key="btn_save_tp"):
                if new_tp_input.strip() and new_tp_input.strip() != selected_topic:
                    old_tp = selected_topic
                    new_tp = new_tp_input.strip()
                    
                    curr_book_dict = st.session_state.curriculum[selected_book]
                    new_book_dict = {}
                    for k, v in curr_book_dict.items():
                        if k == old_tp:
                            new_book_dict[new_tp] = v
                        else:
                            new_book_dict[k] = v
                    st.session_state.curriculum[selected_book] = new_book_dict
                    
                    if "rename_history" not in st.session_state:
                        st.session_state.rename_history = []
                    st.session_state.rename_history.append({
                        "type": "topic",
                        "book": selected_book,
                        "old_name": old_tp,
                        "new_name": new_tp
                    })
                    st.success(f"✅ Đã đổi tên chương thành: **{new_tp}**")
                    st.rerun()

    st.divider()

    # 3. DANH SÁCH & CHỈNH SỬA YÊU CẦU CẦN ĐẠT (YCCĐ)
    st.subheader(f"📋 3. Các Yêu cầu cần đạt (YCCĐ) của: {selected_lesson}")
    yccds = st.session_state.curriculum[selected_book][selected_topic][selected_lesson]

    if not yccds:
        st.info("Bài học này hiện chưa có YCCĐ nào. Hãy bổ sung YCCĐ mới ở khung bên dưới.")
    else:
        for idx, y in enumerate(yccds):
            lvl_color = "🟢" if y['level'] == "Nhận biết" else ("🔵" if y['level'] == "Thông hiểu" else "🔴")
            
            c_y1, c_y2, c_y3, c_y4 = st.columns([0.8, 6.4, 1.4, 1.4])
            with c_y1:
                st.markdown(f"**#{idx+1}**")
            with c_y2:
                # Hiển thị YCCĐ an toàn tránh bị trình duyệt tự động dịch/biến đổi từ ngữ
                st.markdown(f"<div style='font-size: 0.98rem; line-height: 1.5; color: #1a1a1a;'><b>[{y['code']}]</b> {lvl_color} <i>({y['level']})</i>: {y['text']}</div>", unsafe_allow_html=True)
            with c_y3:
                if st.button("✏️ Sửa", key=f"btn_edit_yccd_{idx}"):
                    curr_val = st.session_state.get("editing_yccd_idx")
                    st.session_state["editing_yccd_idx"] = idx if curr_val != idx else None
                    st.rerun()
            with c_y4:
                if st.button("🗑️ Xóa", key=f"btn_del_yccd_{idx}"):
                    yccds.pop(idx)
                    st.success(f"Đã xóa YCCĐ #{idx+1}")
                    st.rerun()

            # Inline form chỉnh sửa YCCĐ
            if st.session_state.get("editing_yccd_idx") == idx:
                with st.form(key=f"form_edit_yccd_{idx}"):
                    st.markdown(f"**📝 Chỉnh sửa YCCĐ #{idx+1}:**")
                    all_codes = list(COMPETENCY_INDICATORS.keys())
                    curr_c_idx = all_codes.index(y['code']) if y['code'] in all_codes else 0
                    e_code = st.selectbox("Mã chỉ báo năng lực:", all_codes, index=curr_c_idx, key=f"edit_code_{idx}")
                    e_text = st.text_area("Nội dung YCCĐ (Trích nguyên văn SGK):", value=y['text'], height=90, key=f"edit_text_{idx}")
                    
                    if st.form_submit_button("💾 Cập nhật YCCĐ này"):
                        if e_text.strip():
                            y['code'] = e_code
                            y['level'] = COMPETENCY_INDICATORS[e_code]['level']
                            y['text'] = e_text.strip()
                            st.session_state["editing_yccd_idx"] = None
                            st.success(f"✅ Đã cập nhật YCCĐ #{idx+1}!")
                            st.rerun()

    # NÚT KHÔI PHỤC CSDL CHƯƠNG TRÌNH GỐC NẾU CÓ BẤT KỲ LỖI TỪ NGỮ NÀO
    with st.expander("🔄 Khôi phục YCCĐ chuẩn GDPT 2018 gốc (Dùng khi cần xóa các chỉnh sửa lỗi)"):
        st.warning("⚠️ Nút này sẽ khôi phục lại toàn bộ 437 YCCĐ của 123 bài học về chuẩn nguyên bản của Bộ GD&ĐT.")
        if st.button("🔄 XÁC NHẬN KHÔI PHỤC TẤT CẢ YCCĐ GỐC", key="btn_reset_curr_full"):
            st.session_state.curriculum = DEFAULT_CURRICULUM
            st.success("✅ Đã khôi phục thành công 100% cơ sở dữ liệu YCCĐ gốc!")
            st.rerun()

    # FORM BỔ SUNG YCCĐ MỚI
    with st.expander("➕ Bổ sung Yêu cầu cần đạt (YCCĐ) mới cho bài học này"):
        with st.form("form_add_yccd"):
            new_code = st.selectbox("Chọn Mã chỉ báo năng lực:", list(COMPETENCY_INDICATORS.keys()))
            new_text = st.text_area("Nội dung Yêu cầu cần đạt (Trích nguyên văn SGK):")
            if st.form_submit_button("💾 Lưu YCCĐ mới"):
                if new_text.strip():
                    lvl = COMPETENCY_INDICATORS[new_code]["level"]
                    st.session_state.curriculum[selected_book][selected_topic][selected_lesson].append({
                        "code": new_code,
                        "level": lvl,
                        "text": new_text.strip()
                    })
                    st.success("✅ Đã thêm YCCĐ mới thành công!")
                    st.rerun()

    st.divider()

    # 4. NÚT ĐỒNG BỘ TẤT CẢ DỮ LIỆU SANG TẤT CẢ CÁC TAB
    st.subheader("🔄 4. Đồng Bộ Dữ Liệu Thay Đổi Sang Tất Cả Các Tab")
    st.caption("Sau khi sửa đổi tên bài học, tên chương hoặc YCCĐ ở Tab 1, bấm nút dưới đây để đồng bộ toàn bộ tên gọi và YCCĐ sang Ma trận đề (Tab 2), Bản đặc tả (Tab 3) và Ngân hàng câu hỏi (Tab 4).")

    if st.button("🔄 BẤM VÀO ĐÂY ĐỂ ĐỒNG BỘ DỮ LIỆU SANG TẤT CẢ CÁC TAB", type="primary", use_container_width=True, key="btn_sync_all_tabs"):
        sync_matrix_count = 0
        sync_q_count = 0
        
        # Đồng bộ lịch sử đổi tên
        rename_history = st.session_state.get("rename_history", [])
        for item in rename_history:
            if item["type"] == "lesson":
                bk, old_n, new_n = item["book"], item["old_name"], item["new_name"]
                # Cập nhật trong matrix_rows
                for r in st.session_state.matrix_rows:
                    if r.get("book") == bk and r.get("lesson") == old_n:
                        r["lesson"] = new_n
                        sync_matrix_count += 1
                # Cập nhật trong questions
                for q in st.session_state.questions:
                    if q.get("book") == bk and q.get("lesson") == old_n:
                        q["lesson"] = new_n
                        sync_q_count += 1
            elif item["type"] == "topic":
                bk, old_n, new_n = item["book"], item["old_name"], item["new_name"]
                for r in st.session_state.matrix_rows:
                    if r.get("book") == bk and r.get("topic") == old_n:
                        r["topic"] = new_n
                        sync_matrix_count += 1
                for q in st.session_state.questions:
                    if q.get("book") == bk and q.get("topic") == old_n:
                        q["topic"] = new_n
                        sync_q_count += 1

        # Xóa lịch sử sau khi đồng bộ
        st.session_state.rename_history = []
        
        st.success(f"""
        ✅ **ĐÃ ĐỒNG BỘ THÀNH CÔNG TẤT CẢ DỮ LIỆU TỪ TAB 1!**
        - 📋 **Ma trận đề (Tab 2):** Đã cập nhật {sync_matrix_count} thông tin tên bài/chủ đề trong ma trận.
        - 📑 **Bản đặc tả (Tab 3):** Đã làm mới tự động theo YCCĐ và tên bài mới từ Tab 1.
        - 📚 **Ngân hàng câu hỏi (Tab 4):** Đã cập nhật {sync_q_count} câu hỏi theo tên bài học mới.
        """)

# -------------------------------------------------------------
# TAB 2: THIẾT KẾ MA TRẬN ĐỀ KIỂM TRA
# -------------------------------------------------------------
with tab2:
    st.header("2. Thiết Kế Ma Trận Đề Kiểm Tra Chuẩn GDPT 2018")
    
    with st.expander("➕ Thêm Nội dung Bài học vào Ma trận", expanded=True):
        c1, c2, c3 = st.columns(3)
        with c1:
            m_book = st.selectbox("Sách / Khối:", list(st.session_state.curriculum.keys()), key="m_bk")
        with c2:
            m_topic = st.selectbox("Chương / Chủ đề:", list(st.session_state.curriculum[m_book].keys()), key="m_tp")
        with c3:
            m_lesson = st.selectbox("Bài học:", list(st.session_state.curriculum[m_book][m_topic].keys()), key="m_ls")
        
        if st.button("➕ Đưa bài học này vào Ma trận", type="primary"):
            new_id = max([r["id"] for r in st.session_state.matrix_rows], default=0) + 1
            st.session_state.matrix_rows.append({
                "id": new_id,
                "book": m_book,
                "topic": m_topic,
                "lesson": m_lesson,
                "mcq_b": 1, "mcq_h": 1, "mcq_vd": 0,
                "tf_b": 0, "tf_h": 1, "tf_vd": 0,
                "sa_b": 0, "sa_h": 0, "sa_vd": 0,
                "es_b": 0, "es_h": 0, "es_vd": 0
            })
            st.rerun()

    st.subheader("📋 Bảng Phân Bổ Số Câu Hỏi Trong Ma Trận")
    for r_idx, r in enumerate(st.session_state.matrix_rows):
        with st.container():
            st.markdown(f"**Bài #{r['id']}: {r['lesson']}** *({r['topic']})*")
            
            mc1, mc2, mc3, mc4, mc5 = st.columns([3, 2, 2, 2, 1])
            with mc1:
                st.caption("1. TNKQ Nhiều Lựa Chọn (0.25đ/câu):")
                r["mcq_b"] = st.number_input("Biết", min_value=0, max_value=20, value=r["mcq_b"], key=f"mb_{r['id']}")
                r["mcq_h"] = st.number_input("Hiểu", min_value=0, max_value=20, value=r["mcq_h"], key=f"mh_{r['id']}")
                r["mcq_vd"] = st.number_input("VD", min_value=0, max_value=20, value=r["mcq_vd"], key=f"mvd_{r['id']}")
            with mc2:
                st.caption("2. TNKQ Đúng/Sai (1.0đ/câu 4 ý):")
                r["tf_b"] = st.number_input("Biết", min_value=0, max_value=10, value=r["tf_b"], key=f"tb_{r['id']}")
                r["tf_h"] = st.number_input("Hiểu", min_value=0, max_value=10, value=r["tf_h"], key=f"th_{r['id']}")
                r["tf_vd"] = st.number_input("VD", min_value=0, max_value=10, value=r["tf_vd"], key=f"tvd_{r['id']}")
            with mc3:
                st.caption("3. TNKQ Trả Lời Ngắn (0.25đ):")
                r["sa_b"] = st.number_input("Biết", min_value=0, max_value=10, value=r["sa_b"], key=f"sb_{r['id']}")
                r["sa_h"] = st.number_input("Hiểu", min_value=0, max_value=10, value=r["sa_h"], key=f"sh_{r['id']}")
                r["sa_vd"] = st.number_input("VD", min_value=0, max_value=10, value=r["sa_vd"], key=f"svd_{r['id']}")
            with mc4:
                st.caption("4. Tự Luận (1.0đ/câu):")
                r["es_b"] = st.number_input("Biết", min_value=0, max_value=5, value=r["es_b"], key=f"eb_{r['id']}")
                r["es_h"] = st.number_input("Hiểu", min_value=0, max_value=5, value=r["es_h"], key=f"eh_{r['id']}")
                r["es_vd"] = st.number_input("VD", min_value=0, max_value=5, value=r["es_vd"], key=f"evd_{r['id']}")
            with mc5:
                st.write("")
                st.write("")
                if st.button("🗑️ Xóa", key=f"del_m_{r['id']}"):
                    st.session_state.matrix_rows.pop(r_idx)
                    st.rerun()

    m_stats = calculate_matrix_stats(st.session_state.matrix_rows)
    st.divider()
    st.subheader("📊 Đo Lường Tỉ Lệ % & Điểm Số Tự Động")
    
    st_col1, st_col2, st_col3, st_col4 = st.columns(4)
    st_col1.metric("🟢 Nhận biết (Chuẩn 40%)", f"{m_stats['pct_b']}%", f"{m_stats['score_b']} điểm")
    st_col2.metric("🔵 Thông hiểu (Chuẩn 30%)", f"{m_stats['pct_h']}%", f"{m_stats['score_h']} điểm")
    st_col3.metric("🔴 Vận dụng (Chuẩn 30%)", f"{m_stats['pct_vd']}%", f"{m_stats['score_vd']} điểm")
    st_col4.metric("🏆 Tổng điểm kiểm tra", f"{m_stats['total_score']} / 10.0đ", f"{m_stats['total_items']} lệnh hỏi")

    col_btn_m1, col_btn_m2 = st.columns(2)
    with col_btn_m1:
        excel_matrix = export_matrix_to_excel(st.session_state.matrix_rows)
        st.download_button(
            label="📥 Xuất Ma Trận ra Excel (.xlsx) Kẻ Bảng Chuẩn",
            data=excel_matrix,
            file_name="Ma_Tran_De_Kiem_Tra_SinhHoc.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )
    with col_btn_m2:
        if st.button("🎲 TỰ ĐỘNG TẠO ĐỀ THI NGẪU NHIÊN NGAY TỪ MA TRẬN NÀY", type="primary", use_container_width=True):
            exam_data, warnings = generate_random_exam(st.session_state.matrix_rows, st.session_state.questions)
            st.session_state.current_generated_exam = exam_data
            st.session_state.exam_warnings = warnings
            st.success("🎉 Đã tạo Đề thi ngẫu nhiên thành công! Vui lòng chuyển sang Tab 4 (mục 4.4) hoặc xem ở bên dưới.")

# -------------------------------------------------------------
# TAB 3: THIẾT KẾ BẢN ĐẶC TẢ
# -------------------------------------------------------------
with tab3:
    st.header("3. Bản Đặc Tả Đề Kiểm Tra Đánh Giá GDPT 2018")
    st.caption("Tự động đồng bộ YCCĐ theo mã chỉ báo năng lực và liên kết số lượng câu từ Tab 2")

    spec_data = []
    for r in st.session_state.matrix_rows:
        book, topic, lesson = r.get("book",""), r.get("topic",""), r.get("lesson","")
        y_list = st.session_state.curriculum.get(book,{}).get(topic,{}).get(lesson,[])
        
        y_nb = [f"• [{y['code']}] {y['text']}" for y in y_list if y.get('level')=="Nhận biết"]
        y_th = [f"• [{y['code']}] {y['text']}" for y in y_list if y.get('level')=="Thông hiểu"]
        y_vd = [f"• [{y['code']}] {y['text']}" for y in y_list if y.get('level')=="Vận dụng"]

        spec_data.append({
            "TT": r["id"],
            "Chương / Chủ đề": topic,
            "Nội dung kiến thức": lesson,
            "YCCĐ Nhận biết": "\n".join(y_nb) if y_nb else "• Nhận biết các khái niệm cơ bản",
            "YCCĐ Thông hiểu": "\n".join(y_th) if y_th else "• Giải thích cơ chế, vai trò",
            "YCCĐ Vận dụng": "\n".join(y_vd) if y_vd else "• Vận dụng kiến thức thực tiễn",
            "Nhiều lựa chọn (B-H-VD)": f"{r['mcq_b']}-{r['mcq_h']}-{r['mcq_vd']}",
            "Đúng/Sai (B-H-VD)": f"{r['tf_b']}-{r['tf_h']}-{r['tf_vd']}",
            "Trả lời ngắn (B-H-VD)": f"{r['sa_b']}-{r['sa_h']}-{r['sa_vd']}",
            "Tự luận (B-H-VD)": f"{r['es_b']}-{r['es_h']}-{r['es_vd']}"
        })

    if spec_data:
        df_spec = pd.DataFrame(spec_data)
        st.dataframe(df_spec, use_container_width=True)

    col_sp1, col_sp2 = st.columns(2)
    with col_sp1:
        excel_spec = export_spec_to_excel(st.session_state.matrix_rows, st.session_state.curriculum)
        st.download_button(
            label="📥 Xuất Bản Đặc Tả ra Excel (.xlsx)",
            data=excel_spec,
            file_name="Ban_Dac_Ta_De_Kiem_Tra_SinhHoc.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )
    with col_sp2:
        doc_sp = docx.Document()
        doc_sp.add_heading("BẢN ĐẶC TẢ ĐỀ KIỂM TRA MÔN SINH HỌC GDPT 2018", level=1)
        table_sp = doc_sp.add_table(rows=1, cols=6)
        hdr_cells = table_sp.rows[0].cells
        hdr_cells[0].text = "TT"
        hdr_cells[1].text = "Chương / Chủ đề"
        hdr_cells[2].text = "Nội dung kiến thức"
        hdr_cells[3].text = "Yêu cầu cần đạt (Chỉ báo)"
        hdr_cells[4].text = "Trắc nghiệm (B-H-VD)"
        hdr_cells[5].text = "Tự luận"

        for row in spec_data:
            r_cells = table_sp.add_row().cells
            r_cells[0].text = str(row["TT"])
            r_cells[1].text = row["Chương / Chủ đề"]
            r_cells[2].text = row["Nội dung kiến thức"]
            r_cells[3].text = f"NHẬN BIẾT:\n{row['YCCĐ Nhận biết']}\n\nTHÔNG HIỂU:\n{row['YCCĐ Thông hiểu']}\n\nVẬN DỤNG:\n{row['YCCĐ Vận dụng']}"
            r_cells[4].text = f"NLC: {row['Nhiều lựa chọn (B-H-VD)']}\nĐS: {row['Đúng/Sai (B-H-VD)']}\nNgắn: {row['Trả lời ngắn (B-H-VD)']}"
            r_cells[5].text = row["Tự luận (B-H-VD)"]

        buf_doc_sp = io.BytesIO()
        doc_sp.save(buf_doc_sp)
        buf_doc_sp.seek(0)
        st.download_button(
            label="📄 Xuất Bản Đặc Tả ra Word (.docx)",
            data=buf_doc_sp,
            file_name="Ban_Dac_Ta_De_Kiem_Tra_SinhHoc.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )

# -------------------------------------------------------------
# TAB 4: NGÂN HÀNG CÂU HỎI & TỰ ĐỘNG TẠO ĐỀ THI NGẪU NHIÊN
# -------------------------------------------------------------
with tab4:
    st.header("4. Biên Soạn & Quản Lý Ngân Hàng Câu Hỏi - Tạo Đề Thi Ngẫu Nhiên")
    st.caption("Xây dựng câu hỏi đa dạng thức, gán ảnh minh họa cho CÂU HỎI và TỪNG ĐÁP ÁN (A, B, C, D); tự động sinh Đề thi ngẫu nhiên theo Ma trận")

    subtab1, subtab2, subtab3, subtab4 = st.tabs([
        "➕ 4.1 Soạn Câu Hỏi Mới (Gán Ảnh Đáp Án A, B, C, D)",
        "📚 4.2 Quản Lý & Sửa / Xóa Trực Tiếp",
        "📥 4.3 Import / Export Excel & Word",
        "🎲 4.4 Tự Động Tạo Đề Thi Ngẫu Nhiên"
    ])

    # SUBTAB 4.1: SOẠN CÂU HỎI MỚI (TÍCH HỢP ẢNH CHO TỪNG ĐÁP ÁN)
    with subtab1:
        st.subheader("➕ Biên soạn câu hỏi mới trực tiếp vào Ngân hàng")

        c_q1, c_q2, c_q3 = st.columns(3)
        with c_q1:
            q_book = st.selectbox("Khối / Sách:", list(st.session_state.curriculum.keys()), key="qb_add")
        with c_q2:
            q_topic = st.selectbox("Chương / Chủ đề:", list(st.session_state.curriculum[q_book].keys()), key="qt_add")
        with c_q3:
            q_lesson = st.selectbox("Bài học:", list(st.session_state.curriculum[q_book][q_topic].keys()), key="ql_add")

        c_q4, c_q5, c_q6 = st.columns(3)
        with c_q4:
            q_type = st.selectbox("Dạng thức câu hỏi:", [
                "TNKQ nhiều lựa chọn",
                "TNKQ đúng - sai",
                "TNKQ trả lời ngắn",
                "Tự luận"
            ], key="qtype_add")
        with c_q5:
            q_level = st.selectbox("Mức độ nhận thức:", ["Nhận biết", "Thông hiểu", "Vận dụng"], key="qlevel_add")
        with c_q6:
            q_ind = st.selectbox("Mã chỉ báo năng lực:", list(COMPETENCY_INDICATORS.keys()),
                                 format_func=lambda x: f"{x} - {COMPETENCY_INDICATORS[x]['name'][:40]}...", key="qind_add")

        q_stem_str, q_ans_str, q_exp_str = "", "A", ""
        opt_a, opt_b, opt_c, opt_d = "", "", "", ""
        opt_a_img_b64, opt_b_img_b64, opt_c_img_b64, opt_d_img_b64 = "", "", "", ""
        tf_a, tf_b, tf_c, tf_d = "", "", "", ""
        tf_ans_a, tf_ans_b, tf_ans_c, tf_ans_d = "Đúng", "Đúng", "Đúng", "Đúng"

        if q_type == "TNKQ nhiều lựa chọn":
            st.markdown("##### 📝 Biên soạn Câu hỏi Trắc nghiệm 4 lựa chọn (MCQ)")
            q_stem_str = st.text_area("1. Nội dung câu hỏi (Câu dẫn):", height=90, placeholder="Nhập câu dẫn...", key="mcq_stem_add")
            
            st.markdown("**2. Nhập nội dung văn bản và/hoặc tải TỆP ẢNH cho từng phương án (A, B, C, D):**")
            col_a1, col_a2 = st.columns([2, 1])
            opt_a = col_a1.text_input("Phương án A (Văn bản):", placeholder="Nội dung A...", key="mcq_opta_add")
            f_img_a = col_a2.file_uploader("🖼️ Tải ảnh cho A:", type=["png","jpg","jpeg","svg"], key="img_a_add")

            col_b1, col_b2 = st.columns([2, 1])
            opt_b = col_b1.text_input("Phương án B (Văn bản):", placeholder="Nội dung B...", key="mcq_optb_add")
            f_img_b = col_b2.file_uploader("🖼️ Tải ảnh cho B:", type=["png","jpg","jpeg","svg"], key="img_b_add")

            col_c1, col_c2 = st.columns([2, 1])
            opt_c = col_c1.text_input("Phương án C (Văn bản):", placeholder="Nội dung C...", key="mcq_optc_add")
            f_img_c = col_c2.file_uploader("🖼️ Tải ảnh cho C:", type=["png","jpg","jpeg","svg"], key="img_c_add")

            col_d1, col_d2 = st.columns([2, 1])
            opt_d = col_d1.text_input("Phương án D (Văn bản):", placeholder="Nội dung D...", key="mcq_optd_add")
            f_img_d = col_d2.file_uploader("🖼️ Tải ảnh cho D:", type=["png","jpg","jpeg","svg"], key="img_d_add")

            mcq_ans = st.radio("3. Đáp án ĐÚNG mặc định:", ["A", "B", "C", "D"], index=0, horizontal=True, key="mcq_ans_add")
            q_ans_str = mcq_ans
            q_exp_str = st.text_area("4. Lời giải thích chi tiết:", height=70, key="mcq_exp_add")

            for img_f, var_name in [(f_img_a, "a"), (f_img_b, "b"), (f_img_c, "c"), (f_img_d, "d")]:
                if img_f is not None:
                    b_d = img_f.getvalue()
                    m_t = img_f.type or "image/png"
                    b64_val = f"data:{m_t};base64,{base64.b64encode(b_d).decode('utf-8')}"
                    if var_name == "a": opt_a_img_b64 = b64_val
                    elif var_name == "b": opt_b_img_b64 = b64_val
                    elif var_name == "c": opt_c_img_b64 = b64_val
                    elif var_name == "d": opt_d_img_b64 = b64_val

        elif q_type == "TNKQ đúng - sai":
            st.markdown("##### 📝 Biên soạn Câu hỏi Trắc nghiệm Đúng - Sai (4 lệnh hỏi a, b, c, d)")
            q_stem_str = st.text_area("1. Bối cảnh / Đoạn văn dẫn / Bảng số liệu câu hỏi:", height=90, key="tf_stem_add")

            t_row1_1, t_row1_2 = st.columns([4, 1.2])
            tf_a = t_row1_1.text_input("a)", placeholder="Mệnh đề a...", key="tfa_add")
            tf_ans_a = t_row1_2.radio("Đáp án a:", ["Đúng", "Sai"], horizontal=True, key="tfansa_add")

            t_row2_1, t_row2_2 = st.columns([4, 1.2])
            tf_b = t_row2_1.text_input("b)", placeholder="Mệnh đề b...", key="tfb_add")
            tf_ans_b = t_row2_2.radio("Đáp án b:", ["Đúng", "Sai"], horizontal=True, key="tfansb_add")

            t_row3_1, t_row3_2 = st.columns([4, 1.2])
            tf_c = t_row3_1.text_input("c)", placeholder="Mệnh đề c...", key="tfc_add")
            tf_ans_c = t_row3_2.radio("Đáp án c:", ["Đúng", "Sai"], horizontal=True, key="tfansc_add")

            t_row4_1, t_row4_2 = st.columns([4, 1.2])
            tf_d = t_row4_1.text_input("d)", placeholder="Mệnh đề d...", key="tfd_add")
            tf_ans_d = t_row4_2.radio("Đáp án d:", ["Đúng", "Sai"], horizontal=True, key="tfansd_add")

            q_ans_str = f"a) {tf_ans_a}, b) {tf_ans_b}, c) {tf_ans_c}, d) {tf_ans_d}"
            q_exp_str = st.text_area("3. Lời giải thích chi tiết:", height=70, key="tf_exp_add")

        elif q_type == "TNKQ trả lời ngắn":
            q_stem_str = st.text_area("1. Nội dung câu hỏi / Yêu cầu điền kết quả:", height=90, key="sa_stem_add")
            q_ans_str = st.text_input("2. Đáp án chuẩn (Từ/Số ngắn):", key="sa_ans_add")
            q_exp_str = st.text_area("3. Lời giải thích chi tiết:", height=70, key="sa_exp_add")

        else: # Tự luận
            q_stem_str = st.text_area("1. Đề bài / Nội dung câu hỏi tự luận:", height=110, key="es_stem_add")
            q_ans_str = st.text_area("2. Hướng dẫn chấm / Thang điểm:", height=90, key="es_ans_add")
            q_exp_str = st.text_area("3. Lời giải chi tiết bổ sung:", height=70, key="es_exp_add")

        col_img1, col_img2 = st.columns(2)
        with col_img1:
            q_img_file = st.file_uploader("🖼️ Ô nhập ảnh chung cho CÂU HỎI (Sơ đồ, đồ thị...):", type=["png", "jpg", "jpeg", "svg"], key="qimg_add")
        with col_img2:
            a_img_file = st.file_uploader("🖼️ Ô nhập ảnh HƯỚNG DẪN GIẢI / ĐÁP ÁN:", type=["png", "jpg", "jpeg", "svg"], key="aimg_add")

        if st.button("💾 Lưu câu hỏi này vào Ngân hàng", type="primary", use_container_width=True):
            if q_stem_str.strip():
                q_img_b64 = ""
                if q_img_file is not None:
                    b_d = q_img_file.getvalue()
                    q_img_b64 = f"data:{q_img_file.type or 'image/png'};base64,{base64.b64encode(b_d).decode('utf-8')}"

                a_img_b64 = ""
                if a_img_file is not None:
                    b_d = a_img_file.getvalue()
                    a_img_b64 = f"data:{a_img_file.type or 'image/png'};base64,{base64.b64encode(b_d).decode('utf-8')}"

                new_id = max([q["id"] for q in st.session_state.questions], default=0) + 1
                st.session_state.questions.append({
                    "id": new_id,
                    "book": q_book,
                    "topic": q_topic,
                    "lesson": q_lesson,
                    "qtype": q_type,
                    "level": q_level,
                    "indicator": q_ind,
                    "content": q_stem_str.strip(),
                    "opt_a": opt_a.strip(), "opt_a_img": opt_a_img_b64,
                    "opt_b": opt_b.strip(), "opt_b_img": opt_b_img_b64,
                    "opt_c": opt_c.strip(), "opt_c_img": opt_c_img_b64,
                    "opt_d": opt_d.strip(), "opt_d_img": opt_d_img_b64,
                    "tf_a": tf_a.strip(), "tf_ans_a": tf_ans_a,
                    "tf_b": tf_b.strip(), "tf_ans_b": tf_ans_b,
                    "tf_c": tf_c.strip(), "tf_ans_c": tf_ans_c,
                    "tf_d": tf_d.strip(), "tf_ans_d": tf_ans_d,
                    "q_image": q_img_b64,
                    "answer": q_ans_str.strip(),
                    "a_image": a_img_b64,
                    "explanation": q_exp_str.strip()
                })
                st.success(f"✅ Đã lưu thành công Câu hỏi mã #{new_id} vào Ngân hàng!")
                st.rerun()
            else:
                st.error("Vui lòng nhập nội dung câu hỏi!")

    # SUBTAB 4.2: QUẢN LÝ & CHỈNH SỬA / XÓA CÂU HỎI TRỰC TIẾP
    with subtab2:
        st.subheader("📚 Quản lý & Chỉnh sửa / Xóa câu hỏi trực tiếp")

        if not st.session_state.questions:
            st.info("Hiện chưa có câu hỏi nào trong ngân hàng.")
        else:
            f_col1, f_col2, f_col3, f_col4 = st.columns(4)
            with f_col1:
                filter_book = st.selectbox("Lọc theo Sách:", ["Tất cả"] + list(st.session_state.curriculum.keys()), key="f_book")
            with f_col2:
                filter_type = st.selectbox("Lọc Dạng thức:", ["Tất cả", "TNKQ nhiều lựa chọn", "TNKQ đúng - sai", "TNKQ trả lời ngắn", "Tự luận"], key="f_type")
            with f_col3:
                filter_level = st.selectbox("Lọc Mức độ:", ["Tất cả", "Nhận biết", "Thông hiểu", "Vận dụng"], key="f_level")
            with f_col4:
                search_kw = st.text_input("🔍 Tìm từ khóa:", "", key="f_kw")

            filtered_q = []
            for q in st.session_state.questions:
                if filter_book != "Tất cả" and q.get("book") != filter_book: continue
                if filter_type != "Tất cả" and q.get("qtype") != filter_type: continue
                if filter_level != "Tất cả" and q.get("level") != filter_level: continue
                if search_kw.strip():
                    kw = search_kw.strip().lower()
                    corpus = (q.get("content","") + " " + q.get("answer","") + " " + q.get("lesson","")).lower()
                    if kw not in corpus: continue
                filtered_q.append(q)

            st.caption(f"Hiển thị **{len(filtered_q)}** / **{len(st.session_state.questions)}** câu hỏi")

            if "editing_q_id" not in st.session_state:
                st.session_state.editing_q_id = None

            for q in filtered_q:
                qid = q["id"]
                st.markdown("---")
                
                b1, b2, b3, b4, b5 = st.columns([1, 3, 2, 2, 2])
                with b1: st.markdown(f"**Mã #{qid}**")
                with b2: st.caption(f"📖 {q.get('book','')} | {q.get('lesson','')[:30]}...")
                with b3: st.markdown(f"🏷️ **{q.get('qtype','')}**")
                with b4:
                    lvl = q.get('level','')
                    badge = "🟢" if lvl=="Nhận biết" else ("🔵" if lvl=="Thông hiểu" else "🔴")
                    st.markdown(f"{badge} **{lvl}** ({q.get('indicator','')})")
                with b5:
                    ec1, ec2 = st.columns(2)
                    if ec1.button("✏️ Sửa", key=f"btn_e_{qid}"):
                        st.session_state.editing_q_id = qid if st.session_state.editing_q_id != qid else None
                        st.rerun()
                    if ec2.button("🗑️ Xóa", key=f"btn_d_{qid}"):
                        st.session_state.questions = [x for x in st.session_state.questions if x["id"] != qid]
                        st.success(f"Đã xóa #{qid}")
                        st.rerun()

                st.markdown(f"**Nội dung:** {q.get('content','')}")
                if q.get("q_image"):
                    st.image(q["q_image"], caption="Ảnh minh họa câu hỏi", width=300)

                qtype_curr = q.get("qtype","")
                if qtype_curr == "TNKQ nhiều lựa chọn":
                    o_a, o_b, o_c, o_d = q.get("opt_a",""), q.get("opt_b",""), q.get("opt_c",""), q.get("opt_d","")
                    st.markdown(f"• **A.** {o_a}")
                    if q.get("opt_a_img"): st.image(q["opt_a_img"], width=180, caption="Ảnh Phương án A")
                    st.markdown(f"• **B.** {o_b}")
                    if q.get("opt_b_img"): st.image(q["opt_b_img"], width=180, caption="Ảnh Phương án B")
                    st.markdown(f"• **C.** {o_c}")
                    if q.get("opt_c_img"): st.image(q["opt_c_img"], width=180, caption="Ảnh Phương án C")
                    st.markdown(f"• **D.** {o_d}")
                    if q.get("opt_d_img"): st.image(q["opt_d_img"], width=180, caption="Ảnh Phương án D")

                st.markdown(f"🎯 **Đáp án:** `{q.get('answer','')}`")

                if st.session_state.editing_q_id == qid:
                    with st.form(key=f"form_edit_{qid}"):
                        st.info(f"📝 **Đang chỉnh sửa câu hỏi #{qid}:**")
                        e_stem = st.text_area("Nội dung câu hỏi:", value=q.get("content",""), height=90, key=f"ecnt_{qid}")
                        
                        if qtype_curr == "TNKQ nhiều lựa chọn":
                            st.markdown("**Sửa văn bản & Ảnh cho từng phương án A, B, C, D:**")
                            e_a_col1, e_a_col2 = st.columns([2, 1])
                            e_opt_a = e_a_col1.text_input("Phương án A:", value=q.get("opt_a",""), key=f"eopta_{qid}")
                            e_f_a = e_a_col2.file_uploader("🖼️ Sửa ảnh A:", type=["png","jpg","jpeg","svg"], key=f"eimga_{qid}")

                            e_b_col1, e_b_col2 = st.columns([2, 1])
                            e_opt_b = e_b_col1.text_input("Phương án B:", value=q.get("opt_b",""), key=f"eoptb_{qid}")
                            e_f_b = e_b_col2.file_uploader("🖼️ Sửa ảnh B:", type=["png","jpg","jpeg","svg"], key=f"eimgb_{qid}")

                            e_c_col1, e_c_col2 = st.columns([2, 1])
                            e_opt_c = e_c_col1.text_input("Phương án C:", value=q.get("opt_c",""), key=f"eoptc_{qid}")
                            e_f_c = e_c_col2.file_uploader("🖼️ Sửa ảnh C:", type=["png","jpg","jpeg","svg"], key=f"eimgc_{qid}")

                            e_d_col1, e_d_col2 = st.columns([2, 1])
                            e_opt_d = e_d_col1.text_input("Phương án D:", value=q.get("opt_d",""), key=f"eoptd_{qid}")
                            e_f_d = e_d_col2.file_uploader("🖼️ Sửa ảnh D:", type=["png","jpg","jpeg","svg"], key=f"eimgd_{qid}")

                            mcq_opts = ["A", "B", "C", "D"]
                            curr_mcq_ans = q.get("answer","A")
                            c_idx = mcq_opts.index(curr_mcq_ans) if curr_mcq_ans in mcq_opts else 0
                            e_ans_val = st.radio("Đáp án ĐÚNG:", mcq_opts, index=c_idx, horizontal=True, key=f"eans_radio_{qid}")

                        else:
                            e_ans_val = st.text_input("Đáp án:", value=q.get("answer",""), key=f"eans_{qid}")

                        e_exp_val = st.text_area("Lời giải thích:", value=q.get("explanation",""), height=70, key=f"eexp_{qid}")

                        if st.form_submit_button("💾 Cập nhật thay đổi"):
                            q["content"] = e_stem.strip()
                            q["answer"] = e_ans_val.strip()
                            q["explanation"] = e_exp_val.strip()
                            
                            if qtype_curr == "TNKQ nhiều lựa chọn":
                                q["opt_a"], q["opt_b"], q["opt_c"], q["opt_d"] = e_opt_a, e_opt_b, e_opt_c, e_opt_d
                                if e_f_a: q["opt_a_img"] = f"data:{e_f_a.type or 'image/png'};base64,{base64.b64encode(e_f_a.getvalue()).decode('utf-8')}"
                                if e_f_b: q["opt_b_img"] = f"data:{e_f_b.type or 'image/png'};base64,{base64.b64encode(e_f_b.getvalue()).decode('utf-8')}"
                                if e_f_c: q["opt_c_img"] = f"data:{e_f_c.type or 'image/png'};base64,{base64.b64encode(e_f_c.getvalue()).decode('utf-8')}"
                                if e_f_d: q["opt_d_img"] = f"data:{e_f_d.type or 'image/png'};base64,{base64.b64encode(e_f_d.getvalue()).decode('utf-8')}"

                            st.session_state.editing_q_id = None
                            st.success("Đã cập nhật câu hỏi!")
                            st.rerun()

    # SUBTAB 4.3: IMPORT / EXPORT EXCEL & WORD
    with subtab3:
        st.subheader("📥 Import / Export Ngân Hàng Câu Hỏi & Tải File Mẫu")

        c_ex1, c_ex2 = st.columns(2)
        with c_ex1:
            st.markdown("##### 1. Tải Mẫu Excel Nhập Câu Hỏi chuẩn 24 cột")
            excel_template = generate_question_excel_template()
            st.download_button(
                label="📥 Tải Mẫu Excel Soạn Câu Hỏi (.xlsx)",
                data=excel_template,
                file_name="Mau_Soan_Cau_Hoi_SinhHoc_2018.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )

            st.divider()
            st.markdown("##### 2. Xuất Toàn bộ Ngân hàng câu hỏi ra Word/Excel")
            buf_q_excel = export_questions_to_excel(st.session_state.questions)
            st.download_button(
                label="📥 Xuất Ngân Hàng ra Excel (.xlsx)",
                data=buf_q_excel,
                file_name="Ngan_Hang_Cau_Hoi_SinhHoc.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )

            buf_q_word = export_questions_to_word(st.session_state.questions)
            st.download_button(
                label="📄 Xuất Ngân Hàng ra Word (.docx)",
                data=buf_q_word,
                file_name="Ngan_Hang_Cau_Hoi_SinhHoc.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )

        with c_ex2:
            st.markdown("##### 3. Nạp (Import) Câu hỏi từ File Excel")
            uploaded_file = st.file_uploader("Chọn file Excel câu hỏi (.xlsx):", type=["xlsx"], key="imp_q_file")
            if uploaded_file is not None:
                try:
                    df_imp = pd.read_excel(uploaded_file)
                    imported_records = df_imp.to_dict(orient="records")
                    max_id = max([q["id"] for q in st.session_state.questions], default=0)
                    for idx, rec in enumerate(imported_records):
                        rec["id"] = max_id + idx + 1
                        st.session_state.questions.append(rec)
                    st.success(f"✅ Đã nạp thành công {len(imported_records)} câu hỏi vào ngân hàng!")
                    st.rerun()
                except Exception as e:
                    st.error(f"Lỗi khi đọc file Excel: {e}")

            st.divider()
            if st.button("🔄 Khôi phục Ngân hàng mẫu ban đầu"):
                st.session_state.questions = DEFAULT_QUESTIONS_DATABASE
                st.success("Đã khôi phục ngân hàng câu hỏi mẫu!")
                st.rerun()

    # SUBTAB 4.4: TỰ ĐỘNG TẠO ĐỀ THI NGẪU NHIÊN THEO MA TRẬN
    with subtab4:
        st.subheader("🎲 Tự Động Tạo Đề Thi Ngẫu Nhiên Từ Ngân Hàng Theo Ma Trận")
        st.caption("Quét Ma trận hiện tại ở Tab 2, chọn ngẫu nhiên các câu hỏi từ Ngân hàng khớp đúng Bài học + Dạng thức + Mức độ")

        eg_c1, eg_c2, eg_c3 = st.columns(3)
        with eg_c1:
            ex_school = st.text_input("Tên Trường / Sở:", value="TRƯỜNG THPT CHUYÊN/CHUẨN", key="ex_sch")
        with eg_c2:
            ex_title = st.text_input("Tên Đề thi:", value="ĐỀ KIỂM TRA GIỮA KỲ I MÔN SINH HỌC 12", key="ex_ttl")
        with eg_c3:
            ex_code = st.text_input("Mã đề thi:", value="101", key="ex_code")

        if st.button("🎲 TỰ ĐỘNG TẠO ĐỀ THI NGẪU NHIÊN BÂY GIỜ", type="primary", use_container_width=True):
            exam_data, warnings = generate_random_exam(st.session_state.matrix_rows, st.session_state.questions)
            st.session_state.current_generated_exam = exam_data
            st.session_state.exam_warnings = warnings
            st.success("🎉 Đã tạo Đề thi ngẫu nhiên thành công!")

        if "current_generated_exam" in st.session_state and st.session_state.current_generated_exam:
            exam_data = st.session_state.current_generated_exam
            warnings = st.session_state.get("exam_warnings", [])

            if warnings:
                for w in warnings:
                    st.warning(w)

            st.divider()
            st.markdown(f"### 📄 HIỂN THỊ ĐỀ THI VỪA TẠO (MÃ ĐỀ: {ex_code})")

            q_num = 1
            for part_key, part_title in [
                ("Part1_MCQ", "PHẦN I. Câu hỏi trắc nghiệm nhiều lựa chọn (0,25đ/câu)"),
                ("Part2_TF", "PHẦN II. Câu hỏi trắc nghiệm Đúng/Sai (1,0đ/câu 4 ý)"),
                ("Part3_SA", "PHẦN III. Câu hỏi trắc nghiệm trả lời ngắn"),
                ("Part4_ES", "PHẦN IV. Câu hỏi tự luận")
            ]:
                if exam_data[part_key]:
                    st.markdown(f"#### {part_title}")
                    for q in exam_data[part_key]:
                        st.markdown(f"**Câu {q_num}** *({q.get('level','')})*: {q.get('content','')}")
                        if q.get("q_image"):
                            st.image(q["q_image"], width=300)

                        if part_key == "Part1_MCQ":
                            o_a, o_b, o_c, o_d = q.get("opt_a",""), q.get("opt_b",""), q.get("opt_c",""), q.get("opt_d","")
                            st.markdown(f"• **A.** {o_a}")
                            if q.get("opt_a_img"): st.image(q["opt_a_img"], width=180)
                            st.markdown(f"• **B.** {o_b}")
                            if q.get("opt_b_img"): st.image(q["opt_b_img"], width=180)
                            st.markdown(f"• **C.** {o_c}")
                            if q.get("opt_c_img"): st.image(q["opt_c_img"], width=180)
                            st.markdown(f"• **D.** {o_d}")
                            if q.get("opt_d_img"): st.image(q["opt_d_img"], width=180)

                        elif part_key == "Part2_TF":
                            if q.get("tf_a"): st.markdown(f"a) {q.get('tf_a')}")
                            if q.get("tf_b"): st.markdown(f"b) {q.get('tf_b')}")
                            if q.get("tf_c"): st.markdown(f"c) {q.get('tf_c')}")
                            if q.get("tf_d"): st.markdown(f"d) {q.get('tf_d')}")

                        q_num += 1

            st.divider()
            with st.expander("🔑 ĐÁP ÁN VÀ HƯỚNG DẪN CHẤM CHI TIẾT"):
                all_q = exam_data["Part1_MCQ"] + exam_data["Part2_TF"] + exam_data["Part3_SA"] + exam_data["Part4_ES"]
                for i, q in enumerate(all_q, start=1):
                    st.markdown(f"**Câu {i}:** Đáp án `{q.get('answer','')}` — *Lời giải:* {q.get('explanation','')}")

            ex_col1, ex_col2 = st.columns(2)
            with ex_col1:
                word_exam_buf = export_exam_to_word(exam_data, ex_school, ex_title, "45 phút", ex_code)
                st.download_button(
                    label="📄 Tải Đề Thi & Đáp Án ra Word (.docx)",
                    data=word_exam_buf,
                    file_name=f"De_Thi_SinhHoc_MaDe_{ex_code}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True
                )

# -------------------------------------------------------------
# TAB 5: HƯỚNG DẪN SỬ DỤNG & TRA CỨU CHỈ BÁO NĂNG LỰC
# -------------------------------------------------------------
with tab5:
    st.header("5. Hướng Dẫn Sử Dụng & Bảng Tra Cứu Chỉ Báo Năng Lực")

    st.subheader("📌 Quy trình 4 bước biên soạn đề kiểm tra Sinh học:")
    st.markdown("""
    1. **Bước 1 (Chọn phạm vi):** Vào **Tab 1**, kiểm tra các bài học nằm trong giới hạn kiểm tra.
    2. **Bước 2 (Thiết kế Ma trận):** Sang **Tab 2**, thêm các bài học vào bảng ma trận, điều chỉnh số lượng câu hỏi.
    3. **Bước 3 (Xuất Bản đặc tả):** Sang **Tab 3**, kiểm tra các YCCĐ đã được tự động gắn mã chỉ báo năng lực tương ứng với số lượng câu ở Tab 2. Nhấn nút xuất file Excel hoặc Word.
    4. **Bước 4 (Tạo Đề thi ngẫu nhiên):** Sang **Subtab 4.4**, nhấn nút **`🎲 Tự động tạo Đề thi`** để phần mềm quét Ngân hàng câu hỏi và sinh Đề thi ngẫu nhiên kèm Đáp án chi tiết, sẵn sàng xuất ra file Word.
    """)

    st.divider()
    st.subheader("🔍 Bảng đối chiếu Chỉ báo Năng lực Sinh học (Chuẩn GDPT 2018)")
    ind_table = []
    for k, v in COMPETENCY_INDICATORS.items():
        ind_table.append({
            "Mã chỉ báo": k,
            "Nhóm năng lực": v["group"],
            "Mức độ nhận thức": v["level"],
            "Biểu hiện hành vi cụ thể": v["name"]
        })
    st.dataframe(pd.DataFrame(ind_table), use_container_width=True)


